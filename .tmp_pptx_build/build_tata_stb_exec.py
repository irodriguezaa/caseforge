#!/usr/bin/env python3
"""Reporte ejecutivo de ejecución TATA STB NA — 1 slide."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
from lxml import etree

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x1B, 0x2A, 0x49)
TEAL = RGBColor(0x0D, 0x8B, 0x7D)
TEAL2 = RGBColor(0x13, 0xA6, 0x96)
GOLD = RGBColor(0xE8, 0x9F, 0x1F)
SLATE = RGBColor(0x51, 0x5F, 0x75)
LINE = RGBColor(0xD0, 0xD6, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF4, 0xF8, 0xFB)
MUTED = RGBColor(0x6B, 0x7A, 0x8D)
BODY = RGBColor(0x3A, 0x47, 0x5C)
RED = RGBColor(0xC4, 0x44, 0x2E)

OUT = Path("/Users/rodriguezisr/Documents/QC/QC_TATA_STB_NA_Ejecutivo.pptx")

SECTIONS = [
    # section, casos, ok, fail, na
    ("Acceso", 8, 2, 0, 6),
    ("Perfiles", 7, 7, 0, 0),
    ("Navegación", 43, 34, 6, 3),
    ("Transacción", 20, 2, 0, 18),
    ("Reproducción", 32, 25, 3, 4),
    ("Rep. Canal Líneal", 63, 53, 2, 8),
    ("NPVR", 10, 7, 3, 0),
    ("TimeShift", 7, 7, 0, 0),
    ("Trickplay", 3, 3, 0, 0),
    ("Buscador", 10, 9, 0, 1),
    ("Social y Settings", 5, 3, 1, 1),
    ("Control Parental", 13, 13, 0, 0),
    ("Control Compras", 5, 5, 0, 0),
    ("APK Performance", 18, 17, 0, 1),
]

TICKETS = [
    "ATSCL-2740", "ATSCL-3025", "ATSCL-3029", "ATSCL-3027",
    "ATSCL-3008", "ATSCL-3007", "ATSCL-2245", "ATSCL-3010",
    "ATSCL-3009", "ATSCL-3087", "ATSCL-3101", "ATSCL-3095",
    "ATSCL-3096", "ATSCL-3094", "ATSCL-3092", "ATSCL-3093",
    "ATSCL-2840", "ATSCL-2825",
    "GCECL-1759", "GCECL-1815", "GCECL-1960", "GCECL-1761",
]


def set_run(run, text, size, bold=False, color=NAVY, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    if radius is not None:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    s.shadow.inherit = False
    return s


def add_tb(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor="t", pad=(0, 0, 0, 0), after=1):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(pad[0])
    tf.margin_right = Pt(pad[1])
    tf.margin_top = Pt(pad[2])
    tf.margin_bottom = Pt(pad[3])
    tf._txBody.bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(after)
        if isinstance(item[0], str):
            item = [item]
        for text, size, bold, color in item:
            run = p.add_run()
            set_run(run, text, size, bold, color)
    return box


def fill_shape_text(shape, paragraphs, align=PP_ALIGN.CENTER, anchor="m", margins=(6, 6, 6, 6)):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(margins[0])
    tf.margin_right = Pt(margins[1])
    tf.margin_top = Pt(margins[2])
    tf.margin_bottom = Pt(margins[3])
    tf._txBody.bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    for i, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        if isinstance(item[0], str):
            item = [item]
        for text, size, bold, color in item:
            run = p.add_run()
            set_run(run, text, size, bold, color)


def add_table(slide, left, top, width, height, headers, data):
    rows = 1 + len(data)
    cols = len(headers)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.first_row = True
    widths = [0.38, 0.13, 0.16, 0.13, 0.13, 0.07]  # fractions
    total = sum(widths)
    for i, w in enumerate(widths):
        table.columns[i].width = int(width * w / total)

    def style_cell(cell, text, header=False, fail=False, center=False):
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if center or header else PP_ALIGN.LEFT
        cell.text_frame.word_wrap = False
        run = p.add_run()
        set_run(run, text, 10 if not header else 10, True if header else False, WHITE if header else (RED if fail else BODY))
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        for tag in ("a:solidFill",):
            pass
        # fill
        fill = NAVY if header else (RGBColor(0xFD, 0xE8, 0xE4) if fail else WHITE)
        solid = etree.SubElement(tcPr, qn("a:solidFill"))
        srgb = etree.SubElement(solid, qn("a:srgbClr"))
        srgb.set("val", str(fill))
        # margins
        for m, v in (("marL", "72000"), ("marR", "72000"), ("marT", "36000"), ("marB", "36000")):
            tcPr.set(m, v)
        bodyPr = cell.text_frame._txBody.find(qn("a:bodyPr"))
        if bodyPr is not None:
            bodyPr.set("anchor", "ctr")

    for i, h in enumerate(headers):
        style_cell(table.cell(0, i), h, header=True, center=True)
    for r, row in enumerate(data, 1):
        fail_row = int(row[3]) > 0
        for c, val in enumerate(row):
            style_cell(table.cell(r, c), str(val), fail=fail_row and c == 3, center=(c > 0))
        # zebra
        if r % 2 == 0:
            for c in range(cols):
                if int(row[3]) == 0:
                    tcPr = table.cell(r, c)._tc.get_or_add_tcPr()
                    solid = etree.SubElement(tcPr, qn("a:solidFill"))
                    srgb = etree.SubElement(solid, qn("a:srgbClr"))
                    srgb.set("val", "F4F8FB")
    return table_shape


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    add_rect(s, 0, 0, Inches(0.11), SLIDE_H, TEAL)

    add_tb(s, Inches(0.42), Inches(0.16), Inches(12.5), Inches(0.32),
           [("Reporte Ejecutivo QC  ·  Regresivo TATA STB NA", 22, True, NAVY)])
    add_tb(s, Inches(0.42), Inches(0.48), Inches(12.5), Inches(0.24),
           [("244 casos  ·  202 ejecutados  ·  187 pasaron  ·  15 fallaron  ·  42 N/A", 13, False, SLATE)])
    add_rect(s, Inches(0.42), Inches(0.76), Inches(12.5), Pt(1.15), LINE)

    kpis = [
        ("244", "Casos", "en la matriz", TEAL),
        ("202", "Ejecutados", "excluye N/A", TEAL2),
        ("187", "Pasaron", "92.6% de ejecutados", TEAL),
        ("15", "Fallaron", "Falla persiste", RED),
        ("22", "Issues", "tickets únicos", GOLD),
    ]
    tw, th, gap = Inches(2.38), Inches(1.05), Inches(0.12)
    for i, (n, a, b, bg) in enumerate(kpis):
        x = Inches(0.42) + i * (tw + gap)
        sh = add_rect(s, x, Inches(0.90), tw, th, bg, radius=0.08)
        fill_shape_text(sh, [(n, 24, True, WHITE), (a, 12, True, WHITE), (b, 10, False, WHITE)])

    add_tb(s, Inches(0.42), Inches(2.08), Inches(7.3), Inches(0.26),
           [("Casos por Section", 13, True, NAVY)])
    headers = ["Section", "Casos", "Ejecutados", "OK", "Falla", "N/A"]
    data = []
    for sec, n, ok, fail, na in SECTIONS:
        data.append((sec, n, ok + fail, ok, fail, na))
    add_table(
        s,
        Inches(0.42), Inches(2.36), Inches(7.35), Inches(4.58),
        headers,
        data,
    )

    add_tb(s, Inches(8.00), Inches(2.08), Inches(4.9), Inches(0.26),
           [("Issues asociados a las fallas  ·  22 tickets", 13, True, NAVY)])
    add_rect(s, Inches(8.00), Inches(2.36), Inches(4.90), Inches(4.58), OFF, radius=0.05)

    col_a = TICKETS[:11]
    col_b = TICKETS[11:]
    add_tb(
        s, Inches(8.18), Inches(2.50), Inches(2.25), Inches(3.90),
        [(t, 13, False, BODY) for t in col_a],
        after=6,
    )
    add_tb(
        s, Inches(10.40), Inches(2.50), Inches(2.30), Inches(3.90),
        [(t, 13, False, BODY) for t in col_b],
        after=6,
    )
    add_tb(
        s, Inches(8.18), Inches(6.42), Inches(4.55), Inches(0.42),
        [("ATSCL-3010, ATSCL-3087 y ATSCL-2825 cerrados; seguimiento TRI-153023 y TRI-153524.", 10, False, MUTED)],
    )

    add_rect(s, Inches(0.42), Inches(7.12), Inches(12.5), Pt(1.1), TEAL)
    add_tb(s, Inches(0.42), Inches(7.16), Inches(12.5), Inches(0.26),
           [("Quality Control · Claro Video    ·    Fuente: NA_STB_TATA    ·    23 sep 2026", 10, False, MUTED)])

    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
