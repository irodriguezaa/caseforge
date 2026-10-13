#!/usr/bin/env python3
"""ATSCL IPTV backlog — 2 gráficos ejecutivos, paleta QC Agosto."""

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x1B, 0x2A, 0x49)
TEAL = RGBColor(0x0D, 0x8B, 0x7D)
TEAL2 = RGBColor(0x13, 0xA6, 0x96)
GOLD = RGBColor(0xE8, 0x9F, 0x1F)
SLATE = RGBColor(0x51, 0x5F, 0x75)
LINE = RGBColor(0xD0, 0xD6, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF4, 0xF8, 0xFB)
MUTED = RGBColor(0x6B, 0x7A, 0x8D)
BODY = RGBColor(0x3A, 0x47, 0x5C)
TRACK = RGBColor(0xE8, 0xEE, 0xF2)

OUT = Path("/Users/rodriguezisr/Documents/QC/ATSCL_IPTV_backlog_QC.pptx")
PREV = Path("/Users/rodriguezisr/Documents/qcpulse/.tmp_pptx_build/preview_atscl.png")

# --- data (sheet 1ZZDVTSkc3y-POZx-J2KDyW0_jY5u0ffJVjSSO7W0Mc0, 21 sep 2026) ---
ESTADO = [
    ("Analysis", 8, GOLD),
    ("To Develop", 4, SLATE),
    ("Validation", 6, TEAL),
    ("QA Validation", 2, TEAL2),
]
TEMAS = [
    ("RTP Valeria", 6, TEAL),
    ("Canales bloqueados / desbloqueo", 6, TEAL2),
    ("Audio y subtítulos pegados", 2, NAVY),
    ("HBO Max", 2, GOLD),
    ("Otros (Live TV, vCard, add-on)", 4, SLATE),
]


def set_run(run, text, size, bold=False, color=NAVY, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill, line=None, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.15)
    if radius is not None:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    s.shadow.inherit = False
    return s


def add_tb(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor="t", pad=(0, 0, 0, 0)):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(pad[0])
    tf.margin_right = Pt(pad[1])
    tf.margin_top = Pt(pad[2])
    tf.margin_bottom = Pt(pad[3])
    tf._txBody.bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        if isinstance(item[0], str):
            item = [item]
        for text, size, bold, color in item:
            run = p.add_run()
            set_run(run, text, size, bold, color)
    return box


def fill_shape_text(shape, paragraphs, align=PP_ALIGN.CENTER, anchor="m", margins=(6, 6, 6, 6)):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(margins[0])
    tf.margin_right = Pt(margins[1])
    tf.margin_top = Pt(margins[2])
    tf.margin_bottom = Pt(margins[3])
    tf._txBody.bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    for i, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        if isinstance(item[0], str):
            item = [item]
        for text, size, bold, color in item:
            run = p.add_run()
            set_run(run, text, size, bold, color)


def hbar_panel(slide, x, y, w, h, title, subtitle, rows, n_total):
    add_rect(slide, x, y, w, h, OFF, radius=0.05)
    add_tb(
        slide,
        x + Inches(0.28),
        y + Inches(0.18),
        w - Inches(0.5),
        Inches(0.32),
        [(title, 15, True, NAVY)],
    )
    add_tb(
        slide,
        x + Inches(0.28),
        y + Inches(0.48),
        w - Inches(0.5),
        Inches(0.28),
        [(subtitle, 11, False, MUTED)],
    )

    max_n = max(n for _, n, _ in rows)
    label_w = Inches(2.55)
    bar_x = x + Inches(0.28) + label_w
    bar_max = w - Inches(0.28) - label_w - Inches(0.70)
    row_h = Inches(0.62)
    start_y = y + Inches(0.92)

    for i, (lab, n, col) in enumerate(rows):
        ry = start_y + i * row_h
        add_tb(
            slide,
            x + Inches(0.22),
            ry,
            label_w,
            Inches(0.42),
            [(lab, 12, False, BODY)],
            anchor="m",
        )
        track = add_rect(slide, bar_x, ry + Inches(0.08), bar_max, Inches(0.28), TRACK, radius=0.5)
        bw = max(Inches(0.12), int(bar_max * n / max_n))
        add_rect(slide, bar_x, ry + Inches(0.08), bw, Inches(0.28), col, radius=0.5)
        pct = f"{n}   {int(round(100 * n / n_total))}%"
        add_tb(
            slide,
            bar_x + bw + Inches(0.08),
            ry,
            Inches(1.1),
            Inches(0.42),
            [(pct, 12, True, NAVY)],
            anchor="m",
        )


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    s = prs.slides.add_slide(prs.slide_layouts[6])

    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    add_rect(s, 0, 0, Inches(0.11), SLIDE_H, TEAL)

    add_tb(s, Inches(0.42), Inches(0.20), Inches(12.5), Inches(0.34),
           [("ATSCL IPTV — backlog de tickets", 26, True, NAVY)])
    add_rect(s, Inches(0.42), Inches(0.58), Inches(12.5), Pt(1.15), LINE)
    add_tb(
        s,
        Inches(0.42),
        Inches(0.64),
        Inches(12.5),
        Inches(0.32),
        [(
            "20 tickets abiertos en ATSCL. La mitad toca operación. Dos temas concentran 12 de 20: RTP Valeria y canales bloqueados.",
            13, False, SLATE,
        )],
    )

    kpis = [
        ("20", "Tickets", "backlog ATSCL", TEAL),
        ("10", "Operación SI", "impacto en campo", GOLD),
        ("8", "En Analysis", "aún sin fix version", NAVY),
        ("8", "Fix 9.13.0", "ya etiquetadas a release", TEAL2),
    ]
    tw, th, gap = Inches(3.02), Inches(1.18), Inches(0.14)
    for i, (n, a, b, bg) in enumerate(kpis):
        x = Inches(0.42) + i * (tw + gap)
        sh = add_rect(s, x, Inches(1.04), tw, th, bg, radius=0.08)
        fill_shape_text(sh, [(n, 28, True, WHITE), (a, 13, True, WHITE), (b, 11, False, WHITE)])

    # operación split
    add_tb(s, Inches(0.42), Inches(2.32), Inches(12.5), Inches(0.24),
           [("Operación  ·  10 SI  /  10 NO", 11, True, NAVY)])
    add_rect(s, Inches(0.42), Inches(2.58), Inches(6.25), Inches(0.22), GOLD, radius=0.4)
    add_rect(s, Inches(6.67), Inches(2.58), Inches(6.25), Inches(0.22), SLATE, radius=0.4)
    add_tb(s, Inches(0.42), Inches(2.80), Inches(6.25), Inches(0.22),
           [("SI — RTP Valeria, CENAM/AUP, desbloqueo en campo", 10, False, MUTED)])
    add_tb(s, Inches(6.67), Inches(2.80), Inches(6.25), Inches(0.22),
           [("NO — Validation interna, HBO 11.0.0, audio/subs", 10, False, MUTED)],
           align=PP_ALIGN.RIGHT)

    hbar_panel(
        s, Inches(0.42), Inches(3.12), Inches(6.20), Inches(3.88),
        "Pipeline por estado",
        "Orden de flujo · n = 20",
        ESTADO, 20,
    )
    hbar_panel(
        s, Inches(6.78), Inches(3.12), Inches(6.14), Inches(3.88),
        "Concentración por tema",
        "12 de 20 caen en dos defectos de producto",
        TEMAS, 20,
    )

    add_rect(s, Inches(0.42), Inches(7.12), Inches(12.5), Pt(1.1), TEAL)
    add_tb(s, Inches(0.42), Inches(7.16), Inches(8.4), Inches(0.28),
           [("Quality Control · Claro Video", 10, False, MUTED)])
    add_tb(
        s, Inches(7.0), Inches(7.16), Inches(5.9), Inches(0.28),
        [("Fuente: ATSCL backlog  ·  21 sep 2026", 10, False, MUTED)],
        align=PP_ALIGN.RIGHT,
    )

    prs.save(OUT)
    print("saved", OUT)
    render_preview()


def font(size, bold=False):
    path = (
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
        if bold
        else "/System/Library/Fonts/Supplemental/Arial.ttf"
    )
    return ImageFont.truetype(path, size)


def render_preview():
    W, H = 1920, 1080
    sc = W / 13.333

    def inch(x):
        return int(x * sc)

    navy, teal, teal2 = (0x1B, 0x2A, 0x49), (0x0D, 0x8B, 0x7D), (0x13, 0xA6, 0x96)
    gold, slate = (0xE8, 0x9F, 0x1F), (0x51, 0x5F, 0x75)
    line, off, muted, body = (0xD0, 0xD6, 0xE1), (0xF4, 0xF8, 0xFB), (0x6B, 0x7A, 0x8D), (0x3A, 0x47, 0x5C)
    track = (0xE8, 0xEE, 0xF2)
    white = (255, 255, 255)

    img = Image.new("RGB", (W, H), white)
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, inch(0.11), H], fill=teal)
    d.text((inch(0.42), inch(0.20)), "ATSCL IPTV — backlog de tickets", font=font(40, True), fill=navy)
    d.rectangle([inch(0.42), inch(0.58), inch(12.92), inch(0.595)], fill=line)
    d.text(
        (inch(0.42), inch(0.66)),
        "20 tickets abiertos en ATSCL. La mitad toca operación. Dos temas concentran 12 de 20: RTP Valeria y canales bloqueados.",
        font=font(20), fill=slate,
    )

    kpis = [
        ("20", "Tickets", "backlog ATSCL", teal),
        ("10", "Operación SI", "impacto en campo", gold),
        ("8", "En Analysis", "aún sin fix version", navy),
        ("8", "Fix 9.13.0", "ya etiquetadas a release", teal2),
    ]
    tw, th, gap = inch(3.02), inch(1.18), inch(0.14)
    for i, (n, a, b, bg) in enumerate(kpis):
        x = inch(0.42) + i * (tw + gap)
        y = inch(1.04)
        d.rounded_rectangle([x, y, x + tw, y + th], radius=14, fill=bg)
        d.text((x + tw // 2, y + inch(0.22)), n, font=font(42, True), fill=white, anchor="ma")
        d.text((x + tw // 2, y + inch(0.64)), a, font=font(18, True), fill=white, anchor="ma")
        d.text((x + tw // 2, y + inch(0.90)), b, font=font(15), fill=white, anchor="ma")

    d.text((inch(0.42), inch(2.32)), "Operación  ·  10 SI  /  10 NO", font=font(16, True), fill=navy)
    d.rounded_rectangle([inch(0.42), inch(2.58), inch(6.67), inch(2.80)], radius=8, fill=gold)
    d.rounded_rectangle([inch(6.67), inch(2.58), inch(12.92), inch(2.80)], radius=8, fill=slate)
    d.text((inch(0.42), inch(2.84)), "SI — RTP Valeria, CENAM/AUP, desbloqueo en campo", font=font(14), fill=muted)
    d.text((inch(12.92), inch(2.84)), "NO — Validation interna, HBO 11.0.0, audio/subs", font=font(14), fill=muted, anchor="ra")

    def panel(px, py, pw, ph, title, sub, rows, total, colors):
        d.rounded_rectangle([px, py, px + pw, py + ph], radius=12, fill=off)
        d.text((px + inch(0.28), py + inch(0.18)), title, font=font(22, True), fill=navy)
        d.text((px + inch(0.28), py + inch(0.48)), sub, font=font(15), fill=muted)
        max_n = max(n for _, n, _ in rows)
        label_w = inch(2.55)
        bar_x = px + inch(0.28) + label_w
        bar_max = pw - inch(0.28) - label_w - inch(0.70)
        start = py + inch(0.92)
        rh = inch(0.62)
        for i, (lab, n, col) in enumerate(rows):
            ry = start + i * rh
            d.text((px + inch(0.28), ry + inch(0.08)), lab, font=font(16), fill=body)
            d.rounded_rectangle([bar_x, ry + inch(0.10), bar_x + bar_max, ry + inch(0.36)], radius=8, fill=track)
            bw = max(inch(0.12), int(bar_max * n / max_n))
            d.rounded_rectangle([bar_x, ry + inch(0.10), bar_x + bw, ry + inch(0.36)], radius=8, fill=col)
            d.text((bar_x + bw + inch(0.10), ry + inch(0.08)), f"{n}   {int(round(100*n/total))}%", font=font(16, True), fill=navy)

    estado = [("Analysis", 8, gold), ("To Develop", 4, slate), ("Validation", 6, teal), ("QA Validation", 2, teal2)]
    temas = [
        ("RTP Valeria", 6, teal),
        ("Canales bloqueados", 6, teal2),
        ("Audio y subtítulos", 2, navy),
        ("HBO Max", 2, gold),
        ("Otros", 4, slate),
    ]
    panel(inch(0.42), inch(3.12), inch(6.20), inch(3.88), "Pipeline por estado", "Orden de flujo · n = 20", estado, 20, None)
    panel(inch(6.78), inch(3.12), inch(6.14), inch(3.88), "Concentración por tema", "12 de 20 caen en dos defectos de producto", temas, 20, None)

    d.rectangle([inch(0.42), inch(7.12), inch(12.92), inch(7.135)], fill=teal)
    d.text((inch(0.42), inch(7.18)), "Quality Control · Claro Video", font=font(15), fill=muted)
    d.text((inch(12.92), inch(7.18)), "Fuente: ATSCL backlog  ·  21 sep 2026", font=font(15), fill=muted, anchor="ra")
    img.save(PREV)
    print("preview", PREV)


if __name__ == "__main__":
    build()
