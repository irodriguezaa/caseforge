#!/usr/bin/env python3
"""Épicas en Validation — QC. Paleta alineada a QC_Agosto2026_v1.3."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pathlib import Path

# Widescreen 16:9 matching the Agosto deck (~13.33" x 7.5")
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x1B, 0x2A, 0x49)
TEAL = RGBColor(0x0D, 0x8B, 0x7D)
TEAL2 = RGBColor(0x13, 0xA6, 0x96)
GOLD = RGBColor(0xE8, 0x9F, 0x1F)
SLATE = RGBColor(0x51, 0x5F, 0x75)
LINE = RGBColor(0xD0, 0xD6, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF4, 0xF8, 0xFB)
MUTED = RGBColor(0x6B, 0x7A, 0x8D)
CYAN = RGBColor(0xA8, 0xD4, 0xEC)
BODY = RGBColor(0x3A, 0x47, 0x5C)

OUT = Path("/Users/rodriguezisr/Documents/QC/Epicas_en_Validation_QC_Ejecutivo.pptx")
PREV = Path("/Users/rodriguezisr/Documents/qcpulse/.tmp_pptx_build/preview_slide1.png")


def rgb_tuple(c: RGBColor):
    return (int(str(c)[0:2], 16), int(str(c)[2:4], 16), int(str(c)[4:6], 16))


def set_run(run, text, size, bold=False, color=NAVY, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _no_autofit(tf):
    tf.word_wrap = True
    try:
        tf.auto_size = None
    except Exception:
        pass


def add_rect(slide, l, t, w, h, fill, line=None, radius=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if radius is not None:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    s.shadow.inherit = False
    return s


def add_tb(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor="t", pad=(6, 6, 6, 6)):
    """lines: list of (text, size, bold, color) or list of list of runs."""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    _no_autofit(tf)
    tf.margin_left = Pt(pad[0])
    tf.margin_right = Pt(pad[1])
    tf.margin_top = Pt(pad[2])
    tf.margin_bottom = Pt(pad[3])
    anchor_map = {"t": "t", "m": "ctr", "b": "b"}
    tf._txBody.bodyPr.set("anchor", anchor_map.get(anchor, "t"))
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        if isinstance(item[0], tuple):
            for j, runspec in enumerate(item):
                run = p.add_run() if j else p.runs[0] if p.runs else p.add_run()
                if j == 0 and not p.runs:
                    run = p.add_run()
                # first run: use existing empty run
            p.clear()
        # normalize
        runs = item if item and isinstance(item[0], (list, tuple)) and len(item[0]) >= 3 and not isinstance(item[0], str) else [item]
        # detect: (text, size, bold, color)
        if isinstance(item[0], str):
            runs = [item]
        else:
            runs = item
        first = True
        for rs in runs:
            text, size, bold, color = rs[0], rs[1], rs[2], rs[3]
            run = p.add_run() if not first or (first and p.runs) else p.add_run()
            if first:
                # python-pptx creates one empty run sometimes; just set on add_run
                first = False
            set_run(run, text, size, bold, color)
        # if we added extra empty first run, clean later
    # cleanup empty runs
    for p in tf.paragraphs:
        for r in list(p.runs):
            if r.text == "" and len(p.runs) > 1:
                r._r.getparent().remove(r._r)
    return box


def clear_paragraphs_set(tf, paragraphs, align, anchor="t"):
    _no_autofit(tf)
    tf._txBody.bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    for i, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(2)
        # item is list of runs or a single run tuple
        if isinstance(item[0], str):
            item = [item]
        # wipe default run
        for r in list(p.runs):
            r.text = ""
        for text, size, bold, color in item:
            run = p.add_run()
            set_run(run, text, size, bold, color)


def fill_shape_text(shape, paragraphs, align=PP_ALIGN.LEFT, anchor="m", margins=(10, 10, 8, 8)):
    tf = shape.text_frame
    _no_autofit(tf)
    tf.margin_left = Pt(margins[0])
    tf.margin_right = Pt(margins[1])
    tf.margin_top = Pt(margins[2])
    tf.margin_bottom = Pt(margins[3])
    tf._txBody.bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    for i, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(2)
        if isinstance(item[0], str):
            item = [item]
        for r in list(p.runs):
            r.text = ""
        for text, size, bold, color in item:
            run = p.add_run()
            set_run(run, text, size, bold, color)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)

    # canvas
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    # left rail
    add_rect(s, 0, 0, Inches(0.11), SLIDE_H, TEAL)

    # header
    add_tb(
        s,
        Inches(0.42),
        Inches(0.22),
        Inches(12.5),
        Inches(0.32),
        [("Épicas en Validation — QC", 26, True, NAVY)],
        pad=(0, 0, 0, 0),
    )
    # underline
    add_rect(s, Inches(0.42), Inches(0.62), Inches(12.5), Pt(1.15), LINE)

    add_tb(
        s,
        Inches(0.42),
        Inches(0.68),
        Inches(12.5),
        Inches(0.36),
        [
            (
                "Se revisaron ~200 EPCs en Validation. Hoy quedan 64 activas, agrupadas por sprint y por causa de permanencia.",
                13,
                False,
                SLATE,
            )
        ],
        pad=(0, 0, 0, 0),
    )

    # KPI tiles
    tiles = [
        ("~200", "Revisadas", "EPCs en Validation", TEAL, WHITE),
        ("64", "Activas hoy", "siguen en Validation", TEAL2, WHITE),
        ("23", "Sprint 44", "cierra mañana", GOLD, WHITE),
        ("7", "Sprint 45 · Nubiral", "ya validadas QA/QC; revisar y ajustar", NAVY, WHITE),
    ]
    x0 = Inches(0.42)
    y0 = Inches(1.12)
    tw = Inches(3.02)
    th = Inches(1.42)
    gap = Inches(0.14)
    for i, (n, a, b, bg, fg) in enumerate(tiles):
        x = x0 + i * (tw + gap)
        sh = add_rect(s, x, y0, tw, th, bg, radius=0.08)
        fill_shape_text(
            sh,
            [
                (n, 32, True, fg),
                (a, 13, True, fg),
                (b, 11, False, fg),
            ],
            align=PP_ALIGN.CENTER,
            anchor="m",
            margins=(8, 8, 8, 8),
        )

    # composition bar of 64
    add_tb(
        s,
        Inches(0.42),
        Inches(2.62),
        Inches(12.5),
        Inches(0.28),
        [("Composición de las 64 activas", 12, True, NAVY)],
        pad=(0, 0, 0, 0),
    )
    bar_y = Inches(2.92)
    bar_h = Inches(0.28)
    bar_x = Inches(0.42)
    bar_w = Inches(12.5)
    segs = [
        (23, GOLD, "23  Sprint 44"),
        (7, NAVY, "7  Nubiral 45"),
        (13, TEAL, "13  ADTCL 10.1.2"),
        (19, TEAL2, "19  En curso"),
        (2, SLATE, "2  AAF"),
    ]
    cursor = bar_x
    total = 64
    for n, col, _lab in segs:
        w = int(bar_w * n / total)
        add_rect(s, cursor, bar_y, w, bar_h, col)
        cursor += w

    # legend
    lx = Inches(0.42)
    ly = Inches(3.28)
    for n, col, lab in segs:
        add_rect(s, lx, ly + Inches(0.05), Inches(0.16), Inches(0.16), col, radius=0.2)
        add_tb(
            s,
            lx + Inches(0.20),
            ly,
            Inches(2.2),
            Inches(0.26),
            [(lab, 11, False, BODY)],
            pad=(0, 0, 0, 0),
        )
        lx += Inches(2.5)

    # section 34
    add_tb(
        s,
        Inches(0.42),
        Inches(3.58),
        Inches(12.5),
        Inches(0.32),
        [
            [
                ("Las otras 34", 14, True, NAVY),
                ("   ·  sprint < 44 o sin sprint asignado", 13, False, SLATE),
            ]
        ],
        pad=(0, 0, 0, 0),
    )

    cards = [
        (
            TEAL,
            "13",
            "ADTCL v10.1.2",
            "Rediseño del player",
            "QA ya validó; QC no. No se retomó: QA lo regresó y la capacidad se fue a HBO Max y PayPal.",
        ),
        (
            TEAL2,
            "19",
            "Proyectos en curso",
            "Splash · PayPal · HBO Max",
            "Dispositivos: SCTCL, ATSCL, iOSPR, ADRPR, TVOSPR, ROKUPR y C9085PR.",
        ),
        (
            SLATE,
            "2",
            "Unificación AAF",
            "Ambiente unificado",
            "Proyecto de unificación de ambiente AAF. En proceso.",
        ),
    ]
    cw = Inches(4.06)
    ch = Inches(2.55)
    cy = Inches(3.95)
    cg = Inches(0.16)
    for i, (accent, num, title, sub, body) in enumerate(cards):
        x = Inches(0.42) + i * (cw + cg)
        add_rect(s, x, cy, cw, ch, OFF, radius=0.06)
        add_rect(s, x, cy, Inches(0.09), ch, accent)
        add_tb(
            s,
            x + Inches(0.22),
            cy + Inches(0.12),
            cw - Inches(0.32),
            Inches(2.3),
            [
                (num, 28, True, accent),
                (title, 15, True, NAVY),
                (sub, 12, True, TEAL if accent != SLATE else MUTED),
                (body, 12, False, BODY),
            ],
            pad=(4, 6, 4, 4),
        )

    # footer
    add_rect(s, Inches(0.42), Inches(7.12), Inches(12.5), Pt(1.1), TEAL)
    add_tb(
        s,
        Inches(0.42),
        Inches(7.16),
        Inches(8.2),
        Inches(0.28),
        [("Quality Control · Claro Video", 10, False, MUTED)],
        pad=(0, 0, 0, 0),
    )
    add_tb(
        s,
        Inches(7.2),
        Inches(7.16),
        Inches(5.7),
        Inches(0.28),
        [("Fuente: Jira · estado Validation  ·  18 sep 2026", 10, False, MUTED)],
        align=PP_ALIGN.RIGHT,
        pad=(0, 0, 0, 0),
    )

    # ---- Slide 2: lectura ejecutiva / próximos pasos ----
    s2 = prs.slides.add_slide(blank)
    add_rect(s2, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    add_rect(s2, 0, 0, Inches(0.11), SLIDE_H, TEAL)
    add_tb(
        s2,
        Inches(0.42),
        Inches(0.22),
        Inches(12.5),
        Inches(0.32),
        [("Épicas en Validation — QC  ·  lectura y siguientes pasos", 24, True, NAVY)],
        pad=(0, 0, 0, 0),
    )
    add_rect(s2, Inches(0.42), Inches(0.62), Inches(12.5), Pt(1.15), LINE)
    add_tb(
        s2,
        Inches(0.42),
        Inches(0.70),
        Inches(12.5),
        Inches(0.36),
        [
            (
                "De las 64 activas, 30 ya están en un sprint vigente (44 o 45). Las 34 restantes no salen de Validation por capacidad o por proyecto en curso.",
                13,
                False,
                SLATE,
            )
        ],
        pad=(0, 0, 0, 0),
    )

    rows = [
        (GOLD, "23", "Sprint 44 — cierran mañana", "Foco inmediato de QC. Cerrar o devolver lo que no quede OK en el sprint que termina."),
        (NAVY, "7", "Sprint 45 — Nubiral", "QA/QC ya las validó en su momento. Revisar con Nubiral y ajustar lo necesario (no son backlog nuevo)."),
        (TEAL, "13", "ADTCL v10.1.2 — player", "QA ok, QC pendiente. Se pausó por el retorno de QA y por HBO Max / PayPal. No está en el flujo actual."),
        (TEAL2, "19", "Splash / PayPal / HBO Max", "Van con los proyectos en curso en SCTCL, ATSCL, iOSPR, ADRPR, TVOSPR, ROKUPR y C9085PR."),
        (SLATE, "2", "Unificación AAF", "Dos épicas de unificación de ambiente. Siguen el proyecto AAF, no un release de app."),
    ]
    y = Inches(1.18)
    for accent, num, title, body in rows:
        add_rect(s2, Inches(0.42), y, Inches(12.5), Inches(0.98), OFF, radius=0.05)
        add_rect(s2, Inches(0.42), y, Inches(0.10), Inches(0.98), accent)
        numbox = add_rect(s2, Inches(0.70), y + Inches(0.18), Inches(1.05), Inches(0.62), accent, radius=0.12)
        fill_shape_text(numbox, [(num, 22, True, WHITE)], align=PP_ALIGN.CENTER, anchor="m", margins=(2, 2, 2, 2))
        add_tb(
            s2,
            Inches(1.95),
            y + Inches(0.12),
            Inches(10.6),
            Inches(0.78),
            [
                (title, 15, True, NAVY),
                (body, 13, False, BODY),
            ],
            pad=(0, 0, 0, 0),
        )
        y += Inches(1.08)

    add_rect(s2, Inches(0.42), Inches(7.12), Inches(12.5), Pt(1.1), TEAL)
    add_tb(
        s2,
        Inches(0.42),
        Inches(7.16),
        Inches(12.5),
        Inches(0.28),
        [("Quality Control · Claro Video    ·    Las 23 de Sprint 44 son la única fecha dura de esta semana.", 10, False, MUTED)],
        pad=(0, 0, 0, 0),
    )

    prs.save(OUT)
    print("saved", OUT)
    render_preview()


def render_preview():
    """PNG preview of slide 1 for visual QA (same geometry, 2x)."""
    W, H = 1920, 1080
    scale = W / 13.333
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)

    def inch(x):
        return int(x * scale)

    def col(rgb):
        return (rgb[0], rgb[1], rgb[2])

    navy = (0x1B, 0x2A, 0x49)
    teal = (0x0D, 0x8B, 0x7D)
    teal2 = (0x13, 0xA6, 0x96)
    gold = (0xE8, 0x9F, 0x1F)
    slate = (0x51, 0x5F, 0x75)
    line = (0xD0, 0xD6, 0xE1)
    off = (0xF4, 0xF8, 0xFB)
    muted = (0x6B, 0x7A, 0x8D)
    body = (0x3A, 0x47, 0x5C)

    def font(size, bold=False):
        name = "Calibri Bold" if bold else "Calibri"
        candidates = [
            f"/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/System/Library/Fonts/Supplemental/Calibri.ttf",
            "/Library/Fonts/Microsoft/Calibri.ttf",
        ]
        for c in candidates:
            p = Path(c)
            if p.exists():
                try:
                    return ImageFont.truetype(str(p), size)
                except Exception:
                    continue
        return ImageFont.load_default()

    d.rectangle([0, 0, inch(0.11), H], fill=teal)
    d.text((inch(0.42), inch(0.22)), "Épicas en Validation — QC", font=font(40, True), fill=navy)
    d.rectangle([inch(0.42), inch(0.62), inch(12.92), inch(0.63)], fill=line)
    d.text(
        (inch(0.42), inch(0.70)),
        "Se revisaron ~200 EPCs en Validation. Hoy quedan 64 activas, agrupadas por sprint y por causa de permanencia.",
        font=font(20),
        fill=slate,
    )

    tiles = [
        ("~200", "Revisadas", "EPCs en Validation", teal),
        ("64", "Activas hoy", "siguen en Validation", teal2),
        ("23", "Sprint 44", "cierra mañana", gold),
        ("7", "Sprint 45 · Nubiral", "ya validadas QA/QC; revisar y ajustar", navy),
    ]
    tw, th, gap = inch(3.02), inch(1.42), inch(0.14)
    x0, y0 = inch(0.42), inch(1.12)
    for i, (n, a, b, bg) in enumerate(tiles):
        x = x0 + i * (tw + gap)
        d.rounded_rectangle([x, y0, x + tw, y0 + th], radius=12, fill=bg)
        d.text((x + tw // 2, y0 + inch(0.22)), n, font=font(48, True), fill=(255, 255, 255), anchor="ma")
        d.text((x + tw // 2, y0 + inch(0.72)), a, font=font(20, True), fill=(255, 255, 255), anchor="ma")
        d.text((x + tw // 2, y0 + inch(1.02)), b, font=font(16), fill=(255, 255, 255), anchor="ma")

    d.text((inch(0.42), inch(2.62)), "Composición de las 64 activas", font=font(20, True), fill=navy)
    segs = [(23, gold), (7, navy), (13, teal), (19, teal2), (2, slate)]
    bx, by, bw, bh = inch(0.42), inch(2.92), inch(12.5), inch(0.28)
    cur = bx
    for n, c in segs:
        w = int(bw * n / 64)
        d.rectangle([cur, by, cur + w, by + bh], fill=c)
        cur += w

    labs = ["23  Sprint 44", "7  Nubiral 45", "13  ADTCL 10.1.2", "19  En curso", "2  AAF"]
    lx = inch(0.42)
    ly = inch(3.30)
    for (n, c), lab in zip(segs, labs):
        d.rounded_rectangle([lx, ly, lx + inch(0.16), ly + inch(0.16)], radius=3, fill=c)
        d.text((lx + inch(0.22), ly - 2), lab, font=font(16), fill=body)
        lx += inch(2.5)

    d.text((inch(0.42), inch(3.60)), "Las otras 34   ·   sprint < 44 o sin sprint asignado", font=font(22, True), fill=navy)

    cards = [
        (teal, "13", "ADTCL v10.1.2", "Rediseño del player", "QA ya validó; QC no. No se retomó: QA lo regresó y la capacidad se fue a HBO Max y PayPal."),
        (teal2, "19", "Proyectos en curso", "Splash · PayPal · HBO Max", "Dispositivos: SCTCL, ATSCL, iOSPR, ADRPR, TVOSPR, ROKUPR y C9085PR."),
        (slate, "2", "Unificación AAF", "Ambiente unificado", "Proyecto de unificación de ambiente AAF. En proceso."),
    ]
    cw, ch, cg = inch(4.06), inch(2.55), inch(0.16)
    cy = inch(3.95)
    for i, (accent, num, title, sub, bodyt) in enumerate(cards):
        x = inch(0.42) + i * (cw + cg)
        d.rounded_rectangle([x, cy, x + cw, cy + ch], radius=10, fill=off)
        d.rectangle([x, cy, x + inch(0.09), cy + ch], fill=accent)
        d.text((x + inch(0.28), cy + inch(0.18)), num, font=font(42, True), fill=accent)
        d.text((x + inch(0.28), cy + inch(0.78)), title, font=font(22, True), fill=navy)
        d.text((x + inch(0.28), cy + inch(1.12)), sub, font=font(18, True), fill=teal if accent != slate else muted)
        # wrap body
        f = font(17)
        words = bodyt.split()
        lines, cur_l = [], ""
        maxw = cw - inch(0.50)
        for w in words:
            trial = (cur_l + " " + w).strip()
            if d.textlength(trial, font=f) <= maxw:
                cur_l = trial
            else:
                lines.append(cur_l)
                cur_l = w
        if cur_l:
            lines.append(cur_l)
        yy = cy + inch(1.50)
        for ln in lines:
            d.text((x + inch(0.28), yy), ln, font=f, fill=body)
            yy += inch(0.28)

    d.rectangle([inch(0.42), inch(7.12), inch(12.92), inch(7.135)], fill=teal)
    d.text((inch(0.42), inch(7.18)), "Quality Control · Claro Video", font=font(15), fill=muted)
    d.text((inch(8.2), inch(7.18)), "Fuente: Jira · estado Validation  ·  18 sep 2026", font=font(15), fill=muted)

    img.save(PREV)
    print("preview", PREV)


if __name__ == "__main__":
    build()
