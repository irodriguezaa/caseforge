#!/usr/bin/env python3
"""Reporte ejecutivo QC — STV TATA vATV-3.8.0. Restaura v1 y genera v2."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x1B, 0x2A, 0x49)
TEAL = RGBColor(0x0D, 0x8B, 0x7D)
TEAL2 = RGBColor(0x13, 0xA6, 0x96)
GOLD = RGBColor(0xE8, 0x9F, 0x1F)
SLATE = RGBColor(0x51, 0x5F, 0x75)
LINE = RGBColor(0xD0, 0xD6, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF4, 0xF8, 0xFB)
MUTED = RGBColor(0x6B, 0x7A, 0x8D)
BODY = RGBColor(0x3A, 0x47, 0x5C)
RED = RGBColor(0xC4, 0x44, 0x2E)

NO_REPLICA = [
    "SCTCL-3191", "SCTCL-3190", "SCTCL-3189", "SCTCL-3188", "SCTCL-3185",
    "SCTCL-3125", "SCTCL-3124", "SCTCL-3123", "SCTCL-3122", "SCTCL-3112",
]


def set_run(run, text, size, bold=False, color=NAVY, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    if radius is not None:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    s.shadow.inherit = False
    return s


def add_tb(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor="t", pad=(0, 0, 0, 0), after=2):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(pad[0])
    tf.margin_right = Pt(pad[1])
    tf.margin_top = Pt(pad[2])
    tf.margin_bottom = Pt(pad[3])
    tf._txBody.bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(after)
        if isinstance(item[0], str):
            item = [item]
        for text, size, bold, color in item:
            run = p.add_run()
            set_run(run, text, size, bold, color)
    return box


def fill_shape_text(shape, paragraphs, align=PP_ALIGN.CENTER, anchor="m", margins=(6, 6, 6, 6)):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(margins[0])
    tf.margin_right = Pt(margins[1])
    tf.margin_top = Pt(margins[2])
    tf.margin_bottom = Pt(margins[3])
    tf._txBody.bodyPr.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    for i, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        if isinstance(item[0], str):
            item = [item]
        for text, size, bold, color in item:
            run = p.add_run()
            set_run(run, text, size, bold, color)


def slide_chrome(s, subtitle):
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    add_rect(s, 0, 0, Inches(0.11), SLIDE_H, TEAL)
    add_tb(s, Inches(0.42), Inches(0.16), Inches(12.5), Inches(0.32),
           [("Reporte Ejecutivo QC  ·  STV TATA  vATV-3.8.0", 22, True, NAVY)])
    add_tb(s, Inches(0.42), Inches(0.48), Inches(12.5), Inches(0.24),
           [(subtitle, 13, False, SLATE)])
    add_rect(s, Inches(0.42), Inches(0.76), Inches(12.5), Pt(1.15), LINE)


def kpis_4(s):
    kpis = [
        ("13", "Revisados", "todos en vATV-3.8.0", TEAL),
        ("10", "No se replican", "evidencia en el ticket", TEAL2),
        ("2", "Se replican", "3186 y 3161", GOLD),
        ("1", "Sin BRF / EPC", "cambio de logo", SLATE),
    ]
    tw, th, gap = Inches(3.02), Inches(1.10), Inches(0.14)
    for i, (n, a, b, bg) in enumerate(kpis):
        x = Inches(0.42) + i * (tw + gap)
        sh = add_rect(s, x, Inches(0.92), tw, th, bg, radius=0.08)
        fill_shape_text(sh, [(n, 26, True, WHITE), (a, 13, True, WHITE), (b, 11, False, WHITE)])


def left_noreplica(s, n_label):
    add_rect(s, Inches(0.42), Inches(2.20), Inches(6.20), Inches(4.72), OFF, radius=0.05)
    add_tb(s, Inches(0.62), Inches(2.32), Inches(5.80), Inches(0.30),
           [(f"No se replican  ·  {n_label} tickets", 14, True, NAVY)])
    add_tb(
        s, Inches(0.62), Inches(2.68), Inches(5.80), Inches(0.36),
        [("QC no reproduce el defecto. Se subió evidencia al ticket.", 12, False, MUTED)],
    )
    col1, col2 = NO_REPLICA[:5], NO_REPLICA[5:]
    add_tb(s, Inches(0.70), Inches(3.12), Inches(2.70), Inches(3.40),
           [(t, 14, False, BODY) for t in col1], after=8)
    add_tb(s, Inches(3.55), Inches(3.12), Inches(2.70), Inches(3.40),
           [(t, 14, False, BODY) for t in col2], after=8)


def footer(s):
    add_rect(s, Inches(0.42), Inches(7.12), Inches(12.5), Pt(1.1), TEAL)
    add_tb(s, Inches(0.42), Inches(7.16), Inches(12.5), Inches(0.26),
           [("Quality Control · Claro Video    ·    STV TATA · SCTCL · vATV-3.8.0    ·    24 sep 2026", 10, False, MUTED)])


def build_v1():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_chrome(s, "13 tickets SCTCL revisados  ·  8 Blocker  ·  5 Crítica  ·  7 Editorial  ·  6 App")
    kpis_4(s)
    left_noreplica(s, "10")
    cards = [
        (GOLD, "Se replica  ·  SCTCL-3186",
         "Full EPG con cuadros en blanco. El comportamiento se ve, pero obedece a eventos de 1–3 min. También en otros dispositivos."),
        (TEAL, "Se replica solo en AR  ·  SCTCL-3161",
         "Time Shift: la barra no muestra carga real. Se replicó en AR con usuario con TS; en otras regiones no."),
        (SLATE, "Fuera de alcance de fix  ·  SCTCL-3193",
         "Logos desactualizados. No hay BRF/EPC de cambio de logo. Blocker editorial."),
    ]
    y = Inches(2.20)
    for col, title, body in cards:
        add_rect(s, Inches(6.80), y, Inches(6.10), Inches(1.48), OFF, radius=0.05)
        add_rect(s, Inches(6.80), y, Inches(0.09), Inches(1.48), col)
        add_tb(s, Inches(7.08), y + Inches(0.12), Inches(5.60), Inches(1.24),
               [(title, 13, True, NAVY), (body, 12, False, BODY)])
        y += Inches(1.62)
    footer(s)
    out = Path("/Users/rodriguezisr/Documents/QC/QC_STV_TATA_3.8.0_Ejecutivo.pptx")
    dl = Path("/Users/rodriguezisr/Downloads/QC_STV_TATA_3.8.0_Ejecutivo.pptx")
    prs.save(out)
    prs.save(dl)
    print("restored", out)


def build_v2():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slide_chrome(s, "13 tickets SCTCL revisados  ·  8 Blocker  ·  5 Crítica  ·  7 editoriales (independiente de réplica)  ·  6 App")
    kpis_4(s)
    left_noreplica(s, "10")
    cards = [
        (RED, "SCTCL-3185  ·  TV en vivo no disponible",
         "QC no ha logrado replicar; se sigue probando para capturar la traza. Negocio sí lo reprodujo."),
        (GOLD, "Se replica  ·  SCTCL-3186",
         "Full EPG con cuadros en blanco. El comportamiento se ve, pero obedece a eventos de 1–3 min. También en otros dispositivos."),
        (TEAL, "Se replica solo en AR  ·  SCTCL-3161",
         "Time Shift: la barra no muestra carga real. Se replicó en AR con usuario con TS; en otras regiones no."),
        (SLATE, "Fuera de alcance de fix  ·  SCTCL-3193",
         "Logos desactualizados. No hay BRF/EPC de cambio de logo. Blocker editorial."),
    ]
    y = Inches(2.20)
    for col, title, body in cards:
        add_rect(s, Inches(6.80), y, Inches(6.10), Inches(1.12), OFF, radius=0.05)
        add_rect(s, Inches(6.80), y, Inches(0.09), Inches(1.12), col)
        add_tb(s, Inches(7.06), y + Inches(0.06), Inches(5.65), Inches(1.02),
               [(title, 12, True, NAVY), (body, 11, False, BODY)])
        y += Inches(1.20)
    footer(s)
    out = Path("/Users/rodriguezisr/Documents/QC/QC_STV_TATA_3.8.0_Ejecutivo_v2.pptx")
    dl = Path("/Users/rodriguezisr/Downloads/QC_STV_TATA_3.8.0_Ejecutivo_v2.pptx")
    prs.save(out)
    prs.save(dl)
    print("v2", out)


if __name__ == "__main__":
    build_v1()
    build_v2()
