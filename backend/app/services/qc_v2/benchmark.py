"""Benchmark de QcPulse v2 contra Luminia (criterio de liberación).

Uso:
    docker compose run --rm --no-deps -v "$PWD/exports:/out" backend \
        python -m app.services.qc_v2.benchmark \
        --luminia /out/Luminia_WEB.xlsx --qcpulse /out/qcv2_WEB.xlsx --out /out --judge

Métricas por épica:
- Granularidad: casos QcPulse ÷ casos Luminia (meta 40%–60%).
- Pasos: % de casos QcPulse con 3 a 6 pasos (meta ≥ 90%).
- Con --judge (usa OpenAI): cobertura de las verificaciones de Luminia (meta ≥ 90% sí+parcial),
  pasos de Desarrollo (meta 0) y duplicados por la regla del bug (meta 0).
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import settings
from app.services.qc_v2.llm import LlmError, chat_json

logger = logging.getLogger(__name__)

META_COBERTURA = 0.90
META_GRANULARIDAD = (0.40, 0.60)
META_PASOS = 0.90


@dataclass
class BenchCase:
    id: str
    epic: str
    title: str
    steps: list[str] = field(default_factory=list)
    expected: list[str] = field(default_factory=list)


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _split_numbered(text: Any) -> list[str]:
    lines = [line.strip() for line in str(text or "").splitlines() if line.strip()]
    return [re.sub(r"^\d+[.)]\s*", "", line) for line in lines]


def read_luminia(path: str | Path) -> list[BenchCase]:
    """Hoja 'Matriz de CP QC': una fila de encabezado de caso y filas de pasos debajo."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = next((n for n in workbook.sheetnames if n.strip().lower() == "matriz de cp qc"), None)
    if sheet_name is None:
        raise ValueError("El archivo de Luminia no tiene la hoja 'Matriz de CP QC'")
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    header = [_norm_header(h) for h in rows[0]]

    def col(*names: str) -> int:
        for name in names:
            if name in header:
                return header.index(name)
        raise ValueError(f"Columna no encontrada en Luminia: {names}")

    c_epic, c_id, c_name = col("technical epic"), col("external id"), col("tc_name")
    c_step, c_exp = col("test step description"), col("expected result")
    cases: list[BenchCase] = []
    current: BenchCase | None = None
    for row in rows[1:]:
        if row[c_id]:
            current = BenchCase(id=str(row[c_id]).strip(), epic=str(row[c_epic] or "").strip(), title=str(row[c_name] or "").strip())
            cases.append(current)
        if current is not None and (row[c_step] or row[c_exp]):
            current.steps.append(str(row[c_step] or "").strip())
            current.expected.append(str(row[c_exp] or "").strip())
    return cases


def read_qcpulse(path: str | Path) -> list[BenchCase]:
    """Acepta el Excel del CLI (hoja 'Casos') o el export de la Release (hoja 'Test Cases')."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Casos" in workbook.sheetnames:
        rows = list(workbook["Casos"].iter_rows(values_only=True))
        header = [_norm_header(h) for h in rows[0]]
        i_epic, i_title = header.index("épica"), header.index("título")
        i_steps, i_exp = header.index("pasos"), header.index("resultado esperado")
        cases = []
        for n, row in enumerate(rows[1:], start=1):
            if not row[i_title]:
                continue
            cases.append(
                BenchCase(
                    id=f"Q{n:03d}",
                    epic=str(row[i_epic] or "").strip(),
                    title=str(row[i_title]).strip(),
                    steps=_split_numbered(row[i_steps]),
                    expected=_split_numbered(row[i_exp]),
                )
            )
        return cases
    if "Test Cases" in workbook.sheetnames:
        rows = list(workbook["Test Cases"].iter_rows(values_only=True))
        header = [_norm_header(h) for h in rows[0]]
        i_id, i_title = header.index("id"), header.index("nombre")
        i_steps, i_exp = header.index("test steps"), header.index("resultado esperado")
        i_epic = header.index("technical epic") if "technical epic" in header else header.index("componente")
        i_type = header.index("issue type") if "issue type" in header else None
        cases = []
        for row in rows[1:]:
            if not row[i_title] or (i_type is not None and str(row[i_type] or "") == "qc_matrix"):
                continue
            cases.append(
                BenchCase(
                    id=str(row[i_id]).strip(),
                    epic=str(row[i_epic] or "").strip(),
                    title=str(row[i_title]).strip(),
                    steps=_split_numbered(row[i_steps]),
                    expected=_split_numbered(row[i_exp]),
                )
            )
        return cases
    raise ValueError("El archivo de QcPulse no tiene hoja 'Casos' ni 'Test Cases'")


# ------------------------------------------------------------------ juez IA

_JUDGE_PROMPT = """Eres un QC Lead que audita casos de prueba. Recibes los casos de LUMINIA (referencia) y los de QCPULSE para la misma épica.
1. Para CADA caso de Luminia decide si su verificación está en QcPulse, ya sea como caso o como paso dentro de un caso:
   "si" (cubierta), "parcial" (cubierta a medias) o "no". Indica los IDs de QcPulse que la cubren.
2. Lista los IDs de QcPulse que tengan algún paso que sea acción de Desarrollo (implementar, configurar llaves, desplegar, consumir un API, revisar logs) y no del tester en el dispositivo.
3. Lista pares de casos de QcPulse que, al fallar, serían el mismo bug (duplicados).
Responde solo JSON: {"cobertura": [{"luminia_id": "CP001", "estado": "si|parcial|no", "qcpulse_ids": ["Q001"], "nota": "..."}], "pasos_desarrollo": [{"qcpulse_id": "Q001", "paso": "..."}], "duplicados": [{"a": "Q001", "b": "Q002", "motivo": "..."}]}"""


def _case_brief(case: BenchCase) -> str:
    steps = " | ".join(f"{i}. {s}" for i, s in enumerate(case.steps, start=1))
    return f"{case.id} — {case.title}\n   Pasos: {steps}"


def judge_epic(
    epic: str,
    luminia: list[BenchCase],
    qcpulse: list[BenchCase],
    chat: Callable[..., Any] = chat_json,
) -> dict[str, Any]:
    user = (
        f"ÉPICA {epic}\n\n=== LUMINIA ({len(luminia)}) ===\n"
        + "\n".join(_case_brief(c) for c in luminia)
        + f"\n\n=== QCPULSE ({len(qcpulse)}) ===\n"
        + "\n".join(_case_brief(c) for c in qcpulse)
    )
    model = getattr(settings, "qc_v2_judge_model", None) or settings.qc_v2_model_b
    payload, _meta = chat(
        [{"role": "system", "content": _JUDGE_PROMPT}, {"role": "user", "content": user}],
        model=model,
        temperature=0,
    )
    return payload


# ------------------------------------------------------------------ métricas


@dataclass
class EpicScore:
    epic: str
    luminia: int
    qcpulse: int
    ratio: float | None
    steps_ok: float | None
    cobertura: float | None = None
    pasos_desarrollo: int | None = None
    duplicados: int | None = None
    detalle: list[dict[str, Any]] = field(default_factory=list)
    solo_qcpulse: list[str] = field(default_factory=list)

    def checks(self) -> dict[str, bool | None]:
        return {
            "granularidad": None if self.ratio is None else META_GRANULARIDAD[0] <= self.ratio <= META_GRANULARIDAD[1],
            "pasos": None if self.steps_ok is None else self.steps_ok >= META_PASOS,
            "cobertura": None if self.cobertura is None else self.cobertura >= META_COBERTURA,
            "desarrollo": None if self.pasos_desarrollo is None else self.pasos_desarrollo == 0,
            "duplicados": None if self.duplicados is None else self.duplicados == 0,
        }


def score(
    luminia: list[BenchCase],
    qcpulse: list[BenchCase],
    judge: bool = False,
    chat: Callable[..., Any] = chat_json,
) -> list[EpicScore]:
    epics = sorted({c.epic for c in luminia} | {c.epic for c in qcpulse})
    results: list[EpicScore] = []
    for epic in epics:
        lum = [c for c in luminia if c.epic == epic]
        qcp = [c for c in qcpulse if c.epic == epic]
        ratio = (len(qcp) / len(lum)) if lum else None
        steps_ok = (sum(1 for c in qcp if 3 <= len(c.steps) <= 6) / len(qcp)) if qcp else None
        result = EpicScore(epic, len(lum), len(qcp), ratio, steps_ok)
        if judge and lum and qcp:
            try:
                verdict = judge_epic(epic, lum, qcp, chat=chat)
            except LlmError as exc:
                logger.warning("juez falló en %s: %s", epic, exc)
                verdict = None
            if verdict:
                rows = [r for r in verdict.get("cobertura") or [] if isinstance(r, dict)]
                good = sum(1 for r in rows if str(r.get("estado")).lower() in {"si", "sí", "parcial"})
                result.cobertura = good / len(lum)
                result.detalle = rows
                result.pasos_desarrollo = len(verdict.get("pasos_desarrollo") or [])
                result.duplicados = len(verdict.get("duplicados") or [])
                referenced = {qid for r in rows for qid in (r.get("qcpulse_ids") or [])}
                result.solo_qcpulse = [f"{c.id} — {c.title}" for c in qcp if c.id not in referenced]
        results.append(result)
    return results


def _pct(value: float | None) -> str:
    return "—" if value is None else f"{value * 100:.0f}%"


def _mark(value: bool | None) -> str:
    return "—" if value is None else ("✔" if value else "✘")


def write_report(results: list[EpicScore], path: str | Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resumen"
    headers = [
        "Épica", "Casos Luminia", "Casos QcPulse", "Granularidad (40–60%)", "✔",
        "Pasos 3–6 (≥90%)", "✔", "Cobertura Luminia (≥90%)", "✔", "Pasos de Desarrollo (0)", "✔", "Duplicados (0)", "✔",
    ]
    sheet.append(headers)
    for r in results:
        c = r.checks()
        sheet.append([
            r.epic, r.luminia, r.qcpulse, _pct(r.ratio), _mark(c["granularidad"]),
            _pct(r.steps_ok), _mark(c["pasos"]), _pct(r.cobertura), _mark(c["cobertura"]),
            "—" if r.pasos_desarrollo is None else r.pasos_desarrollo, _mark(c["desarrollo"]),
            "—" if r.duplicados is None else r.duplicados, _mark(c["duplicados"]),
        ])
    total_l = sum(r.luminia for r in results)
    total_q = sum(r.qcpulse for r in results)
    sheet.append(["Total", total_l, total_q, _pct(total_q / total_l) if total_l else "—"])
    detalle = workbook.create_sheet("Detalle Luminia")
    detalle.append(["Épica", "Caso Luminia", "Estado", "Casos QcPulse", "Nota"])
    for r in results:
        for row in r.detalle:
            detalle.append([r.epic, row.get("luminia_id"), row.get("estado"), ", ".join(row.get("qcpulse_ids") or []), row.get("nota")])
    solo = workbook.create_sheet("Solo QcPulse")
    solo.append(["Épica", "Caso QcPulse sin equivalente en Luminia (revisar: brecha real o relleno)"])
    for r in results:
        for item in r.solo_qcpulse:
            solo.append([r.epic, item])
    for ws in workbook.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 18
    workbook.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Benchmark QcPulse v2 vs Luminia")
    parser.add_argument("--luminia", required=True)
    parser.add_argument("--qcpulse", required=True)
    parser.add_argument("--out", default="/out")
    parser.add_argument("--judge", action="store_true", help="usa OpenAI para cobertura, pasos de Desarrollo y duplicados")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    results = score(read_luminia(args.luminia), read_qcpulse(args.qcpulse), judge=args.judge)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    report = out / "benchmark_qcv2_vs_luminia.xlsx"
    write_report(results, report)
    (out / "benchmark_qcv2_vs_luminia.json").write_text(
        json.dumps([r.__dict__ for r in results], ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    for r in results:
        print(
            f"{r.epic}: Luminia={r.luminia} QcPulse={r.qcpulse} granularidad={_pct(r.ratio)} "
            f"pasos3-6={_pct(r.steps_ok)} cobertura={_pct(r.cobertura)} "
            f"desarrollo={r.pasos_desarrollo} duplicados={r.duplicados}"
        )
    print(f"Reporte: {report}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
