"""Corre la v2 para una o varias Technical Epics sin tocar la base de datos.

Uso (desde la raíz del proyecto, con Docker):
    docker compose run --rm --no-deps -v "$PWD/exports:/out" backend \
        python -m app.services.qc_v2.cli WEBCL-2060 WEBCL-1586 --device WEB --out /out

Escribe en --out:
    qcv2_<etiqueta>.json   todo lo que pasó (entrada, Etapa A, casos, avisos, tokens)
    qcv2_<etiqueta>.xlsx   hojas Resumen, Escenarios, Casos y Cruce para revisión
"""

from __future__ import annotations

import argparse
import dataclasses
import json
import logging
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services.qc_v2.engine import EpicRun, run_epics

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_GAP_FILL = PatternFill("solid", fgColor="FFF2CC")
_ERR_FILL = PatternFill("solid", fgColor="F8CBAD")


def _sheet(workbook: Workbook, title: str, headers: list[str], widths: list[int]):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for index, width in enumerate(widths, start=1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"
    return sheet


def _wrap(sheet) -> None:
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(runs: list[EpicRun]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    resumen = _sheet(
        workbook,
        "Resumen",
        ["Épica", "CRN", "COV", "Escenarios", "Casos", "Requieren revisión", "Llamadas IA", "Tokens entrada", "Tokens salida", "Segundos", "Error / avisos"],
        [14, 7, 7, 11, 8, 12, 11, 14, 13, 10, 70],
    )
    for run in runs:
        scenarios = len(run.stage_a.scenarios) if run.stage_a else 0
        review = sum(1 for c in run.candidates if c.review_required)
        notes = "; ".join(([run.error] if run.error else []) + run.bundle.warnings)
        resumen.append([
            run.bundle.key, len(run.bundle.crns), len(run.bundle.covs), scenarios, len(run.candidates),
            review, run.llm_calls, run.prompt_tokens, run.completion_tokens, round(run.ms / 1000), notes,
        ])
    _wrap(resumen)

    escenarios = _sheet(
        workbook,
        "Escenarios",
        ["Épica", "ID", "Escenario", "Descripción", "Condiciones", "CRN", "COV", "Brecha", "No afectación", "Casos"],
        [14, 8, 45, 60, 35, 20, 20, 12, 12, 7],
    )
    for run in runs:
        if not run.stage_a:
            continue
        counts: dict[str, int] = {}
        for c in run.candidates:
            sid = next((r.split("=", 1)[1] for r in c.applied_rules if r.startswith("qc-v2:escenario=")), "")
            counts[sid] = counts.get(sid, 0) + 1
        for s in run.stage_a.scenarios:
            escenarios.append([
                run.bundle.key, s.id, s.titulo, s.descripcion, "; ".join(s.condiciones),
                ", ".join(s.cubre_negocio), ", ".join(s.cubre_tecnico), s.brecha,
                "sí" if s.es_no_afectacion else "", counts.get(s.id, 0),
            ])
    _wrap(escenarios)

    casos = _sheet(
        workbook,
        "Casos",
        ["Épica", "Escenario", "Título", "Tipo", "Prioridad", "Complejidad", "Precondiciones", "Pasos", "Resultado esperado", "Variantes", "Referencia técnica", "Cubre", "Brecha", "Requiere revisión"],
        [14, 9, 50, 7, 10, 11, 45, 60, 60, 40, 35, 22, 12, 30],
    )
    for run in runs:
        for c in run.candidates:
            rules = {r.split("=", 1)[0]: r.split("=", 1)[1] for r in c.applied_rules if "=" in r}
            casos.append([
                run.bundle.key,
                rules.get("qc-v2:escenario", ""),
                c.name,
                rules.get("qc-v2:tipo", ""),
                c.priority,
                rules.get("qc-v2:complejidad", ""),
                "\n".join(f"- {p.strip()}" for p in (c.precondition or "").split(";") if p.strip()),
                "\n".join(f"{s.step_number}. {s.action}" for s in c.steps),
                "\n".join(f"{s.step_number}. {s.expected_result}" for s in c.steps),
                c.variant_table or "",
                c.technical_reference or "",
                ", ".join(c.covers),
                rules.get("qc-v2:brecha", ""),
                c.applicability_reason or "",
            ])
    _wrap(casos)

    cruce = _sheet(
        workbook,
        "Cruce",
        ["Épica", "Historia", "ID / Escenario", "Situación", "Casos que lo cubren", "Detalle"],
        [14, 14, 60, 22, 60, 50],
    )
    for run in runs:
        for row in run.ledger:
            cruce.append([run.bundle.key, row.get("story"), row.get("scenario"), row.get("estado"), row.get("case_id"), row.get("motivo")])
            estado = str(row.get("estado") or "")
            fill = None
            if estado in {"solo negocio", "solo técnico", "implícito"}:
                fill = _GAP_FILL
            elif estado.startswith("error") or estado.startswith("pendiente") or estado == "aviso":
                fill = _ERR_FILL
            if fill:
                for cell in cruce[cruce.max_row]:
                    cell.fill = fill
    _wrap(cruce)
    return workbook


def _jsonable(run: EpicRun) -> dict[str, Any]:
    return {
        "epic": run.bundle.key,
        "summary": run.summary(),
        "warnings": run.bundle.warnings,
        "error": run.error,
        "bundle": dataclasses.asdict(run.bundle),
        "stage_a": dataclasses.asdict(run.stage_a) if run.stage_a else None,
        "stage_b_errors": {r.scenario_id: r.error for r in run.stage_b if r.error},
        "candidates": [c.model_dump() for c in run.candidates],
        "ledger": run.ledger,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="QcPulse v2: genera casos QC sin tocar la base de datos")
    parser.add_argument("epics", nargs="+", help="Technical Epics, ej. WEBCL-2060 WEBCL-1586")
    parser.add_argument("--device", required=True, help="Dispositivo de la Release, ej. WEB, ADT, ADR")
    parser.add_argument("--region", default=None)
    parser.add_argument("--out", default="/out")
    parser.add_argument("--label", default=None)
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    label = args.label or f"{args.device}_{datetime.now():%Y%m%d_%H%M}"
    runs: list[EpicRun] = []
    for key in [k.strip().upper() for k in args.epics]:
        current = run_epics([key], args.device, args.region)
        runs.extend(current)
        # Se guarda después de cada épica: si la corrida se corta, lo avanzado no se pierde.
        (out / f"qcv2_{label}.json").write_text(
            json.dumps([_jsonable(r) for r in runs], ensure_ascii=False, indent=2, default=str), encoding="utf-8"
        )
        build_workbook(runs).save(out / f"qcv2_{label}.xlsx")
        for run in current:
            print(run.summary(), flush=True)
            for warning in run.bundle.warnings:
                print(f"   aviso: {warning}", flush=True)
        if any((r.error or "").startswith("corrida detenida") for r in current):
            print("Corrida detenida (OpenAI sin créditos). Lo generado hasta aquí quedó guardado.")
            break
    print(f"Listo: {out / f'qcv2_{label}.xlsx'}")
    return 1 if any(r.error for r in runs) else 0


if __name__ == "__main__":
    sys.exit(main())
