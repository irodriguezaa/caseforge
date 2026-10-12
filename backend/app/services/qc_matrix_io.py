"""Import and export a QC base matrix (JSON or the Excel workbook this module writes)."""

from __future__ import annotations

import io
import json
import re
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models.qc_matrix import (
    QcMatrixCase,
    QcMatrixCaseDevice,
    QcMatrixDevice,
    QcMatrixModule,
    QcMatrixModuleLink,
    QcMatrixVersion,
)

_NOTE_PREFIX = re.compile(r"^\[([^\]]+)\]\s*(.*)$", re.DOTALL)


def parse_upload(filename: str, raw: bytes) -> dict[str, Any]:
    name = (filename or "").lower()
    if name.endswith(".json"):
        payload = json.loads(raw.decode("utf-8"))
        if not isinstance(payload, dict) or "matrix" not in payload:
            raise ValueError("El JSON no tiene el objeto matrix.")
        return payload
    if name.endswith(".xlsx"):
        return workbook_to_payload(raw)
    raise ValueError("El archivo debe ser .json o .xlsx.")


def workbook_to_payload(raw: bytes) -> dict[str, Any]:
    book = load_workbook(io.BytesIO(raw), data_only=True)
    matrix_rows = _sheet_dicts(book, "Matrix")
    if not matrix_rows:
        raise ValueError("El Excel no tiene la hoja Matrix.")
    header = matrix_rows[0]
    modules = _sheet_dicts(book, "Modules")
    devices = _sheet_dicts(book, "Devices")
    cases = _sheet_dicts(book, "Cases")
    return {
        "matrix": {
            "code": header.get("code") or "",
            "version": header.get("version") or "",
            "name": header.get("name") or "",
            "generated": header.get("generated") or "",
            "platforms": _split_pipe(header.get("platforms")),
        },
        "scope_rules": _loads_obj(header.get("scope_rules")),
        "providers": _split_pipe(header.get("providers")),
        "modules": [
            {
                "id": row.get("id") or "",
                "group": row.get("group") or "",
                "name": row.get("name") or "",
                "description": row.get("description") or "",
                "related": _split_pipe(row.get("related")),
            }
            for row in modules
            if (row.get("id") or "").strip()
        ],
        "devices": [
            {
                "code": row.get("code") or "",
                "label": row.get("label") or "",
                "screen": row.get("screen") or "",
                "mdp_mx": _split_pipe(row.get("mdp_mx")),
                "mdp_other_regions": _split_pipe(row.get("mdp_other_regions")),
            }
            for row in devices
            if (row.get("code") or "").strip()
        ],
        "cases": [_case_from_sheet(row) for row in cases if (row.get("title") or "").strip()],
    }


def import_matrix(db: Session, payload: dict[str, Any]) -> QcMatrixVersion:
    matrix = payload.get("matrix") or {}
    code = str(matrix.get("code") or "").strip()
    version = str(matrix.get("version") or "").strip()
    name = str(matrix.get("name") or code or "Matriz QC").strip()
    if not code or not version:
        raise ValueError("La matriz necesita código y versión.")

    row = db.execute(
        select(QcMatrixVersion)
        .where(QcMatrixVersion.code == code, QcMatrixVersion.version == version)
        .options(
            selectinload(QcMatrixVersion.modules).selectinload(QcMatrixModule.relations),
            selectinload(QcMatrixVersion.devices),
            selectinload(QcMatrixVersion.cases).selectinload(QcMatrixCase.device_links),
        )
    ).scalar_one_or_none()
    if row is None:
        row = QcMatrixVersion(code=code, version=version, name=name, status="draft")
        db.add(row)
        db.flush()
    row.name = name
    row.generated_on = _parse_date(matrix.get("generated"))
    row.providers = list(payload.get("providers") or [])
    row.scope_rules = dict(payload.get("scope_rules") or {})

    modules_by_key = {module.module_key: module for module in row.modules}
    for index, item in enumerate(payload.get("modules") or []):
        key = str(item.get("id") or "").strip()
        if not key:
            continue
        module = modules_by_key.get(key)
        if module is None:
            module = QcMatrixModule(matrix_version_id=row.id, module_key=key, name=key)
            row.modules.append(module)
            db.flush()
            modules_by_key[key] = module
        module.group_name = str(item.get("group") or "")
        module.name = str(item.get("name") or key)
        module.description = str(item.get("description") or "")
        related = [str(value).strip() for value in (item.get("related") or []) if str(value).strip()]
        module.relations.clear()
        db.flush()
        for related_key in dict.fromkeys(related):
            module.relations.append(QcMatrixModuleLink(related_key=related_key))

    devices_by_code = {device.code: device for device in row.devices}
    device_items = list(payload.get("devices") or [])
    if not device_items:
        device_items = [
            {"code": code_name, "label": code_name, "screen": "", "mdp_mx": [], "mdp_other_regions": []}
            for code_name in (matrix.get("platforms") or [])
            if str(code_name).strip()
        ]
    for index, item in enumerate(device_items):
        device_code = str(item.get("code") or "").strip()
        if not device_code:
            continue
        device = devices_by_code.get(device_code)
        if device is None:
            device = QcMatrixDevice(matrix_version_id=row.id, code=device_code, label=device_code)
            row.devices.append(device)
            db.flush()
            devices_by_code[device_code] = device
        device.label = str(item.get("label") or device_code)
        device.screen = str(item.get("screen") or "")
        device.mdp_mx = list(item.get("mdp_mx") or [])
        device.mdp_other_regions = list(item.get("mdp_other_regions") or [])
        device.sort_order = index

    cases_by_base = {case.base_id: case for case in row.cases if case.base_id}
    for item in payload.get("cases") or []:
        base_id = str(item.get("id") or "").strip() or None
        case = cases_by_base.get(base_id) if base_id else None
        if case is None:
            case = QcMatrixCase(matrix_version_id=row.id, base_id=base_id, title="")
            row.cases.append(case)
            db.flush()
            if base_id:
                cases_by_base[base_id] = case
        module_key = str(item.get("module_id") or "").strip()
        module = modules_by_key.get(module_key)
        case.module_id = module.id if module is not None else None
        case.title = str(item.get("title") or "")
        case.steps = str(item.get("steps") or "")
        case.expected = str(item.get("expected") or "")
        case.severity = str(item.get("severity") or "Critical")
        case.smoke = bool(item.get("smoke"))
        region = item.get("region")
        case.region = str(region).strip() if region else None
        case.parametrized = bool(item.get("parametrized"))
        case.param_values = str(item.get("param_values") or "")
        case.user_types = list(item.get("user_types") or [])
        case.providers_in_source = list(item.get("providers_in_source") or [])
        case.case_status = str(item.get("status") or "")
        case.sources = list(item.get("sources") or [])
        case.origin = str(item.get("origin") or "")
        case.ott_equivalent = str(item.get("ott_equivalent") or "")
        case.review_reasons = list(item.get("review_reasons") or [])
        case.approvals = list(item.get("approvals") or [])
        _sync_case_devices(db, case, item, devices_by_code, row)

    db.flush()
    return row


def publish_version(db: Session, matrix: QcMatrixVersion, published_by: str) -> QcMatrixVersion:
    others = db.execute(
        select(QcMatrixVersion).where(
            QcMatrixVersion.code == matrix.code,
            QcMatrixVersion.id != matrix.id,
            QcMatrixVersion.status == "published",
        )
    ).scalars().all()
    for other in others:
        other.status = "historical"
    matrix.status = "published"
    matrix.published_at = datetime.now().astimezone()
    matrix.published_by = published_by
    db.flush()
    return matrix


def version_to_workbook(matrix: QcMatrixVersion) -> bytes:
    book = Workbook()
    matrix_sheet = book.active
    matrix_sheet.title = "Matrix"
    matrix_sheet.append(
        ["code", "version", "name", "generated", "platforms", "providers", "scope_rules", "status"]
    )
    platforms = [device.code for device in sorted(matrix.devices, key=lambda item: item.sort_order)]
    matrix_sheet.append(
        [
            matrix.code,
            matrix.version,
            matrix.name,
            matrix.generated_on.isoformat() if matrix.generated_on else "",
            _join_pipe(platforms),
            _join_pipe(matrix.providers or []),
            json.dumps(matrix.scope_rules or {}, ensure_ascii=False),
            matrix.status,
        ]
    )

    modules_sheet = book.create_sheet("Modules")
    modules_sheet.append(["id", "group", "name", "description", "related"])
    for module in sorted(matrix.modules, key=lambda item: item.module_key):
        modules_sheet.append(
            [
                module.module_key,
                module.group_name,
                module.name,
                module.description,
                _join_pipe(link.related_key for link in module.relations),
            ]
        )

    devices_sheet = book.create_sheet("Devices")
    devices_sheet.append(["code", "label", "screen", "mdp_mx", "mdp_other_regions"])
    for device in sorted(matrix.devices, key=lambda item: item.sort_order):
        devices_sheet.append(
            [
                device.code,
                device.label,
                device.screen,
                _join_pipe(device.mdp_mx or []),
                _join_pipe(device.mdp_other_regions or []),
            ]
        )

    cases_sheet = book.create_sheet("Cases")
    cases_sheet.append(
        [
            "id",
            "module_id",
            "title",
            "user_types",
            "steps",
            "expected",
            "severity",
            "smoke",
            "platforms",
            "region",
            "parametrized",
            "param_values",
            "platform_notes",
            "status",
            "sources",
            "origin",
            "ott_equivalent",
            "providers_in_source",
            "review_reasons",
            "approvals",
        ]
    )
    for case in sorted(matrix.cases, key=lambda item: item.base_id or ""):
        platform_codes = [link.device.code for link in case.device_links]
        notes = [
            f"[{link.device.code}] {link.note}" if link.note else ""
            for link in case.device_links
            if link.note
        ]
        cases_sheet.append(
            [
                case.base_id or "",
                case.module.module_key if case.module else "",
                case.title,
                _join_pipe(case.user_types or []),
                case.steps,
                case.expected,
                case.severity,
                "true" if case.smoke else "false",
                _join_pipe(platform_codes),
                case.region or "",
                "true" if case.parametrized else "false",
                case.param_values,
                "\n".join(notes),
                case.case_status,
                json.dumps(case.sources or [], ensure_ascii=False),
                case.origin,
                case.ott_equivalent,
                _join_pipe(case.providers_in_source or []),
                json.dumps(case.review_reasons or [], ensure_ascii=False),
                json.dumps(case.approvals or [], ensure_ascii=False),
            ]
        )

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _sync_case_devices(
    db: Session,
    case: QcMatrixCase,
    item: dict[str, Any],
    devices_by_code: dict[str, QcMatrixDevice],
    matrix: QcMatrixVersion,
) -> None:
    notes_by_code: dict[str, list[str]] = {}
    for raw_note in item.get("platform_notes") or []:
        for code, text in _split_platform_note(str(raw_note)):
            notes_by_code.setdefault(code, []).append(text)
    platforms = [str(code).strip() for code in (item.get("platforms") or []) if str(code).strip()]
    for code in notes_by_code:
        if code not in platforms and code in devices_by_code:
            platforms.append(code)
    case.device_links.clear()
    db.flush()
    seen: set[str] = set()
    for code in platforms:
        if code in seen:
            continue
        seen.add(code)
        device = devices_by_code.get(code)
        if device is None:
            device = QcMatrixDevice(
                matrix_version_id=matrix.id,
                code=code,
                label=code,
                sort_order=len(devices_by_code),
            )
            matrix.devices.append(device)
            db.flush()
            devices_by_code[code] = device
        case.device_links.append(
            QcMatrixCaseDevice(device_id=device.id, note="\n".join(notes_by_code.get(code) or []))
        )


def _split_platform_note(note: str) -> list[tuple[str, str]]:
    match = _NOTE_PREFIX.match(note.strip())
    if match is None:
        return []
    text = match.group(2).strip()
    return [(code.strip(), text) for code in match.group(1).split("/") if code.strip()]


def _sheet_dicts(book, title: str) -> list[dict[str, str]]:
    if title not in book.sheetnames:
        return []
    sheet = book[title]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell or "").strip() for cell in rows[0]]
    parsed: list[dict[str, str]] = []
    for row in rows[1:]:
        if not any(cell not in (None, "") for cell in row):
            continue
        parsed.append(
            {
                headers[index]: "" if value is None else str(value)
                for index, value in enumerate(row)
                if index < len(headers) and headers[index]
            }
        )
    return parsed


def _case_from_sheet(row: dict[str, str]) -> dict[str, Any]:
    return {
        "id": (row.get("id") or "").strip(),
        "module_id": row.get("module_id") or "",
        "title": row.get("title") or "",
        "user_types": _split_pipe(row.get("user_types")),
        "steps": row.get("steps") or "",
        "expected": row.get("expected") or "",
        "severity": row.get("severity") or "Critical",
        "smoke": str(row.get("smoke") or "").strip().lower() in {"1", "true", "yes", "si", "sí"},
        "platforms": _split_pipe(row.get("platforms")),
        "region": (row.get("region") or "").strip() or None,
        "parametrized": str(row.get("parametrized") or "").strip().lower() in {"1", "true", "yes", "si", "sí"},
        "param_values": row.get("param_values") or "",
        "platform_notes": [line for line in str(row.get("platform_notes") or "").splitlines() if line.strip()],
        "status": row.get("status") or "",
        "sources": _loads_list(row.get("sources")),
        "origin": row.get("origin") or "",
        "ott_equivalent": row.get("ott_equivalent") or "",
        "providers_in_source": _split_pipe(row.get("providers_in_source")),
        "review_reasons": _loads_list(row.get("review_reasons")),
        "approvals": _loads_list(row.get("approvals")),
    }


def _split_pipe(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item) for item in value if str(item).strip()]
    return [part.strip() for part in str(value).split("|") if part.strip()]


def _join_pipe(values) -> str:
    return " | ".join(str(value) for value in values if str(value).strip())


def _loads_list(value: Any) -> list:
    if not value:
        return []
    if isinstance(value, list):
        return value
    try:
        loaded = json.loads(str(value))
    except json.JSONDecodeError:
        return []
    return loaded if isinstance(loaded, list) else []


def _loads_obj(value: Any) -> dict:
    if not value:
        return {}
    if isinstance(value, dict):
        return value
    try:
        loaded = json.loads(str(value))
    except json.JSONDecodeError:
        return {}
    return loaded if isinstance(loaded, dict) else {}


def _parse_date(value: Any) -> date | None:
    if not value:
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None
