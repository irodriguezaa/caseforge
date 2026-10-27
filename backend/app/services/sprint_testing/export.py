"""Excel of execution issues and executive PPTX for Sprint Testing.

Each SWF gets two slides: epic status bars (same reading as the dashboard)
and open blockers grouped by parent epic, top 5 per device.
"""

from __future__ import annotations

import io
from typing import Any, NamedTuple

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from app.schemas.sprint_testing import ExecutionBlockerRow, ExecutionIssueRow, SprintTestingRead

ISSUE_HEADERS = [
    "Key",
    "Summary",
    "Issue Type",
    "Prioridad",
    "Estado",
    "Creación",
    "Principal",
    "Cambio de estado",
    "Último comentario",
    "Anidados",
]

_HEADER_FILL = PatternFill("solid", fgColor="1B2A49")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_CELL_FONT = Font(name="Calibri", size=11)
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_EXCEL_MAX_CELL = 32767

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
NAVY = RGBColor(0x1B, 0x2A, 0x49)
TEAL = RGBColor(0x0D, 0x8B, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF4, 0xF8, 0xFB)
MUTED = RGBColor(0x6B, 0x7A, 0x8D)
BODY = RGBColor(0x3A, 0x47, 0x5C)
LINE = RGBColor(0xE1, 0xE7, 0xEE)
RED = RGBColor(0xC4, 0x5C, 0x5C)
DEV_COLOR = RGBColor(0xC4, 0xB5, 0x9A)
INTEGRATION_COLOR = RGBColor(0x7B, 0x6D, 0x8D)
QA_COLOR = RGBColor(0x4C, 0x7D, 0x9E)
VALIDATION_COLOR = RGBColor(0x7D, 0x87, 0x94)
CLOSED_COLOR = RGBColor(0x4A, 0x9B, 0x82)
OTHER_COLOR = RGBColor(0x4C, 0x6A, 0x8A)
_TOP_EPICS = 5
_TESTING_ORDER = ("QA Validation", "Validation", "Validate QC")
_STATUS_BARS = (
    ("development", "Desarrollo"),
    ("integration", "Integración"),
    ("testing", "Testing"),
    ("closed", "Cerrado"),
)


def _excel_value(value: object) -> str:
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    if len(text) > _EXCEL_MAX_CELL:
        return text[: _EXCEL_MAX_CELL - 1]
    return text


def _sheet_title(label: str) -> str:
    cleaned = "".join(ch if ch not in r":\/?*[]" else " " for ch in (label or "Issues")).strip()
    return (cleaned or "Issues")[:31]


def _write_issues_sheet(sheet: Any, rows: list[ExecutionIssueRow]) -> None:
    for index, header in enumerate(ISSUE_HEADERS, start=1):
        cell = sheet.cell(1, index, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = _THIN
    widths = [16, 48, 14, 16, 16, 16, 14, 36, 32, 22]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    last_row = max(len(rows) + 1, 1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(ISSUE_HEADERS))}{last_row}"

    for row_index, row in enumerate(rows, start=2):
        values = [
            row.key,
            row.summary,
            row.issue_type,
            row.priority,
            row.status,
            row.created,
            row.principal,
            row.status_change,
            row.last_comment,
            row.nested_issues,
        ]
        wrap_cols = {2, 8, 9, 10}
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col, _excel_value(value))
            cell.font = _CELL_FONT
            cell.border = _THIN
            cell.alignment = Alignment(wrap_text=col in wrap_cols, vertical="top")
        sheet.row_dimensions[row_index].height = 36


def build_issues_workbook(rows: list[ExecutionIssueRow], sheet_title: str = "Issues") -> bytes:
    return build_issues_workbook_sheets([(sheet_title, rows)])


def build_issues_workbook_sheets(sheets: list[tuple[str, list[ExecutionIssueRow]]]) -> bytes:
    workbook = Workbook()
    if not sheets:
        sheets = [("Issues", [])]
    for index, (title, rows) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = _sheet_title(title)
        _write_issues_sheet(sheet, rows)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class _EpicRank(NamedTuple):
    key: str
    title: str
    count: int


_STATUS_LEGEND = (
    ("Desarrollo", DEV_COLOR),
    ("Integración", INTEGRATION_COLOR),
    ("QA Validation", QA_COLOR),
    ("Validation", VALIDATION_COLOR),
    ("Cerrado", CLOSED_COLOR),
)
_ISSUE_LEGEND = (
    ("Blocker", RED),
    ("No Blocker", OTHER_COLOR),
)


def _clip(text: str, limit: int) -> str:
    compact = " ".join((text or "").split())
    if len(compact) <= limit:
        return compact
    return compact[: limit - 1].rstrip() + "…"


def _sum_bucket(values: dict[str, int]) -> int:
    return sum(values.values())


def _closed_pct(open_count: int, closed_count: int) -> str:
    total = open_count + closed_count
    if total <= 0:
        return "—"
    return f"{round(closed_count * 100 / total, 1)}%"


def _count_label(count: int, noun: str) -> str:
    return f"1 {noun}" if count == 1 else f"{count} {noun}s"


def _rank_blocker_epics(
    blockers: list[ExecutionBlockerRow],
    limit: int = _TOP_EPICS,
) -> tuple[list[_EpicRank], str]:
    """Top parent epics by open blocker count. Singles past the cap are called out."""
    buckets: dict[str, dict[str, Any]] = {}
    for row in blockers:
        key = (row.principal or "").strip().upper()
        slot = buckets.setdefault(key, {"count": 0, "title": ""})
        slot["count"] += 1
        title = (row.principal_summary or "").strip()
        if title and not slot["title"]:
            slot["title"] = title
    ranked = [
        _EpicRank(
            key or "—",
            slot["title"] or ("Sin principal" if not key else ""),
            int(slot["count"]),
        )
        for key, slot in buckets.items()
    ]
    ranked.sort(key=lambda item: (-item.count, item.key))
    shown = ranked[:limit]
    omitted = ranked[limit:]
    if not omitted:
        note = ""
    elif all(item.count == 1 for item in omitted):
        note = f"+ {len(omitted)} con 1 blocker"
    elif len(omitted) == 1:
        note = "+ 1 épica más"
    else:
        note = f"+ {len(omitted)} épicas más"
    return shown, note


def _add_rect(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    fill: RGBColor,
    line: RGBColor | None = None,
) -> Any:
    if int(width) <= 0 or int(height) <= 0:
        return None
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(left), int(top), int(width), int(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.shadow.inherit = False
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def _style_run(run: Any, text: str, size: int, bold: bool, color: RGBColor) -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _prepare_frame(frame: Any, margin: float = 0.02) -> None:
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0)


def _add_text(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    *,
    size: int,
    bold: bool = False,
    color: RGBColor = WHITE,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    frame = box.text_frame
    _prepare_frame(frame)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    _style_run(paragraph.add_run(), text, size, bold, color)


def _metric_card(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    label: str,
    value: str,
    value_color: RGBColor,
) -> None:
    _add_rect(slide, left, top, width, height, OFF, line=LINE)
    box = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    label_p = frame.paragraphs[0]
    label_p.alignment = PP_ALIGN.CENTER
    label_p.space_before = Pt(0)
    label_p.space_after = Pt(0)
    _style_run(label_p.add_run(), label, 10, False, MUTED)
    value_p = frame.add_paragraph()
    value_p.alignment = PP_ALIGN.CENTER
    value_p.space_before = Pt(2)
    value_p.space_after = Pt(0)
    _style_run(value_p.add_run(), value, 20, True, value_color)


def _legend(slide: Any, left: float, top: float, items: tuple[tuple[str, RGBColor], ...]) -> None:
    x = int(left)
    y = int(top)
    swatch = int(Inches(0.12))
    for label, color in items:
        _add_rect(slide, x, y + int(Inches(0.05)), swatch, swatch, color)
        text_w = int(Inches(0.075 * max(len(label), 1) + 0.16))
        _add_text(slide, x + int(Inches(0.16)), y, text_w, int(Inches(0.24)), label, size=10, color=MUTED)
        x += int(Inches(0.18)) + text_w


def _header(slide: Any, payload: SprintTestingRead, section: str) -> None:
    _add_rect(slide, 0, 0, SLIDE_W, Inches(0.66), NAVY)
    _add_text(slide, Inches(0.36), Inches(0.05), Inches(7.2), Inches(0.22), "Reporte ejecutivo", size=11, color=TEAL)
    _add_text(
        slide,
        Inches(0.36),
        Inches(0.26),
        Inches(8.4),
        Inches(0.34),
        f"{payload.sprint.label}  ·  {payload.swf}",
        size=20,
        bold=True,
    )
    _add_text(
        slide,
        Inches(8.4),
        Inches(0.2),
        Inches(4.55),
        Inches(0.3),
        section,
        size=14,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def _bar_segments(bucket: str, counts: dict[str, int]) -> list[tuple[int, RGBColor]]:
    total = _sum_bucket(counts)
    if total <= 0:
        return []
    if bucket == "development":
        return [(total, DEV_COLOR)]
    if bucket == "integration":
        return [(total, INTEGRATION_COLOR)]
    if bucket == "closed":
        return [(total, CLOSED_COLOR)]
    colors = {
        "QA Validation": QA_COLOR,
        "Validation": VALIDATION_COLOR,
        "Validate QC": VALIDATION_COLOR,
    }
    segments: list[tuple[int, RGBColor]] = []
    seen: set[str] = set()
    for label in _TESTING_ORDER:
        value = counts.get(label, 0)
        if value:
            segments.append((value, colors[label]))
            seen.add(label)
    for label, value in counts.items():
        if label not in seen and value:
            segments.append((value, VALIDATION_COLOR))
    return segments


def _group_origin(count: int, left: int, width: int, pad: int) -> tuple[int, int]:
    inner_left = left + pad
    inner_w = max(width - 2 * pad, 0)
    if count <= 0:
        return inner_left, inner_w
    group_w = min(inner_w // count, int(Inches(2.7)))
    row_w = group_w * count
    return inner_left + max(inner_w - row_w, 0) // 2, group_w


def _draw_status_chart(
    slide: Any,
    programs: list[Any],
    left: float,
    top: float,
    width: float,
    height: float,
) -> None:
    _add_rect(slide, left, top, width, height, OFF, line=LINE)
    if not programs:
        _add_text(
            slide,
            int(left) + int(Inches(0.2)),
            int(top) + int(Inches(0.35)),
            int(width) - int(Inches(0.4)),
            int(Inches(0.36)),
            "Sin Technical Epics para este SWF.",
            size=12,
            color=MUTED,
        )
        return
    pad = int(Inches(0.14))
    name_h = int(Inches(0.3))
    number_h = int(Inches(0.2))
    plot_bottom = int(top) + int(height) - pad - name_h
    plot_top = int(top) + pad + number_h
    plot_h = max(plot_bottom - plot_top, 1)
    origin, group_w = _group_origin(len(programs), int(left), int(width), pad)
    max_value = max(
        1,
        max(
            _sum_bucket(getattr(program, bucket) or {})
            for program in programs
            for bucket, _name in _STATUS_BARS
        ),
    )
    bar_w = min(int(Inches(0.34)), max(int(Inches(0.14)), (group_w - int(Inches(0.36))) // 4))
    gap = int(Inches(0.06))
    _add_rect(slide, origin, plot_bottom, group_w * len(programs), int(Inches(0.012)), LINE)
    for index, program in enumerate(programs):
        group_left = origin + index * group_w
        if index:
            _add_rect(slide, group_left, plot_top, int(Inches(0.01)), plot_bottom - plot_top + name_h, LINE)
        used = 4 * bar_w + 3 * gap
        x0 = group_left + max(group_w - used, 0) // 2
        for bar_index, (bucket, _name) in enumerate(_STATUS_BARS):
            segments = _bar_segments(bucket, getattr(program, bucket) or {})
            total = sum(value for value, _color in segments)
            x = x0 + bar_index * (bar_w + gap)
            y = plot_bottom
            for value, color in segments:
                seg_h = int(plot_h * value / max_value)
                if seg_h <= 0:
                    continue
                y -= seg_h
                _add_rect(slide, x, y, bar_w, seg_h, color)
            if total:
                _add_text(
                    slide,
                    x - int(Inches(0.08)),
                    y - number_h,
                    bar_w + int(Inches(0.16)),
                    number_h,
                    str(total),
                    size=10,
                    bold=True,
                    color=NAVY,
                    align=PP_ALIGN.CENTER,
                )
        _add_text(
            slide,
            group_left,
            plot_bottom + int(Inches(0.04)),
            group_w,
            name_h,
            program.display_name,
            size=11,
            bold=True,
            color=NAVY,
            align=PP_ALIGN.CENTER,
        )


def _draw_issue_chart(
    slide: Any,
    programs: list[Any],
    left: float,
    top: float,
    width: float,
    height: float,
    empty: str | None = None,
) -> None:
    _add_rect(slide, left, top, width, height, OFF, line=LINE)
    if empty or not programs:
        _add_text(
            slide,
            int(left) + int(Inches(0.2)),
            int(top) + int(Inches(0.28)),
            int(width) - int(Inches(0.4)),
            int(Inches(0.36)),
            empty or "No hay issues de ejecución para este SWF.",
            size=12,
            color=MUTED,
        )
        return
    pad = int(Inches(0.12))
    name_h = int(Inches(0.28))
    number_h = int(Inches(0.2))
    plot_bottom = int(top) + int(height) - pad - name_h
    plot_top = int(top) + pad + number_h
    plot_h = max(plot_bottom - plot_top, 1)
    origin, group_w = _group_origin(len(programs), int(left), int(width), pad)
    max_value = max(1, max(max(program.blocker, program.non_blocker) for program in programs))
    bar_w = min(int(Inches(0.42)), max(int(Inches(0.18)), (group_w - int(Inches(0.5))) // 2))
    gap = int(Inches(0.14))
    _add_rect(slide, origin, plot_bottom, group_w * len(programs), int(Inches(0.012)), LINE)
    for index, program in enumerate(programs):
        group_left = origin + index * group_w
        if index:
            _add_rect(slide, group_left, plot_top, int(Inches(0.01)), plot_bottom - plot_top + name_h, LINE)
        bars = ((program.blocker, RED), (program.non_blocker, OTHER_COLOR))
        used = 2 * bar_w + gap
        x0 = group_left + max(group_w - used, 0) // 2
        for bar_index, (value, color) in enumerate(bars):
            x = x0 + bar_index * (bar_w + gap)
            seg_h = int(plot_h * value / max_value) if value else 0
            if seg_h > 0:
                _add_rect(slide, x, plot_bottom - seg_h, bar_w, seg_h, color)
                _add_text(
                    slide,
                    x - int(Inches(0.1)),
                    plot_bottom - seg_h - number_h,
                    bar_w + int(Inches(0.2)),
                    number_h,
                    str(value),
                    size=10,
                    bold=True,
                    color=color,
                    align=PP_ALIGN.CENTER,
                )
        _add_text(
            slide,
            group_left,
            plot_bottom + int(Inches(0.03)),
            group_w,
            name_h,
            program.display_name,
            size=11,
            bold=True,
            color=NAVY,
            align=PP_ALIGN.CENTER,
        )


def _epic_item(slide: Any, left: int, top: int, width: int, height: int, item: _EpicRank, clip: int) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    _prepare_frame(frame)
    head = frame.paragraphs[0]
    head.alignment = PP_ALIGN.LEFT
    head.space_before = Pt(0)
    head.space_after = Pt(0)
    _style_run(head.add_run(), str(item.count), 13, True, NAVY)
    if item.key != "—":
        _style_run(head.add_run(), f"  {item.key}", 11, True, NAVY)
    title = _clip(item.title, clip)
    if title:
        line = frame.add_paragraph()
        line.alignment = PP_ALIGN.LEFT
        line.space_before = Pt(0)
        line.space_after = Pt(0)
        _style_run(line.add_run(), title, 10, False, BODY)


def _draw_epic_columns(slide: Any, execution: Any) -> None:
    programs = list(execution.programs)
    by_key: dict[str, list[ExecutionBlockerRow]] = {}
    for row in execution.blockers:
        by_key.setdefault((row.program_key or "").upper(), []).append(row)
    count = len(programs)
    left = int(Inches(0.36))
    width = int(Inches(12.6))
    gap = int(Inches(0.1))
    col_w = max((width - gap * (count - 1)) // count, int(Inches(1.4)))
    top = int(Inches(1.46))
    clip = 72 if count <= 3 else 54 if count <= 5 else 36
    item_h = int(Inches(0.56))
    for index, program in enumerate(programs):
        x = left + index * (col_w + gap)
        rows = by_key.get(program.program_key.upper(), [])
        _add_text(slide, x, top, col_w, int(Inches(0.26)), program.display_name, size=12, bold=True, color=NAVY)
        _add_text(
            slide,
            x,
            top + int(Inches(0.24)),
            col_w,
            int(Inches(0.2)),
            _count_label(program.blocker, "blocker"),
            size=10,
            color=MUTED,
        )
        _add_rect(slide, x, top + int(Inches(0.46)), max(col_w - int(Inches(0.08)), 1), int(Inches(0.015)), TEAL)
        ranked, note = _rank_blocker_epics(rows)
        if not ranked:
            _add_text(slide, x, top + int(Inches(0.56)), col_w, int(Inches(0.28)), "Sin blockers", size=11, color=MUTED)
            continue
        y = top + int(Inches(0.54))
        for item in ranked:
            _epic_item(slide, x, y, col_w, item_h, item, clip)
            y += item_h
        if note:
            _add_text(slide, x, y, col_w, int(Inches(0.22)), note, size=9, color=MUTED)


def _notes_box(slide: Any) -> None:
    _add_text(slide, Inches(0.36), Inches(5.28), Inches(6), Inches(0.26), "Observaciones", size=13, bold=True, color=NAVY)
    shape = _add_rect(slide, Inches(0.36), Inches(5.56), Inches(12.6), Inches(1.66), OFF, line=LINE)
    if shape is None:
        return
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.14)
    frame.margin_right = Inches(0.14)
    frame.margin_top = Inches(0.1)
    frame.margin_bottom = Inches(0.08)


def _build_status_slide(slide: Any, payload: SprintTestingRead) -> None:
    _header(slide, payload, "Technical Epics")
    open_count = sum(row.open for row in payload.programs)
    closed_count = sum(row.closed_total for row in payload.programs)
    execution = payload.execution
    blocker = sum(row.blocker for row in execution.programs) if execution else None
    non_blocker = sum(row.non_blocker for row in execution.programs) if execution else None
    metrics = [
        ("Technical Epics", str(payload.technical_epic_count), NAVY),
        ("Abiertos", str(open_count), NAVY),
        ("Cerrados", str(closed_count), CLOSED_COLOR),
        ("% cerrado", _closed_pct(open_count, closed_count), TEAL),
        ("Issues", "—" if execution is None else str(execution.issue_count), NAVY),
        ("Blocker", "—" if blocker is None else str(blocker), RED if blocker else NAVY),
        ("No Blocker", "—" if non_blocker is None else str(non_blocker), NAVY),
    ]
    left = int(Inches(0.32))
    gap = int(Inches(0.08))
    card_w = (int(SLIDE_W) - left * 2 - gap * (len(metrics) - 1)) // len(metrics)
    card_h = int(Inches(0.76))
    card_top = int(Inches(0.78))
    for index, (label, value, color) in enumerate(metrics):
        _metric_card(slide, left + index * (card_w + gap), card_top, card_w, card_h, label, value, color)

    _add_text(slide, Inches(0.36), Inches(1.66), Inches(8), Inches(0.26), "Estado de Technical Epics", size=14, bold=True, color=NAVY)
    _add_text(
        slide,
        Inches(0.36),
        Inches(1.92),
        Inches(12.4),
        Inches(0.24),
        "Por dispositivo. En Testing, cada color es el estado real en Jira.",
        size=11,
        color=MUTED,
    )
    _draw_status_chart(slide, payload.programs, Inches(0.32), Inches(2.2), Inches(12.7), Inches(2.16))
    _legend(slide, Inches(0.36), Inches(4.42), _STATUS_LEGEND)

    _add_text(slide, Inches(0.36), Inches(4.74), Inches(8), Inches(0.24), "Issues en ejecución", size=14, bold=True, color=NAVY)
    if execution is None:
        _draw_issue_chart(
            slide,
            [],
            Inches(0.32),
            Inches(5.08),
            Inches(12.7),
            Inches(1.15),
            empty="Sin filtro de issues en ejecución.",
        )
    else:
        _add_text(
            slide,
            Inches(0.36),
            Inches(4.98),
            Inches(12.4),
            Inches(0.22),
            "Blocker y no Blocker, por dispositivo.",
            size=11,
            color=MUTED,
        )
        _draw_issue_chart(slide, execution.programs, Inches(0.32), Inches(5.24), Inches(12.7), Inches(1.48))
        _legend(slide, Inches(0.36), Inches(6.76), _ISSUE_LEGEND)
    execution_filter = execution.filter_id if execution else "—"
    _add_text(
        slide,
        Inches(0.36),
        Inches(7.14),
        Inches(12.6),
        Inches(0.24),
        f"Filtros Jira {payload.sprint.filter_id} / {execution_filter}",
        size=9,
        color=MUTED,
    )


def _build_blocker_slide(slide: Any, payload: SprintTestingRead) -> None:
    _header(slide, payload, "Blockers por épica")
    _add_text(
        slide,
        Inches(0.36),
        Inches(0.76),
        Inches(12.6),
        Inches(0.28),
        "Blockers por Technical Epic",
        size=16,
        bold=True,
        color=NAVY,
    )
    _add_text(
        slide,
        Inches(0.36),
        Inches(1.06),
        Inches(12.6),
        Inches(0.32),
        "Blockers abiertos por épica (issue principal), por dispositivo. Top 5, con el título de la épica.",
        size=12,
        color=MUTED,
    )
    execution = payload.execution
    if execution is None or not execution.programs:
        _add_text(
            slide,
            Inches(0.36),
            Inches(1.6),
            Inches(12),
            Inches(0.36),
            "Sin issues de ejecución para este SWF.",
            size=13,
            color=MUTED,
        )
    else:
        _draw_epic_columns(slide, execution)
    _notes_box(slide)


def build_executive_pptx(payload: SprintTestingRead) -> bytes:
    return build_executive_pptx_pack([payload])


def build_executive_pptx_pack(payloads: list[SprintTestingRead]) -> bytes:
    presentation = Presentation()
    presentation.slide_width = SLIDE_W
    presentation.slide_height = SLIDE_H
    if not payloads:
        presentation.slides.add_slide(presentation.slide_layouts[6])
    for payload in payloads:
        blank = presentation.slide_layouts[6]
        _build_status_slide(presentation.slides.add_slide(blank), payload)
        _build_blocker_slide(presentation.slides.add_slide(blank), payload)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()

