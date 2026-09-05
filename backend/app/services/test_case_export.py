"""Excel export of persisted Test Cases: QC sheet + Zephyr-flat sheet."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload
from sqlalchemy import select

from app.models.test_case import TestCase


def _display_case_name(case: TestCase) -> str:
    name = case.test_case_name or ""
    device = (case.device or "").strip()
    if not device:
        return name
    for sep in (f" · {device}", f" - {device}", f" | {device}"):
        if name.endswith(sep):
            return name[: -len(sep)].rstrip()
    return name

QC_HEADERS = [
    "ID",
    "Nombre",
    "Componente",
    "Ecosistema",
    "Dispositivo",
    "Fuente dispositivo",
    "Prioridad",
    "Tipo",
    "Tipo de Usuario",
    "Estado",
    "Test Steps",
    "Resultado Esperado",
    "Datos de Prueba",
    "Requiere Condición",
    "Evidencia",
    "Justificación",
    "Technical Epic",
    "Technical Story",
    "Scenario / origen",
    "Confianza IA",
    "Complejidad IA",
    "Estimación IA (h)",
]

# Matches the FLAT Zephyr layout that CaseForge already imports (see services/imports.py).
ZEPHYR_HEADERS = [
    "Test Case ID",
    "Component",
    "Test Case Name",
    "Description",
    "User Type",
    "Step",
    "Test Step",
    "Expected Result",
    "Priority",
    "Test Type",
    "Status",
]

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_CELL_FONT = Font(name="Calibri", size=11)
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)


_EXCEL_MAX_CELL = 32767


def _enum_value(value: object) -> str:
    return value.value if hasattr(value, "value") else str(value or "")


def _excel_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if len(text) > _EXCEL_MAX_CELL:
        return text[: _EXCEL_MAX_CELL - 1]
    return text


def _steps_text(case: TestCase, expected: bool) -> str:
    lines = []
    for step in sorted(case.steps, key=lambda item: item.step_number):
        body = step.expected_result if expected else step.test_step
        lines.append(f"{step.step_number}. {body}")
    return "\n".join(lines)


def _style_header(sheet, headers: list[str], widths: list[int]) -> None:
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(1, index, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = _THIN
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _style_cell(cell, wrap: bool = True) -> None:
    cell.font = _CELL_FONT
    cell.border = _THIN
    cell.alignment = Alignment(wrap_text=wrap, vertical="top")


def load_release_cases(db: Session, release_id: int) -> list[TestCase]:
    stmt = (
        select(TestCase)
        .where(TestCase.release_id == release_id)
        .options(selectinload(TestCase.steps))
        .order_by(TestCase.test_case_id)
    )
    return list(db.execute(stmt).scalars().all())


def build_test_cases_workbook(cases: list[TestCase]) -> bytes:
    workbook = Workbook()
    qc = workbook.active
    qc.title = "Test Cases"
    _style_header(
        qc,
        QC_HEADERS,
        [10, 42, 16, 12, 18, 24, 12, 14, 16, 14, 40, 40, 28, 16, 28, 32, 16, 18, 28, 12, 14, 14],
    )
    for row_index, case in enumerate(cases, start=2):
        values = [
            case.test_case_id,
            _display_case_name(case),
            case.component,
            case.ecosystem or "",
            case.device or "",
            case.device_source or "",
            _enum_value(case.priority),
            _enum_value(case.test_type),
            case.user_type or "",
            _enum_value(case.status),
            _steps_text(case, expected=False),
            _steps_text(case, expected=True),
            case.test_data or "",
            "Sí" if case.requires_condition else "No",
            case.evidence or "",
            case.justification or "",
            case.technical_epic or "",
            case.technical_story or "",
            case.scenario_origin or "",
            case.confidence or "",
            case.complexity or "",
            "" if case.estimation_hours is None else float(case.estimation_hours),
        ]
        for col, value in enumerate(values, start=1):
            cell = qc.cell(row_index, col, _excel_value(value))
            _style_cell(cell)

    zephyr = workbook.create_sheet("Zephyr")
    _style_header(
        zephyr,
        ZEPHYR_HEADERS,
        [14, 18, 42, 28, 16, 8, 40, 40, 12, 14, 14],
    )
    zephyr_row = 2
    for case in cases:
        steps = sorted(case.steps, key=lambda item: item.step_number)
        if not steps:
            steps = [None]
        description_parts = [part for part in (case.description, case.test_data) if part]
        description = "\n".join(description_parts)
        for step in steps:
            values = [
                case.test_case_id,
                case.component,
                _display_case_name(case),
                description,
                case.user_type or "",
                getattr(step, "step_number", 1) if step is not None else 1,
                getattr(step, "test_step", "") if step is not None else "",
                getattr(step, "expected_result", "") if step is not None else "",
                _enum_value(case.priority),
                _enum_value(case.test_type),
                _enum_value(case.status),
            ]
            for col, value in enumerate(values, start=1):
                cell = zephyr.cell(zephyr_row, col, _excel_value(value))
                _style_cell(cell)
            zephyr_row += 1

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
