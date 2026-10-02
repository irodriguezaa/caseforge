"""Excel of RN scope tickets: Technical Epic, NCO, QA Bug, QC Bug, TRI.

Does not change RN counting or CaseForge generation. Keys come from the same
table walk as analyze-rn; Prioridad / Estado / versiones are filled from Jira
when configured.
"""

from __future__ import annotations

import io
import re
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.jira_generation import fetch_scope_fields_for_keys
from app.services.release_note_analyzer import iter_rn_ticket_rows
from app.services.rn_source_type import source_type_for_rn_bucket

SCOPE_HEADERS = [
    "Key",
    "Actividad",
    "Descripción",
    "Prioridad",
    "Estado",
    "Versión afectada",
    "Version correctora",
]

_ACTIVIDAD_LABEL = {
    "functionality": "Technical Epic",
    "nco": "NCO",
    "tri": "TRI",
    "qa_bug": "QA Bug",
    "qc_bug": "QC Bug",
    "qa_qc": "QA/QC Bug",
}

_ACTIVIDAD_ORDER = {
    "Technical Epic": 0,
    "NCO": 1,
    "QA Bug": 2,
    "QC Bug": 3,
    "QA/QC Bug": 4,
    "TRI": 5,
}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_CELL_FONT = Font(name="Calibri", size=11)
_THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_EXCEL_MAX_CELL = 32767


def _excel_value(value: object) -> str:
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    if len(text) > _EXCEL_MAX_CELL:
        return text[: _EXCEL_MAX_CELL - 1]
    return text


def _description_from_cell(ticket_id: str, cell_text: str) -> str:
    text = re.sub(r"[ \t]*\n[ \t]*", " ", cell_text or "").strip()
    text = re.sub(rf"^{re.escape(ticket_id)}\s*[:\-–—]?\s*", "", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def extract_rn_scope_rows(
    pdf_bytes: bytes,
    *,
    jira_fields: dict[str, dict[str, str]] | None = None,
    fetch_jira: bool = True,
) -> list[dict[str, str]]:
    hits = iter_rn_ticket_rows(pdf_bytes)
    fields = dict(jira_fields or {})
    if fetch_jira and hits:
        missing = [ticket_id for ticket_id, _cell, _bucket in hits if ticket_id.upper() not in fields]
        if missing:
            fields.update(fetch_scope_fields_for_keys(missing))

    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for ticket_id, cell_text, bucket in hits:
        key = ticket_id.strip().upper()
        if not key or key in seen:
            continue
        seen.add(key)
        jira = fields.get(key) or {}
        source = source_type_for_rn_bucket(bucket, jira.get("issuetype"))
        descripcion = _description_from_cell(ticket_id, cell_text)
        if not descripcion:
            descripcion = (jira.get("summary") or "").strip()
        if not descripcion:
            descripcion = (jira.get("description") or "").strip()
        rows.append(
            {
                "key": key,
                "actividad": _ACTIVIDAD_LABEL.get(source, "Technical Epic"),
                "descripcion": descripcion,
                "prioridad": (jira.get("priority") or "").strip(),
                "estado": (jira.get("status") or "").strip(),
                "version_afectada": (jira.get("affected_versions") or "").strip(),
                "version_correctora": (jira.get("fix_versions") or "").strip(),
            }
        )
    rows.sort(
        key=lambda row: (
            _ACTIVIDAD_ORDER.get(row["actividad"], 9),
            row["key"],
        )
    )
    return rows


def build_rn_scope_workbook(
    pdf_bytes: bytes,
    *,
    jira_fields: dict[str, dict[str, str]] | None = None,
    fetch_jira: bool = True,
) -> bytes:
    rows = extract_rn_scope_rows(pdf_bytes, jira_fields=jira_fields, fetch_jira=fetch_jira)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Alcance RN"
    for index, header in enumerate(SCOPE_HEADERS, start=1):
        cell = sheet.cell(1, index, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = _THIN
    widths = [16, 18, 70, 14, 18, 22, 22]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(SCOPE_HEADERS))}1"

    for row_index, row in enumerate(rows, start=2):
        values = [
            row["key"],
            row["actividad"],
            row["descripcion"],
            row["prioridad"],
            row["estado"],
            row["version_afectada"],
            row["version_correctora"],
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col, _excel_value(value))
            cell.font = _CELL_FONT
            cell.border = _THIN
            cell.alignment = Alignment(wrap_text=col == 3, vertical="top")
        sheet.row_dimensions[row_index].height = 32

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
