"""QcPulse v2: recolección, etapas y mapeo, con Jira real como fixture y la IA simulada."""

from __future__ import annotations

import json

import pytest

from app.services.qc_v2 import collect, engine, stages
from app.services.qc_v2.collect import EpicBundle, Item, TechStory, UserStory
from app.services.qc_v2.context import device_context
from app.services.qc_v2.llm import LlmMeta, QuotaExhausted

# ------------------------------------------------------------------ fixtures (texto real de Jira)

BRFRE_3529_AC = """1. Cuando el usuario Externo se registra en Claro video e ingresa por primera vez a la home de inicio, cuente o no con alguna suscripción, el Centro de notificaciones debe visualizarse desplegado mostrando los siguientes elementos:

    1. Se debe mostrar la comunicación Descubre más planes
    2. Se debe mostrar el botón accionable que dirija al usuario al Selector de planes respetando el comportamiento actualmente implementado en producción

2. Cuando el usuario Externo inicia sesión en Claro video e ingresa a la home de inicio, cuente o no con alguna suscripción, el Centro de notificaciones debe visualizarse replegado.

    1. Al desplegarse, este debe mostrar los elementos mencionados en el CA 1

3. El Centro de notificaciones debe visualizarse replegado si:

    1. Si el usuario selecciona uno de los botones del Centro de Notificaciones, independientemente si termina o no las acciones que se detonan de los botones
    2. Si el usuario cierra el Centro de Notificaciones antes salir de la home de inicio y posteriormente regresa al mismo

4. Si el usuario no realizó alguna de las interacciones previamente definidas con el Centro de Notificaciones y sale y vuelve a ingresar a la home de inicio, este se debe visualizar desplegado.
5. El Centro de Notificaciones debe ocultarse completamente solo si el usuario ha adquirido todas las suscripciones disponibles en la operación.
6. Todos los textos deben de ser configurables por operación.
7. Todos los elementos deben cumplir con los textos, posición y elementos definidos en los insumos de diseño compartidos.
"""

WEBCL_2061 = """Feature: Ajuste de espaciado vertical en las páginas Home de la UI

  Background:
    Given que tengo acceso a los insumos de diseño aprobados
    And me encuentro en una de las páginas Home definidas para el ajuste

  Scenario: Ajustar distancia entre el menú de navegación principal y los elementos de UI
    When identifico el menú de navegación principal
    Then la distancia vertical debe ser igual al valor definido en Figma

  Scenario: Ajustar espaciado entre elementos intermedios dentro de la Home
    Given que existen elementos intermedios como Brand Header
    Then cada distancia debe respetar los valores establecidos

  Scenario: Validación por resolución
  When se renderiza la Home en las resoluciones soportadas
  Then el espaciado debe mantenerse conforme al layout definido para cada breakpoint
```"""

WEBCL_1876 = """# Motivo: Actualizar el escenario "Identificar contenido bloqueado" al grabar o visualizar
Feature: Actualizar los assets de eventos bloqueados - WEB

    Background: El usuario ingresa a sección "Mis Contenidos"
        Given el usuario está autentificado en la plataforma
        Then se obtiene la lista de contenidos grabados, grabando y por grabar del usuario

    Scenario: Identificar contenido bloqueado
        Given se obtienen las listas de contenidos grabados
        When el "group id" del canal del evento se encuentra en la lista "data.pin_channel.info[]"
        Then la tarjeta del evento contiene el ícono de bloqueado y el overlay en el logo del canal
"""


# ------------------------------------------------------------------ collect


def test_split_criteria_keeps_subitems_inside_top_level():
    criteria = collect.split_criteria(BRFRE_3529_AC)
    assert len(criteria) == 7
    assert "Descubre más planes" in criteria[0]
    assert criteria[2].startswith("El Centro de notificaciones debe visualizarse replegado si")
    assert "cierra el Centro" in criteria[2]


def test_split_criteria_from_adf_ordered_list():
    adf = {
        "type": "doc",
        "content": [
            {"type": "orderedList", "content": [
                {"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Primero"}]}]},
                {"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Segundo"}]}]},
            ]}
        ],
    }
    assert collect.split_criteria(adf) == ["Primero", "Segundo"]


def test_split_criteria_from_adf_split_lists():
    li = lambda text: {"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}]}
    adf = {"type": "doc", "content": [
        {"type": "orderedList", "content": [li("Uno"), li("Dos")]},
        {"type": "paragraph", "content": [{"type": "text", "text": "nota"}]},
        {"type": "orderedList", "attrs": {"order": 3}, "content": [li("Tres")]},
    ]}
    assert collect.split_criteria(adf) == ["Uno", "Dos", "Tres"]


def test_split_criteria_without_numbers_is_one_criterion():
    assert collect.split_criteria("Texto libre sin numeración") == ["Texto libre sin numeración"]


def test_split_gherkin_background_is_context_not_scenario():
    context, blocks = collect.split_gherkin(WEBCL_2061)
    assert "Background" in context
    assert [title for title, _ in blocks] == [
        "Ajustar distancia entre el menú de navegación principal y los elementos de UI",
        "Ajustar espaciado entre elementos intermedios dentro de la Home",
        "Validación por resolución",
    ]
    assert "```" not in blocks[-1][1]


def test_split_gherkin_keeps_quotes_and_comments():
    context, blocks = collect.split_gherkin(WEBCL_1876)
    assert "al grabar o visualizar" in context
    assert len(blocks) == 1
    assert '"data.pin_channel.info[]"' in blocks[0][1]


def test_tech_story_drops_metrics_and_key_scenarios():
    issue = {
        "key": "WEBCL-4541",
        "fields": {
            "summary": "Setup de métricas WEB",
            "issuetype": {"name": "Technical Story"},
            "description": (
                "Feature: WEB\n"
                "  Scenario: Carga el esquema de métricas\n"
                "    Then el esquema de métricas cumple el pipeline\n"
                "  Scenario: El usuario ve el home\n"
                "    Then se muestra el home\n"
                "  Scenario: No se logra obtener una llave\n"
                "    Given una llave no se encuentra en apa/metadata\n"
                "    Then el espacio donde debería mostrarse la leyenda debe mostrar la llave\n"
            ),
        },
    }
    story = collect._tech_story(issue, collect._counter("COV"))
    assert [item.title for item in story.scenarios] == ["El usuario ve el home"]


def test_title_match_technical_brief_vs_product_epic():
    tb = "FE: Onboarding: Actualizar el comportamiento del Centro de Notificaciones de acuerdo al tipo de usuario (Pantallas medianas)"
    pe = "Actualizar el comportamiento del Centro de Notificaciones de acuerdo al tipo de usuario"
    spotlight = "Implementar estado Spotlight en el Centro de Notificaciones"
    assert collect.title_similarity(tb, pe) >= collect.TITLE_MATCH_THRESHOLD
    assert collect.title_similarity(tb, spotlight) < collect.TITLE_MATCH_THRESHOLD


class _Resp:
    def __init__(self, status: int, payload: dict | None = None):
        self.status_code = status
        self._payload = payload or {}

    def json(self):
        return self._payload

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(self.status_code)


def _issue(key, issuetype, summary="", parent=None, **fields):
    data = {"summary": summary, "issuetype": {"name": issuetype}, **fields}
    if parent:
        data["parent"] = {"key": parent}
    return {"key": key, "fields": data}


class _FakeJira:
    def __init__(self):
        self.issues = {
            "WEBCL-1586": _issue("WEBCL-1586", "Technical Epic", "Onboarding CN", parent="TBRFRE-895"),
            "TBRFRE-895": _issue(
                "TBRFRE-895", "Technical Brief",
                "FE: Onboarding: Actualizar el comportamiento del Centro de Notificaciones de acuerdo al tipo de usuario (Pantallas medianas)",
                parent="BRFRE-3492",
            ),
            "BRFRE-3492": _issue("BRFRE-3492", "Product Brief", "MX | Onboarding"),
        }
        self.children_map = {
            "WEBCL-1586": [
                _issue("WEBCL-1587", "Technical Story", "TS externo", description=WEBCL_2061),
                _issue("WEBCL-2296", "QA Bug", "bug", description="Scenario: no debe contar"),
                _issue("WEBCL-2677", "NCO System Demo", "nco", description="Scenario: tampoco"),
            ],
            "BRFRE-3492": [
                _issue("BRFRE-3510", "Product Epic", "Actualizar el comportamiento del Centro de Notificaciones de acuerdo al tipo de usuario"),
                _issue("BRFRE-3382", "Product Epic", "Implementar estado Spotlight en el Centro de Notificaciones"),
                _issue("TBRFRE-895", "Technical Brief", "FE: ..."),
            ],
            "BRFRE-3510": [
                _issue("BRFRE-3529", "User Story", "CN usuarios Externos", **{
                    collect.F_AC: BRFRE_3529_AC,
                    collect.F_YO_COMO: "negocio de Claro video",
                    collect.F_TIPO_USUARIO: [{"value": "Registrado"}, {"value": "Suscripto"}],
                }),
            ],
            "BRFRE-3382": [_issue("BRFRE-3502", "User Story", "Spotlight", **{collect.F_AC: "1. Spotlight"})],
        }

    def get(self, url, params=None):
        key = url.rsplit("/", 1)[-1]
        issue = self.issues.get(key)
        return _Resp(200, issue) if issue else _Resp(404)

    def post(self, url, json=None):
        key = json["jql"].split('"')[1]
        return _Resp(200, {"issues": self.children_map.get(key, [])})


def test_collect_epic_follows_business_branch_and_filters_by_title():
    bundle = collect.collect_epic(collect.JiraReader(_FakeJira()), "WEBCL-1586")
    assert [s.key for s in bundle.user_stories] == ["BRFRE-3529"]  # Spotlight fuera
    assert len(bundle.crns) == 7
    assert [c.id for c in bundle.covs] == ["COV-001", "COV-002", "COV-003"]
    assert bundle.user_stories[0].tipo_usuario == "Registrado, Suscripto"
    assert bundle.product_brief_key == "BRFRE-3492"
    assert not [w for w in bundle.warnings if "Sin User Story" in w]
    assert any("1 NCO System Demo, 1 QA Bug" in w for w in bundle.warnings)


def test_story_type_filter():
    assert collect.is_story_type("Technical Story")
    assert collect.is_story_type("Historia técnica")
    assert not collect.is_story_type("QA Bug")
    assert not collect.is_story_type("QC Bug")
    assert not collect.is_story_type("NCO System Demo")
    assert not collect.is_story_type("Subtarea")


# ------------------------------------------------------------------ stages + engine


def _bundle() -> EpicBundle:
    return EpicBundle(
        key="WEBCL-2060",
        summary="Espaciado vertical",
        technical_stories=[
            TechStory("WEBCL-2061", "Ajuste", "Background ...", [
                Item("COV-001", "Distancia menú", "Scenario: Distancia menú ...", "WEBCL-2061"),
                Item("COV-002", "Por resolución", "Scenario: Por resolución ...", "WEBCL-2061"),
                Item("COV-003", "Telemetría", "Scenario: se envía evento", "WEBCL-2061"),
            ])
        ],
        user_stories=[UserStory("BRFRE-1", "Espaciado", criteria=[Item("CRN-001", "Espaciado", "1. Espaciado", "BRFRE-1")])],
    )


_STAGE_A = {
    "resumen_funcionalidad": "El espaciado entre menú y contenido cambia",
    "escenarios": [
        {"id": "ESC-01", "titulo": "Espaciado entre menú y primer componente", "descripcion": "d",
         "condiciones": [], "cubre_negocio": ["CRN-001"], "cubre_tecnico": ["COV-001", "COV-002"],
         "brecha": "ninguna", "es_no_afectacion": False},
        {"id": "ESC-02", "titulo": "No afectación del menú", "descripcion": "d", "condiciones": [],
         "cubre_negocio": [], "cubre_tecnico": [], "brecha": "implicito", "es_no_afectacion": True},
    ],
    "excluidos": [{"id": "COV-003", "motivo": "telemetria", "explicacion": "solo evento"}],
}

_CASE = {
    "titulo": "Visualización del espaciado entre menú y primer componente",
    "tipo": "ERG",
    "prioridad": "Critical",
    "precondiciones": ["Pantalla mediana", "Diseño aprobado"],
    "pasos": [
        {"accion": "Seleccionar la Home", "resultado_esperado": "Se muestra la Home"},
        {"accion": "Observar la separación", "resultado_esperado": "Igual al diseño"},
        {"accion": "Recorrer las demás Homes", "resultado_esperado": "Igual al diseño en todas"},
    ],
    "variantes": [],
    "referencia_tecnica": "Figma",
    "riesgo": "espaciado incorrecto",
    "confianza": "alta",
}


def _fake_chat(responses):
    calls = []

    def chat(messages, *, model, temperature=None, **kwargs):
        calls.append(messages)
        system = messages[0]["content"]
        payload = responses["a"] if "ESCENARIOS DE PRUEBA" in system else responses["b"]
        if callable(payload):
            payload = payload(messages)
        return payload, LlmMeta(model=model, prompt_tokens=10, completion_tokens=5)

    chat.calls = calls
    return chat


def test_stage_a_parses_and_detects_missing():
    result = stages.parse_stage_a({**_STAGE_A, "excluidos": []}, _bundle())
    assert result.missing == ["COV-003"]
    assert result.scenarios[0].cubre_tecnico == ["COV-001", "COV-002"]


def test_stage_a_retries_once_when_missing():
    first = {**_STAGE_A, "excluidos": []}
    answers = iter([first, _STAGE_A])
    chat = _fake_chat({"a": lambda m: next(answers), "b": {"casos": [_CASE]}})
    result = stages.run_stage_a(_bundle(), device_context("WEB"), chat=chat)
    assert result.missing == []
    assert len(chat.calls) == 2


def test_engine_end_to_end_with_fake_llm():
    chat = _fake_chat({"a": _STAGE_A, "b": {"casos": [_CASE]}})
    run = engine.run_epic(_bundle(), device_context("WEB"), chat=chat)
    assert run.error is None
    assert len(run.candidates) == 2  # un caso por escenario en este fake
    first = run.candidates[0]
    assert first.generation_origin == "qc-v2"
    assert first.covers == ["CRN-001", "COV-001", "COV-002"]
    assert "qc-v2:tipo=ERG" in first.applied_rules
    assert "qc-v2:complejidad=BAJA" in first.applied_rules
    assert first.review_required is False
    estados = {row["scenario"].split()[0]: row["estado"] for row in run.ledger}
    assert estados["CRN-001"] == "cruce"
    assert estados["COV-003"] == "excluido: telemetria"
    assert any(row["estado"] == "no afectación" for row in run.ledger)
    assert run.llm_calls == 3


def test_case_with_two_steps_is_kept_and_flagged():
    short = {**_CASE, "pasos": _CASE["pasos"][:2]}
    chat = _fake_chat({"a": _STAGE_A, "b": {"casos": [short]}})
    run = engine.run_epic(_bundle(), device_context("WEB"), chat=chat)
    assert run.candidates, "nunca se descarta un caso"
    assert run.candidates[0].review_required is True
    assert "2 pasos" in (run.candidates[0].applicability_reason or "")


def test_empty_quotes_flagged():
    bad = {**_CASE, "pasos": [{"accion": 'Seleccionar "< >"', "resultado_esperado": "x"}] + _CASE["pasos"][1:]}
    case = stages.parse_stage_b({"casos": [bad]}, "ESC-01")[0]
    assert "texto entre comillas vacío" in stages.validate_case(case)


def test_complexity_rules():
    case = stages.parse_stage_b({"casos": [_CASE]}, "ESC-01")[0]
    assert stages.complexity(case) == "BAJA"
    case.tipo = "E2E"
    assert stages.complexity(case) == "ALTA"
    case.tipo = "HPA"
    case.variantes = [{"condicion": "a", "resultado_esperado": "b"}]
    assert stages.complexity(case) == "MEDIA"


def test_quota_stops_the_run():
    def chat(messages, *, model, temperature=None, **kwargs):
        raise QuotaExhausted("sin créditos")

    runs = engine.run_epics(["WEBCL-1586"], "WEB", reader=collect.JiraReader(_FakeJira()), chat=chat)
    assert runs[0].error and runs[0].error.startswith("corrida detenida")


def test_device_context_groups():
    assert device_context("ADT").modo_interaccion == "control remoto"
    assert device_context("WEB").grupo_pantalla == "Pantalla mediana"
    assert device_context("WEB").modo_interaccion == "cursor (mouse)"
    assert device_context("ADR").modo_interaccion == "táctil"
    assert device_context("WIN").metodos_pago == "No transacciona"


def test_prompts_exist_and_are_loaded():
    assert "ESCENARIOS DE PRUEBA" in stages.load_prompt("etapa_a.md")
    assert "REGLA DEL BUG" in stages.load_prompt("etapa_b.md")


# ------------------------------------------------------------------ export


def test_export_adds_v2_columns_only_for_v2_cases():
    from types import SimpleNamespace

    from openpyxl import Workbook

    from app.services.test_case_export import _append_v2_columns, _v2_fields

    v2 = SimpleNamespace(
        justification="QcPulse v2 · ESC-03 · tipo=TTF · brecha=sin_tecnico · complejidad=MEDIA covers=CRN-001",
        applicability_reason="requiere revisión: 2 pasos (esperado 3 a 6)",
    )
    legacy = SimpleNamespace(justification="Materialización determinista", applicability_reason=None)
    assert _v2_fields(legacy) is None
    assert _v2_fields(v2) == ["ESC-03", "TTF", "Solo negocio (sin Technical Story)", "2 pasos (esperado 3 a 6)"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ID", "Nombre"])
    sheet.append(["QC-001", "x"])
    _append_v2_columns(sheet, [legacy])
    assert sheet.max_column == 2  # sin casos v2 la hoja no cambia
    _append_v2_columns(sheet, [v2])
    assert [sheet.cell(1, c).value for c in range(3, 7)] == ["Escenario", "Tipo de prueba", "Brecha", "Requiere revisión"]
    assert sheet.cell(2, 4).value == "TTF"


# ------------------------------------------------------------------ benchmark


def _luminia_book(path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz de CP QC"
    ws.append(["Technical Epic", "Technical Story", "External id", "TC_Name", "#", "Step", "Test Step Description", "Expected Result"])
    for n in range(1, 11):
        ws.append(["WEBCL-1", "WEBCL-2", f"CP{n:03d}", f"Caso {n}", n, "Step 1", "Abrir", "Se abre"])
        ws.append([None, None, None, None, None, "Step 2", "Ver", "Se ve"])
    wb.save(path)


def _qcpulse_book(path, cases):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Casos"
    ws.append(["Épica", "Escenario", "Título", "Tipo", "Prioridad", "Complejidad", "Precondiciones", "Pasos", "Resultado esperado"])
    for title, steps in cases:
        ws.append(["WEBCL-1", "ESC-01", title, "HPP", "CRITICAL", "BAJA", "", "\n".join(f"{i}. p" for i in range(1, steps + 1)), "\n".join(f"{i}. r" for i in range(1, steps + 1))])
    wb.save(path)


def test_benchmark_deterministic_metrics(tmp_path):
    from app.services.qc_v2 import benchmark

    _luminia_book(tmp_path / "l.xlsx")
    _qcpulse_book(tmp_path / "q.xlsx", [("A", 3), ("B", 4), ("C", 2), ("D", 5), ("E", 6)])
    lum = benchmark.read_luminia(tmp_path / "l.xlsx")
    qcp = benchmark.read_qcpulse(tmp_path / "q.xlsx")
    assert len(lum) == 10 and lum[0].steps == ["Abrir", "Ver"]
    result = benchmark.score(lum, qcp)[0]
    assert result.ratio == 0.5 and result.checks()["granularidad"] is True
    assert result.steps_ok == 0.8 and result.checks()["pasos"] is False
    benchmark.write_report([result], tmp_path / "r.xlsx")


def test_benchmark_judge_parsing(tmp_path):
    from app.services.qc_v2 import benchmark

    _luminia_book(tmp_path / "l.xlsx")
    _qcpulse_book(tmp_path / "q.xlsx", [("A", 3), ("B", 4)])
    verdict = {
        "cobertura": [{"luminia_id": f"CP{n:03d}", "estado": "si" if n <= 9 else "no", "qcpulse_ids": ["Q001"]} for n in range(1, 11)],
        "pasos_desarrollo": [],
        "duplicados": [],
    }

    def chat(messages, *, model, temperature=None, **kwargs):
        return verdict, LlmMeta(model=model)

    result = benchmark.score(
        benchmark.read_luminia(tmp_path / "l.xlsx"), benchmark.read_qcpulse(tmp_path / "q.xlsx"), judge=True, chat=chat
    )[0]
    assert result.cobertura == 0.9 and result.checks()["cobertura"] is True
    assert result.solo_qcpulse == ["Q002 — B"]


def test_cli_workbook_end_to_end(tmp_path):
    from openpyxl import load_workbook

    from app.services.qc_v2 import cli

    stage_a = {
        "resumen_funcionalidad": "r",
        "escenarios": [
            {"id": "ESC-01", "titulo": "Centro desplegado", "descripcion": "d", "condiciones": [],
             "cubre_negocio": [f"CRN-{n:03d}" for n in range(1, 8)], "cubre_tecnico": ["COV-001", "COV-002"],
             "brecha": "ninguna", "es_no_afectacion": False},
            {"id": "ESC-02", "titulo": "Resolución", "descripcion": "d", "condiciones": [],
             "cubre_negocio": [], "cubre_tecnico": ["COV-003"], "brecha": "sin_negocio", "es_no_afectacion": False},
        ],
        "excluidos": [],
    }
    chat = _fake_chat({"a": stage_a, "b": {"casos": [_CASE]}})
    runs = engine.run_epics(["WEBCL-1586"], "WEB", reader=collect.JiraReader(_FakeJira()), chat=chat)
    path = tmp_path / "out.xlsx"
    cli.build_workbook(runs).save(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Resumen", "Escenarios", "Casos", "Cruce"]
    assert wb["Casos"].max_row == 3
    estados = [row[3] for row in wb["Cruce"].iter_rows(min_row=2, values_only=True)]
    assert "solo técnico" in estados and "cruce" in estados
    json_ready = cli._jsonable(runs[0])
    assert json_ready["epic"] == "WEBCL-1586"



# ------------------------------------------------------------------ cliente OpenAI


def _transport(responses, seen):
    import httpx

    queue = list(responses)

    def handler(request):
        seen.append(json.loads(request.content))
        status, body = queue.pop(0)
        return httpx.Response(status, json=body)

    return httpx.MockTransport(handler)


@pytest.fixture()
def fake_key():
    from app.config import settings

    previous = settings.openai_api_key
    object.__setattr__(settings, "openai_api_key", "k")
    yield
    object.__setattr__(settings, "openai_api_key", previous)


def test_llm_drops_unsupported_params_and_parses(fake_key):
    from app.services.qc_v2 import llm

    ok = {"choices": [{"message": {"content": json.dumps({"casos": []})}}], "usage": {"prompt_tokens": 3, "completion_tokens": 2}}
    seen: list[dict] = []
    transport = _transport(
        [
            (400, {"error": {"message": "Unsupported parameter: 'temperature'"}}),
            (400, {"error": {"message": "Unrecognized request argument: reasoning_effort"}}),
            (200, ok),
        ],
        seen,
    )
    payload, meta = llm.chat_json([{"role": "user", "content": "json"}], model="m", temperature=0.2, reasoning_effort="low", transport=transport)
    assert payload == {"casos": []}
    assert "temperature" in seen[0] and "temperature" not in seen[1]
    assert "reasoning_effort" in seen[1] and "reasoning_effort" not in seen[2]
    assert meta.prompt_tokens == 3


def test_llm_quota_raises(fake_key):
    from app.services.qc_v2 import llm

    transport = _transport([(429, {"error": {"code": "insufficient_quota", "message": "You have no credits remaining"}})], [])
    with pytest.raises(QuotaExhausted):
        llm.chat_json([{"role": "user", "content": "json"}], model="m", transport=transport)
