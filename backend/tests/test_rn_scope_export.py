from io import BytesIO

from openpyxl import load_workbook

from app.services.rn_scope_export import SCOPE_HEADERS, build_rn_scope_workbook, extract_rn_scope_rows
from tests.test_case_generation import WEB_RN, _create_app_release_with_rn, _store_dir

_WEB_FEATURES = {"WEBCL-3721", "WEBCL-3153", "WEBCL-3779", "WEBCL-3762", "WEBCL-3767"}


def test_extract_web_rn_scope_without_jira() -> None:
    rows = extract_rn_scope_rows(WEB_RN.read_bytes(), fetch_jira=False)
    keys = {row["key"] for row in rows}
    assert _WEB_FEATURES <= keys
    epics = [row for row in rows if row["actividad"] == "Technical Epic"]
    assert {row["key"] for row in epics} >= _WEB_FEATURES
    assert all(row["descripcion"] for row in epics)
    assert {row["actividad"] for row in rows} <= {
        "Technical Epic",
        "NCO",
        "QA Bug",
        "QC Bug",
        "QA/QC Bug",
        "TRI",
    }


def test_extract_splits_qa_qc_from_jira_issuetype(monkeypatch) -> None:
    monkeypatch.setattr(
        "app.services.rn_scope_export.iter_rn_ticket_rows",
        lambda _pdf: [
            ("WEBCL-3849", "WEBCL-3849: login fallido", "qa_qc"),
            ("WEBCL-3900", "WEBCL-3900: crash al pagar", "qa_qc"),
            ("WEBCL-3721", "WEBCL-3721: TE-activacion", "functionality"),
        ],
    )
    fields = {
        "WEBCL-3849": {
            "issuetype": "QA Bug",
            "priority": "High",
            "status": "Done",
            "affected_versions": "16.9.0",
            "fix_versions": "16.9.1",
        },
        "WEBCL-3900": {
            "issuetype": "QC Bug",
            "priority": "Highest",
            "status": "In Progress",
            "affected_versions": "",
            "fix_versions": "16.9.1",
        },
    }
    by_key = {
        row["key"]: row
        for row in extract_rn_scope_rows(b"pdf", jira_fields=fields, fetch_jira=False)
    }
    assert by_key["WEBCL-3849"]["actividad"] == "QA Bug"
    assert by_key["WEBCL-3849"]["prioridad"] == "High"
    assert by_key["WEBCL-3849"]["estado"] == "Done"
    assert by_key["WEBCL-3849"]["version_afectada"] == "16.9.0"
    assert by_key["WEBCL-3849"]["version_correctora"] == "16.9.1"
    assert by_key["WEBCL-3900"]["actividad"] == "QC Bug"
    assert by_key["WEBCL-3721"]["actividad"] == "Technical Epic"


def test_extract_jira_nco_overrides_functionality_bucket(monkeypatch) -> None:
    """WEBCL keys are shared by Technical Epic and NCO; Jira issuetype wins."""
    monkeypatch.setattr(
        "app.services.rn_scope_export.iter_rn_ticket_rows",
        lambda _pdf, _filename="": [
            ("WEBCL-450", "WEBCL-450: ajuste CMS", "functionality"),
            ("WEBCL-1900", "WEBCL-1900: banner cookies", "functionality"),
        ],
    )
    fields = {
        "WEBCL-450": {
            "issuetype": "NCO",
            "priority": "",
            "status": "",
            "affected_versions": "",
            "fix_versions": "",
        },
        "WEBCL-1900": {
            "issuetype": "Technical Epic",
            "priority": "",
            "status": "",
            "affected_versions": "",
            "fix_versions": "",
        },
    }
    by_key = {
        row["key"]: row
        for row in extract_rn_scope_rows(b"pdf", jira_fields=fields, fetch_jira=False)
    }
    assert by_key["WEBCL-450"]["actividad"] == "NCO"
    assert by_key["WEBCL-1900"]["actividad"] == "Technical Epic"


def test_workbook_headers_and_rows() -> None:
    payload = build_rn_scope_workbook(WEB_RN.read_bytes(), fetch_jira=False)
    book = load_workbook(BytesIO(payload))
    sheet = book.active
    assert [cell.value for cell in sheet[1]] == SCOPE_HEADERS
    assert sheet.max_row > 1
    assert sheet["A2"].value


def test_export_rn_scope_from_release(client, monkeypatch, tmp_path) -> None:
    monkeypatch.setattr("app.services.rn_scope_export.fetch_scope_fields_for_keys", lambda _keys: {})

    def boom(*_args, **_kwargs):
        raise AssertionError("export must not re-read the RN PDF")

    release_id, _analysis = _create_app_release_with_rn(client, monkeypatch, tmp_path)
    monkeypatch.setattr("app.services.release_note_analyzer.iter_rn_ticket_rows", boom)
    monkeypatch.setattr("app.services.tata_release_note_parser.parse_tata_release_note", boom)
    monkeypatch.setattr("app.routers.releases.read_release_note_pdf", boom)
    response = client.get(f"/api/v1/releases/{release_id}/rn-scope/export")
    assert response.status_code == 200, response.text
    assert "spreadsheet" in response.headers["content-type"]
    book = load_workbook(BytesIO(response.content))
    keys = {row[0] for row in book.active.iter_rows(min_row=2, values_only=True)}
    assert _WEB_FEATURES <= keys


def test_export_rn_scope_from_analyzed_pdf(client, monkeypatch, tmp_path) -> None:
    monkeypatch.setattr("app.services.rn_scope_export.fetch_scope_fields_for_keys", lambda _keys: {})

    def boom(*_args, **_kwargs):
        raise AssertionError("export must not re-read the RN PDF")

    _store_dir(monkeypatch, tmp_path)
    analyzed = client.post(
        "/api/v1/releases/analyze-rn",
        files={"file": (WEB_RN.name, WEB_RN.read_bytes(), "application/pdf")},
    )
    assert analyzed.status_code == 200
    analysis = analyzed.json()["analysis"]
    monkeypatch.setattr("app.services.release_note_analyzer.iter_rn_ticket_rows", boom)
    monkeypatch.setattr("app.services.tata_release_note_parser.parse_tata_release_note", boom)
    monkeypatch.setattr("app.routers.releases.read_release_note_pdf", boom)
    response = client.post(
        "/api/v1/releases/rn-scope/export",
        json={"filename": WEB_RN.name, "raw_analysis": analysis["raw_analysis"]},
    )
    assert response.status_code == 200, response.text
    assert "spreadsheet" in response.headers["content-type"]


def test_fetch_scope_fields_timeout_returns_empty(monkeypatch) -> None:
    import httpx
    from app.services.jira_generation import fetch_scope_fields_for_keys

    class Boom:
        def __enter__(self):
            raise httpx.TimeoutException("slow")

        def __exit__(self, *_args):
            return False

    monkeypatch.setattr("app.services.jira_generation._client", lambda timeout=None: Boom())
    assert fetch_scope_fields_for_keys(["WEBCL-3721"]) == {}


def test_export_rn_scope_requires_stored_analysis(client, monkeypatch, tmp_path) -> None:
    _store_dir(monkeypatch, tmp_path)
    response = client.post(
        "/api/v1/releases/rn-scope/export",
        json={"pdf_file_path": str(WEB_RN)},
    )
    assert response.status_code == 422


def _assert_stored_matches_pdf(pdf_path) -> None:
    from app.services.ai_case_engine import _tickets_by_section
    from app.services.epc_paste import tickets_from_normalized_technical_epics
    from app.services.release_note_analyzer import RuleBasedPdfAnalyzer
    from app.services.rn_scope_export import extract_rn_scope_rows_from_analysis

    pdf = pdf_path.read_bytes()
    analysis = RuleBasedPdfAnalyzer().analyze(pdf_path.name, pdf)
    normalized = analysis.raw_analysis["normalized"]
    assert isinstance(normalized.get("technical_epics"), list)
    assert isinstance(normalized.get("nco"), list)
    assert isinstance(normalized.get("tri"), list)
    assert isinstance(normalized.get("qa_qc"), list)
    assert isinstance(normalized.get("qa_bugs"), list)
    assert isinstance(normalized.get("qc_bugs"), list)
    assert isinstance(normalized.get("incidents"), list)
    stored = tickets_from_normalized_technical_epics(analysis.raw_analysis)
    walked = _tickets_by_section(pdf, pdf_path.name)
    for bucket in ("functionality", "nco", "tri", "qa_qc"):
        assert [(key.upper(), text) for key, text in stored[bucket]] == [
            (key.upper(), text) for key, text in walked[bucket]
        ]
    assert extract_rn_scope_rows_from_analysis(analysis.raw_analysis, fetch_jira=False) == (
        extract_rn_scope_rows(pdf, fetch_jira=False, filename=pdf_path.name)
    )
    if analysis.raw_analysis.get("vendor") == "tata":
        for item in normalized["incidents"]:
            assert item.get("source_category")
            if item.get("source_category") == "INCIDENT":
                assert item.get("normalized_category") == "TRI"
                assert item["id"] in {row["id"] for row in normalized["tri"]}
        for bucket in ("known_issues", "qa_evidence", "qa_bugs", "qc_bugs", "qco"):
            assert all(item.get("normalized_category") != "TRI" for item in normalized[bucket])
    else:
        for bucket in ("technical_epics", "nco", "tri", "qa_qc"):
            for item in normalized[bucket]:
                assert item.get("source_category", "") == ""
                assert item.get("normalized_category", "") == ""


def test_stored_analysis_matches_pdf_scope_without_second_walk(monkeypatch) -> None:
    import app.services.release_note_analyzer as analyzer_mod

    pdf = WEB_RN.read_bytes()
    walks = {"n": 0}
    real_walk = analyzer_mod._iter_damco_ticket_rows

    def counting_walk(pdf_bytes: bytes):
        walks["n"] += 1
        return real_walk(pdf_bytes)

    monkeypatch.setattr(analyzer_mod, "_iter_damco_ticket_rows", counting_walk)
    from app.services.release_note_analyzer import RuleBasedPdfAnalyzer

    RuleBasedPdfAnalyzer().analyze(WEB_RN.name, pdf)
    assert walks["n"] == 1
    _assert_stored_matches_pdf(WEB_RN)
    tata = WEB_RN.parent / "tata_launcher_12.0.0.pdf"
    _assert_stored_matches_pdf(tata)


def test_keep_technical_epic_keys_drops_nco_webcl() -> None:
    from app.services.jira_generation import keep_technical_epic_keys

    kept, dropped = keep_technical_epic_keys(
        ["WEBCL-1900", "WEBCL-450", "WEBCL-451", "WEBCL-2103", "WEBCL-2116"],
        {
            "WEBCL-1900": "Technical Epic",
            "WEBCL-450": "NCO",
            "WEBCL-451": "NCO",
            "WEBCL-2103": "NCO",
            "WEBCL-2116": "NCO",
        },
    )
    assert kept == ["WEBCL-1900"]
    assert dropped == ["WEBCL-450", "WEBCL-451", "WEBCL-2103", "WEBCL-2116"]
    open_kept, open_dropped = keep_technical_epic_keys(["WEBCL-450"], {})
    assert open_kept == ["WEBCL-450"] and open_dropped == []
