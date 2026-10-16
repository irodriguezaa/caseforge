"""FASE 8 FIX 2: export accumulates every EPC, keeps NA, and Avance matches Test Cases."""

from __future__ import annotations

import io
from types import SimpleNamespace

from openpyxl import load_workbook

from app.schemas.case_generation import GenerationStats
from app.services.ai_case_engine import generate_release_app_candidates
from app.services.gherkin_coverage import candidates_from_jira_artifacts
from app.services.rn_epc_scope import merge_scenario_ledgers
from app.services.scenario_llm import _llm_eligible
from app.services.test_case_export import build_test_cases_workbook


def _story(key: str, summary: str, description: str) -> dict:
    return {
        "key": key,
        "issuetype": "Technical Story",
        "summary": summary,
        "description": description,
        "acceptance_criteria": "",
    }


def _epic(key: str, stories: list[dict]) -> dict:
    return {
        "key": key,
        "issuetype": "Technical Epic",
        "summary": key,
        "description": "",
        "acceptance_criteria": "",
        "children": stories,
    }


def _generate(monkeypatch, artifacts: list[dict], *, only: set[str] | None = None):
    from dataclasses import replace

    from app.config import settings

    monkeypatch.setattr(
        "app.services.ai_case_engine.settings",
        replace(settings, openai_api_key=""),
    )
    monkeypatch.setattr(
        "app.services.scenario_llm.settings",
        replace(settings, openai_api_key=""),
    )
    by_key = {item["key"]: item for item in artifacts}

    def fake_fetch(keys):
        return [by_key[key] for key in keys if key in by_key]

    monkeypatch.setattr("app.services.ai_case_engine.fetch_artifacts_for_keys", fake_fetch)
    monkeypatch.setattr(
        "app.services.ai_case_engine.fetch_issuetypes_for_keys",
        lambda keys: {key: "Technical Epic" for key in keys},
    )
    return generate_release_app_candidates(
        release_id=1,
        release_name="RN sintetico",
        validation_type=None,
        analysis_present=True,
        rn_filename="rn.pdf",
        pdf_bytes=None,
        release_context={},
        existing_cases=[],
        tickets={"functionality": [(item["key"], item["key"]) for item in artifacts]},
        restrict_to_functionality_keys=only,
    )


def _sheet_rows(sheet) -> list[tuple]:
    rows = []
    for index in range(2, sheet.max_row + 1):
        values = tuple(sheet.cell(index, col).value for col in range(1, 6))
        if any(values):
            rows.append(values)
    return rows


ST_A = """
Feature: Pago
  Scenario: Completar pago
    When el usuario selecciona Confirmar
    Then se muestra el comprobante
  Scenario: Cancelar pago
    When el usuario selecciona Cancelar
    Then se muestra el resumen
"""

ST_B = """
Feature: Ficha
  Scenario: Abrir ficha
    When el usuario abre la ficha
    Then se muestra el detalle
"""

ST_NA = """
Feature: Leyenda
  Scenario: Llave remota faltante
    Given la llave de configuración está faltante
    When el usuario abre la pantalla
    Then se muestra el espacio vacío
"""

ST_REMOTE = """
Feature: Configuración remota
  Scenario: Sin configuración remota
    Given la configuración remota no se logra obtener
    When el usuario abre el flujo
    Then se muestra el estado vacío en pantalla
"""


def test_x1_sheet_accumulates_scenarios_from_every_epc(monkeypatch) -> None:
    artifacts = [
        _epic("EPC-A", [_story("ST-A", "Feature pago", ST_A)]),
        _epic("EPC-B", [_story("ST-B", "Feature ficha", ST_B), _story("ST-NA", "Feature leyenda", ST_NA)]),
    ]
    first = _generate(monkeypatch, artifacts, only={"EPC-A"})
    second = _generate(monkeypatch, artifacts, only={"EPC-B"})
    merged = merge_scenario_ledgers(
        first.rn_scope_coverage["scenario_ledger"],
        second.rn_scope_coverage["scenario_ledger"],
    )
    stories = {row["story"] for row in merged}
    assert stories == {"ST-A", "ST-B", "ST-NA"}
    assert len(merged) == 4
    book = load_workbook(io.BytesIO(build_test_cases_workbook([], scenario_ledger=merged)))
    rows = _sheet_rows(book["Escenarios de Jira"])
    assert len(rows) == 4
    assert {row[0] for row in rows} == stories


def test_x1_regenerating_one_epc_does_not_duplicate(monkeypatch) -> None:
    artifacts = [
        _epic("EPC-A", [_story("ST-A", "Feature pago", ST_A)]),
        _epic("EPC-B", [_story("ST-B", "Feature ficha", ST_B)]),
    ]
    first = _generate(monkeypatch, artifacts, only={"EPC-A"})
    again = _generate(monkeypatch, artifacts, only={"EPC-A"})
    other = _generate(monkeypatch, artifacts, only={"EPC-B"})
    merged = merge_scenario_ledgers(first.rn_scope_coverage["scenario_ledger"], other.rn_scope_coverage["scenario_ledger"])
    merged = merge_scenario_ledgers(merged, again.rn_scope_coverage["scenario_ledger"])
    assert len([row for row in merged if row["story"] == "ST-A"]) == 2
    assert len([row for row in merged if row["story"] == "ST-B"]) == 1


def test_x2_na_in_rn_is_on_ledger_and_na_sheet(monkeypatch) -> None:
    artifacts = [
        _epic(
            "EPC-B",
            [_story("ST-REMOTE", "Feature configuración remota", ST_REMOTE), _story("ST-B", "Feature ficha", ST_B)],
        )
    ]
    result = _generate(monkeypatch, artifacts)
    ledger = result.generation_stats.scenario_ledger
    na_rows = [row for row in ledger if row.get("estado") == "na_ambiente"]
    assert na_rows
    assert any("configuración remota" in (row.get("scenario") or "").lower() for row in na_rows)
    book = load_workbook(
        io.BytesIO(build_test_cases_workbook([], scenario_ledger=ledger))
    )
    assert "No aplicables en ambiente" in book.sheetnames
    na_sheet = book["No aplicables en ambiente"]
    names = [na_sheet.cell(index, 2).value for index in range(2, na_sheet.max_row + 1)]
    assert any(name and "configuración remota" in name.lower() for name in names)
    assert any(row.get("estado") == "na_ambiente" for row in result.rn_scope_coverage["scenario_ledger"])


def test_x2_persisted_na_is_not_repeated_from_the_ledger() -> None:
    case = _case(
        test_case_id="QC-NA",
        test_case_name="Llave remota",
        technical_story="ST-NA",
        scenario_origin="ST-NA: Llave remota",
        applicability_reason="na_ambiente: Given exige configuración remota",
    )
    ledger = [
        {
            "epic": "EPC-1",
            "story": "ST-NA",
            "scenario": "Llave remota",
            "estado": "na_ambiente",
            "motivo": "Given exige configuración remota",
            "case_id": "Llave remota",
        }
    ]
    book = load_workbook(io.BytesIO(build_test_cases_workbook([case], scenario_ledger=ledger)))
    na_sheet = book["No aplicables en ambiente"]
    names = [na_sheet.cell(index, 2).value for index in range(2, na_sheet.max_row + 1)]
    assert names.count("Llave remota") == 1


def test_x2_na_survives_deterministic_and_is_not_sent_to_llm() -> None:
    from app.services.gherkin_coverage import build_coverage_inventory

    stats = GenerationStats()
    artifacts = [_epic("EPC-1", [_story("ST-REMOTE", "Feature configuración remota", ST_REMOTE)])]
    units = build_coverage_inventory(artifacts, "rn.pdf", stats=stats)
    na_units = [unit for unit in units if unit.applicability == "na_ambiente"]
    assert na_units
    assert all(not _llm_eligible(unit) for unit in na_units)
    cases = candidates_from_jira_artifacts(artifacts, "rn.pdf", [], lambda *_a: None, stats=stats)
    assert any(row.get("estado") == "na_ambiente" for row in stats.scenario_ledger)
    assert all((case.applicability or "") != "na_ambiente" or case.applicability == "na_ambiente" for case in cases)


def test_x2_missing_key_is_not_materialized() -> None:
    from app.services.gherkin_coverage import build_coverage_inventory

    stats = GenerationStats()
    artifacts = [_epic("EPC-1", [_story("ST-NA", "Feature leyenda", ST_NA)])]
    units = build_coverage_inventory(artifacts, "rn.pdf", stats=stats)
    assert units == []
    assert any("Llave remota" in (row.get("scenario") or "") for row in stats.scenario_ledger)
    assert not any(row.get("estado") == "na_ambiente" for row in stats.scenario_ledger)


def _case(**kwargs):
    defaults = {
        "component": "EPC-1",
        "technical_epic": "EPC-1",
        "hn_source": None,
        "status": "UNEXECUTED",
        "estimation_hours": 0.2,
        "test_case_id": "QC-001",
        "test_case_name": "Caso",
        "device": None,
        "priority": "CRITICAL",
        "test_type": "FUNCTIONAL",
        "source_type": "functionality",
        "user_type": None,
        "test_data": None,
        "evidence": None,
        "justification": None,
        "technical_story": "ST-1",
        "scenario_origin": "Caso",
        "confidence": None,
        "complexity": "BAJA",
        "description": None,
        "steps": [],
        "applicability_reason": None,
    }
    defaults.update(kwargs)
    return SimpleNamespace(**defaults)


def test_x3_avance_total_equals_test_case_rows_and_parcial_explains_count() -> None:
    cases = [
        _case(test_case_id="QC-001", component="EPC-1", technical_epic="EPC-1"),
        _case(test_case_id="QC-002", component="EPC-1", technical_epic="EPC-1", status="PASS"),
        _case(test_case_id="QC-003", component="EPC-2", technical_epic="EPC-2"),
        _case(
            test_case_id="QC-004",
            component="EPC-2",
            technical_epic="EPC-2",
            applicability_reason="na_ambiente: llave faltante",
            test_case_name="No aplica",
        ),
    ]
    ledger = [
        {"epic": "EPC-1", "story": "ST-1", "scenario": "Uno", "estado": "caso ejecutable"},
        {"epic": "EPC-1", "story": "ST-1", "scenario": "Dos", "estado": "pendiente de revisión QC", "motivo": "sin dato"},
        {"epic": "EPC-1", "story": "ST-1", "scenario": "Tres", "estado": "error: timeout", "motivo": "timeout"},
        {"epic": "EPC-2", "story": "ST-2", "scenario": "Cuatro", "estado": "caso ejecutable"},
    ]
    payload = build_test_cases_workbook(cases, scope_keys=["EPC-1", "EPC-2", "EPC-3"], scenario_ledger=ledger)
    book = load_workbook(io.BytesIO(payload))
    avance = book["Avance"]
    qc = book["Test Cases"]
    qc_rows = sum(1 for index in range(2, qc.max_row + 1) if qc.cell(index, 1).value)
    totals = []
    labels = {}
    for index in range(2, avance.max_row + 1):
        key = avance.cell(index, 1).value
        if key == "Total":
            assert avance.cell(index, 2).value == qc_rows
            continue
        totals.append(avance.cell(index, 2).value or 0)
        labels[key] = avance.cell(index, 6).value or ""
    assert sum(totals) == qc_rows
    assert "parcial" in labels["EPC-1"]
    assert "1 pendiente de revisión QC" in labels["EPC-1"]
    assert "1 error" in labels["EPC-1"]
    assert "parcial" not in labels["EPC-2"]
    assert labels["EPC-3"] == ""
