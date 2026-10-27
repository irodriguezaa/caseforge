"""Sprint Testing: B/A/C-style buckets for Technical Epics from Jira filters."""

from tests.conftest import login_as
from app.services.sprint_testing.catalog import get_swf
from app.services.sprint_testing.classify import (
    classify_status,
    is_blocker_priority,
    is_closed_status,
    is_technical_epic,
)
from app.services.sprint_testing.service import aggregate_execution_issues, aggregate_issues, build_sprint_testing


def _issue(key: str, project: str, status: str, issuetype: str = "Technical Epic") -> dict:
    return {
        "key": key,
        "fields": {
            "issuetype": {"name": issuetype},
            "status": {"name": status},
            "project": {"key": project},
            "summary": key,
        },
    }


def test_classify_aliases() -> None:
    assert classify_status("Tareas por hacer") == ("todo", "To Do")
    assert classify_status("To Do") == ("todo", "To Do")
    assert classify_status("Validation") == ("testing", "Validation")
    assert classify_status("Validate QC") == ("testing", "Validate QC")
    assert classify_status("QC Validation") == ("testing", "Validation")
    assert classify_status("QA Validation") == ("testing", "QA Validation")
    assert classify_status("Integration") == ("integration", "Integration")
    assert classify_status("Cancelled") == ("closed", "Canceled")
    assert classify_status("Canceled") == ("closed", "Canceled")
    assert classify_status("Finalizada") == ("closed", "Finalizada")
    assert classify_status("Data Validation") == ("closed", "Data Validation")
    assert classify_status("In Progress") == ("development", "In Progress")
    assert classify_status("Blocked") == ("development", "Blocked")
    assert classify_status("Released", "done") == ("closed", "Released")
    assert classify_status("QA Validation", "done") == ("testing", "QA Validation")


def test_only_technical_epic_is_counted() -> None:
    assert is_technical_epic("Technical Epic")
    assert is_technical_epic(" technical epic ")
    assert not is_technical_epic("Epic")
    assert not is_technical_epic("Story")
    assert not is_technical_epic("QA Bug")


def test_blocker_priority_aliases() -> None:
    assert is_blocker_priority("Blocker") is True
    assert is_blocker_priority("Supone un impedimento") is True
    assert is_blocker_priority("Impedimento") is True
    assert is_blocker_priority("Critical") is False
    assert is_blocker_priority("Major") is False


def test_closed_status_excludes_done_rollout_canceled() -> None:
    assert is_closed_status("Done") is True
    assert is_closed_status("Roll Out") is True
    assert is_closed_status("Cancelled") is True
    assert is_closed_status("Canceled") is True
    assert is_closed_status("Cancelado") is True
    assert is_closed_status("QA Validation") is False
    assert is_closed_status("Pending Resolution") is False
    assert is_closed_status("Released", "done") is True


def test_adt_example_graph1_and_graph2() -> None:
    statuses = (
        ["To Do"] * 2
        + ["In Progress"] * 5
        + ["Blocked"]
        + ["Development"] * 4
        + ["Integration"] * 3
        + ["QA Validation"] * 2
        + ["QC Validation"]
        + ["Roll Out"] * 2
        + ["Done"] * 4
        + ["Canceled"]
    )
    issues = [_issue(f"ADTCL-{index}", "ADTCL", status) for index, status in enumerate(statuses, start=1)]
    issues.append(_issue("ADTCL-999", "ADTCL", "In Progress", issuetype="Story"))
    issues.append(_issue("WINCL-1", "WINCL", "Done"))
    hitss = get_swf("hitss")
    assert hitss is not None
    payload = aggregate_issues("44", "Sprint 44", "117698", hitss, issues)
    adt = next(row for row in payload.programs if row.program_key == "ADTCL")
    assert adt.todo == {"To Do": 2}
    assert adt.development == {"In Progress": 5, "Blocked": 1, "Development": 4}
    assert adt.integration == {"Integration": 3}
    assert adt.testing == {"QA Validation": 2, "Validation": 1}
    assert adt.closed == {"Roll Out": 2, "Done": 4, "Canceled": 1}
    assert adt.total == 25
    assert adt.open == 18
    assert adt.closed_total == 7
    assert adt.consistency_ok is True
    assert (
        sum(adt.development.values())
        + sum(adt.integration.values())
        + sum(adt.testing.values())
        + sum(adt.closed.values())
        + sum(adt.todo.values())
        == adt.total
    )
    assert adt.open + adt.closed_total == adt.total
    win = next(row for row in payload.programs if row.program_key == "WINCL")
    assert win.total == 1
    assert all(row.program_key not in {"AAFCL", "STVCL", "WEBCL"} for row in payload.programs)
    assert payload.other_issue_count == 1
    assert payload.technical_epic_count == 26
    assert payload.consistency_ok is True


def test_done_category_moves_open_to_closed() -> None:
    hitss = get_swf("hitss")
    assert hitss is not None
    issues = [_issue(f"ADTCL-{i}", "ADTCL", "In Progress") for i in range(1, 8)]
    extra = []
    for index in range(8, 15):
        extra.append(
            {
                "key": f"ADTCL-{index}",
                "fields": {
                    "issuetype": {"name": "Technical Epic"},
                    "status": {"name": "Finalizada", "statusCategory": {"key": "done"}},
                    "project": {"key": "ADTCL"},
                    "summary": f"ADTCL-{index}",
                },
            }
        )
    payload = aggregate_issues("44", "Sprint 44", "117698", hitss, issues + extra)
    adt = payload.programs[0]
    assert adt.open == 7
    assert adt.closed_total == 7
    assert adt.closed.get("Finalizada") == 7


def test_zero_programs_are_omitted() -> None:
    hitss = get_swf("hitss")
    assert hitss is not None
    payload = aggregate_issues(
        "44",
        "Sprint 44",
        "117698",
        hitss,
        [_issue("WEBCL-1", "WEBCL", "Integration")],
    )
    assert [row.program_key for row in payload.programs] == ["WEBCL"]


def test_execution_counts_blocker_by_program() -> None:
    hitss = get_swf("hitss")
    assert hitss is not None

    def _exec(key: str, project: str, priority: str) -> dict:
        return {
            "key": key,
            "fields": {
                "issuetype": {"name": "Bug"},
                "status": {"name": "In Progress"},
                "project": {"key": project},
                "priority": {"name": priority},
                "summary": key,
            },
        }

    payload = aggregate_execution_issues(
        "117704",
        hitss,
        [
            _exec("ADTCL-1", "ADTCL", "Blocker"),
            _exec("ADTCL-2", "ADTCL", "Major"),
            _exec("ADTCL-3", "ADTCL", "Impedimento"),
            _exec("WINCL-1", "WINCL", "Critical"),
        ],
    )
    assert payload.filter_id == "117704"
    assert payload.issue_count == 4
    adt = next(row for row in payload.programs if row.program_key == "ADTCL")
    assert adt.blocker == 2
    assert adt.non_blocker == 1
    assert adt.total == 3
    win = next(row for row in payload.programs if row.program_key == "WINCL")
    assert win.blocker == 0
    assert win.non_blocker == 1
    assert [row.key for row in payload.blockers] == ["ADTCL-1", "ADTCL-3"]
    assert payload.blockers[0].summary == "ADTCL-1"
    assert payload.blockers[0].status == "In Progress"
    assert payload.blockers[0].program_key == "ADTCL"
    assert payload.blockers[0].display_name == "ADT / FireTV"
    assert payload.blockers[0].issue_type == "Bug"
    assert all(row.program_key not in {"WEBCL", "AAFCL"} for row in payload.programs)


def test_execution_blockers_grouped_by_program_match_bar_counts() -> None:
    hitss = get_swf("hitss")
    assert hitss is not None

    def _exec(key: str, project: str, priority: str, issuetype: str = "QA Bug") -> dict:
        return {
            "key": key,
            "fields": {
                "issuetype": {"name": issuetype},
                "status": {"name": "In Progress"},
                "project": {"key": project},
                "priority": {"name": priority},
                "summary": key,
            },
        }

    payload = aggregate_execution_issues(
        "117704",
        hitss,
        [
            _exec("AAFCL-1", "AAFCL", "Blocker", "QA Bug"),
            _exec("AAFCL-2", "AAFCL", "Blocker", "QC Bug"),
            _exec("STVCL-1", "STVCL", "Blocker"),
            _exec("STVCL-2", "STVCL", "Blocker"),
            _exec("STVCL-3", "STVCL", "Major"),
        ],
    )
    stale = next(row for row in payload.programs if row.program_key == "AAFCL")
    evo = next(row for row in payload.programs if row.program_key == "STVCL")
    assert stale.blocker == 2
    assert evo.blocker == 2
    assert evo.non_blocker == 1
    stale_keys = [row.key for row in payload.blockers if row.program_key == "AAFCL"]
    evo_keys = [row.key for row in payload.blockers if row.program_key == "STVCL"]
    assert stale_keys == ["AAFCL-1", "AAFCL-2"]
    assert evo_keys == ["STVCL-1", "STVCL-2"]
    assert [row.issue_type for row in payload.blockers if row.program_key == "AAFCL"] == [
        "QA Bug",
        "QC Bug",
    ]


def test_execution_blockers_omit_done_rollout_and_canceled() -> None:
    hitss = get_swf("hitss")
    assert hitss is not None

    def _exec(key: str, status: str, category: str | None = None) -> dict:
        status_field: dict = {"name": status}
        if category:
            status_field["statusCategory"] = {"key": category}
        return {
            "key": key,
            "fields": {
                "issuetype": {"name": "QC Bug"},
                "status": status_field,
                "project": {"key": "ADTCL"},
                "priority": {"name": "Blocker"},
                "summary": key,
            },
        }

    payload = aggregate_execution_issues(
        "117704",
        hitss,
        [
            _exec("ADTCL-1", "QA Validation"),
            _exec("ADTCL-2", "Done"),
            _exec("ADTCL-3", "Roll Out"),
            _exec("ADTCL-4", "Cancelled"),
            _exec("ADTCL-5", "Released", "done"),
            _exec("ADTCL-6", "Pending Resolution"),
        ],
    )
    adt = next(row for row in payload.programs if row.program_key == "ADTCL")
    assert adt.total == 6
    assert adt.blocker == 2
    assert adt.non_blocker == 4
    assert [row.key for row in payload.blockers] == ["ADTCL-1", "ADTCL-6"]


def test_execution_blockers_include_created_principal_status_comment_and_nested() -> None:
    hitss = get_swf("hitss")
    assert hitss is not None

    issue = {
        "key": "ADTCL-2483",
        "changelog": {
            "histories": [
                {
                    "created": "2026-07-28T16:47:56.584-0300",
                    "items": [{"field": "status", "fromString": "To Do", "toString": "Analysis"}],
                },
                {
                    "created": "2026-10-09T21:23:45.759-0300",
                    "items": [{"field": "status", "fromString": "QA Validation", "toString": "Analysis"}],
                },
                {
                    "created": "2026-10-01T10:00:00.000-0300",
                    "author": {"displayName": "Someone Else"},
                    "items": [{"field": "Comment"}],
                },
            ]
        },
        "fields": {
            "issuetype": {"name": "QA Bug"},
            "status": {"name": "Analysis"},
            "project": {"key": "ADTCL"},
            "priority": {"name": "Supone un impedimento"},
            "summary": "Panel de Audio",
            "created": "2026-07-28T16:38:26.826-0300",
            "parent": {"key": "ADTCL-541", "fields": {"summary": "Feature"}},
            "comment": {
                "total": 2,
                "comments": [
                    {
                        "created": "2026-08-01T10:00:00.000-0300",
                        "updated": "2026-08-01T10:00:00.000-0300",
                        "author": {"displayName": "Other"},
                    },
                    {
                        "created": "2026-10-13T13:32:40.365-0300",
                        "updated": "2026-10-13T13:32:40.365-0300",
                        "author": {"displayName": "Itzayana Castro"},
                    },
                ],
            },
            "issuelinks": [
                {
                    "type": {"name": "Blocks", "outward": "blocks"},
                    "outwardIssue": {"key": "QA-5453"},
                },
                {
                    "type": {"name": "Nested for Resolution", "outward": "Nested for Resolution"},
                    "outwardIssue": {"key": "PLYCL-763"},
                },
                {
                    "type": {"name": "Nested for Resolution", "inward": "Nested for Resolution"},
                    "inwardIssue": {"key": "CONCL-1167"},
                },
            ],
        },
    }
    payload = aggregate_execution_issues("117704", hitss, [issue])
    row = payload.blockers[0]
    assert row.key == "ADTCL-2483"
    assert row.priority == "Supone un impedimento"
    assert row.created == "2026-07-28T16:38:26.826-0300"
    assert row.principal == "ADTCL-541"
    assert row.principal_summary == "Feature"
    assert row.status_changed_at == "2026-10-09T21:23:45.759-0300"
    assert row.status_changed_from == "QA Validation"
    assert row.status_changed_to == "Analysis"
    assert row.last_comment_at == "2026-10-13T13:32:40.365-0300"
    assert row.last_comment_author == "Itzayana Castro"
    assert row.nested_issue_keys == ["PLYCL-763", "CONCL-1167"]


def test_options_mark_future_sprints(client) -> None:
    response = client.get("/api/v1/sprint-testing/options")
    assert response.status_code == 200
    body = response.json()
    by_id = {item["id"]: item for item in body["sprints"]}
    assert by_id["44"]["actionable"] is True
    assert by_id["44"]["filter_id"] == "117698"
    assert by_id["44"]["execution_filter_id"] == "117704"
    assert by_id["45"]["actionable"] is True
    assert by_id["45"]["filter_id"] == "117703"
    assert by_id["45"]["execution_filter_id"] == "117796"
    assert by_id["46"]["actionable"] is True
    assert by_id["46"]["filter_id"] == "117859"
    assert by_id["46"]["execution_filter_id"] is None
    assert body["pack_available"] is False
    assert body["pack_swfs"] == []


def test_tester_and_consulta_cannot_read(client) -> None:
    login_as(client, "tester@test.com", "tester-pass")
    denied = client.get("/api/v1/sprint-testing/options")
    assert denied.status_code == 403
    login_as(client, "consulta@test.com", "consulta-pass")
    denied_consulta = client.get("/api/v1/sprint-testing", params={"sprint": "44", "swf": "hitss"})
    assert denied_consulta.status_code == 403


def test_build_paginates_and_dedupes(monkeypatch) -> None:
    from app.services import jira_client

    pages = [
        {
            "issues": [
                _issue("ADTCL-1", "ADTCL", "In Progress"),
                _issue("ADTCL-1", "ADTCL", "In Progress"),
            ],
            "nextPageToken": "page-2",
        },
        {
            "issues": [_issue("ADTCL-2", "ADTCL", "To Do")],
            "nextPageToken": None,
        },
    ]

    class _FakeResponse:
        status_code = 200

        def __init__(self, payload: dict) -> None:
            self._payload = payload

        def json(self) -> dict:
            return self._payload

    class _FakeClient:
        def __enter__(self):
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def get(self, path: str) -> _FakeResponse:
            if path.endswith("117704"):
                return _FakeResponse({"jql": "execution"})
            assert path == "/rest/api/3/filter/117698"
            return _FakeResponse({"jql": "project in (ADTCL, WINCL)"})

        def post(self, path: str, json: dict) -> _FakeResponse:
            assert path == "/rest/api/3/search/jql"
            if json.get("jql") == "execution":
                return _FakeResponse({"issues": [], "nextPageToken": None})
            if json.get("nextPageToken") == "page-2":
                return _FakeResponse(pages[1])
            return _FakeResponse(pages[0])

    monkeypatch.setattr(jira_client, "ensure_authenticated", lambda: None)
    monkeypatch.setattr(jira_client, "_client", lambda: _FakeClient())
    payload = build_sprint_testing("44", "hitss")
    assert payload.programs[0].total == 2
    assert payload.programs[0].open == 2
    assert payload.filter_issue_count == 2


def _exec_issue(
    key: str,
    project: str,
    status: str,
    priority: str,
    summary: str,
    description: object,
    issuetype: str = "QA Bug",
) -> dict:
    return {
        "key": key,
        "fields": {
            "issuetype": {"name": issuetype},
            "status": {"name": status},
            "project": {"key": project},
            "priority": {"name": priority},
            "summary": summary,
            "description": description,
        },
    }


def test_sprint_device_filters_parse_env(monkeypatch) -> None:
    from app.services.sprint_testing.device_filters import sprint_device_filters

    monkeypatch.setenv(
        "SPRINT45_DEVICE_FILTERS",
        "ADTCL=99868,117752; WEBCL=117767; C9085PR=112076,; ROKUPR=,117755",
    )
    mapping = sprint_device_filters("45")
    assert mapping["ADTCL"].qa_filter_id == "99868"
    assert mapping["ADTCL"].qc_filter_id == "117752"
    assert mapping["WEBCL"].filter_ids() == ("117767",)
    assert mapping["C9085PR"].qa_filter_id == "112076"
    assert mapping["C9085PR"].qc_filter_id is None
    assert mapping["ROKUPR"].qa_filter_id is None
    assert mapping["ROKUPR"].qc_filter_id == "117755"
    assert sprint_device_filters("44") == {}
    assert sprint_device_filters("46") == {}


def test_device_qa_qc_filters_own_nested_issues(monkeypatch) -> None:
    from app.services.sprint_testing.catalog import get_sprint
    from app.services.sprint_testing import service as sprint_service

    monkeypatch.setenv("SPRINT45_DEVICE_FILTERS", "ADTCL=99868,117752")

    def fake_fetch(filter_id: str, _fields: list[str], expand: list[str] | None = None):
        del expand
        if filter_id == "99868":
            return [
                _exec_issue("ADTCL-1", "ADTCL", "Analysis", "Blocker", "ADT qa", "x"),
                _exec_issue("GCECL-2190", "GCECL", "Analysis", "Blocker", "nested qa", "x"),
            ]
        if filter_id == "117752":
            return [_exec_issue("GCECL-2177", "GCECL", "Pending Resolution", "Blocker", "nested qc", "x")]
        if filter_id == "117796":
            return [_exec_issue("ATSCL-2717", "ATSCL", "Analysis", "Blocker", "tata", "x")]
        return []

    monkeypatch.setattr(sprint_service, "fetch_raw_issues_by_filter", fake_fetch)
    sprint = get_sprint("45")
    hitss = get_swf("hitss")
    assert sprint is not None and hitss is not None
    label, grouped = sprint_service.collect_execution_issues(sprint, hitss)
    assert "99868" in label and "117752" in label
    adt_keys = {str(issue.get("key")) for issue in grouped["ADTCL"]}
    assert adt_keys == {"ADTCL-1", "GCECL-2177", "GCECL-2190"}
    assert grouped["WEBCL"] == []
    payload = aggregate_execution_issues(label, hitss, grouped=grouped)
    assert [row.key for row in payload.blockers] == ["ADTCL-1", "GCECL-2177", "GCECL-2190"]
    assert all(row.program_key == "ADTCL" for row in payload.blockers)
    assert all(row.key != "ATSCL-2717" for row in payload.blockers)


def test_execution_issue_rows_filter_and_truncate() -> None:
    from app.services.sprint_testing.service import list_execution_issue_rows

    hitss = get_swf("hitss")
    assert hitss is not None
    rows = list_execution_issue_rows(
        hitss,
        [
            _exec_issue(
                "ADTCL-1",
                "ADTCL",
                "In Progress",
                "Blocker",
                "Playback VOD",
                {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "falla"}]}]},
            ),
            _exec_issue("ATSCL-9", "ATSCL", "To Do", "Major", "Otro SWF", "fuera"),
            _exec_issue("WINCL-2", "WINCL", "QA Validation", "Minor", "Layout", "x" * 600, issuetype="QC Bug"),
        ],
    )
    assert [row.key for row in rows] == ["ADTCL-1", "WINCL-2"]
    assert rows[0].device == "ADT / FireTV"
    assert rows[0].summary == "Playback VOD"
    assert rows[0].status == "In Progress"
    assert rows[0].priority == "Blocker"
    assert rows[0].issue_type == "QA Bug"
    assert rows[0].created == ""
    assert rows[0].status_change == ""
    assert rows[1].summary == "Layout"
    assert rows[1].issue_type == "QC Bug"


def test_issues_workbook_headers() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    from app.schemas.sprint_testing import ExecutionIssueRow
    from app.services.sprint_testing.export import ISSUE_HEADERS, build_issues_workbook

    payload = build_issues_workbook(
        [
            ExecutionIssueRow(
                key="ADTCL-1",
                summary="Playback",
                status="In Progress",
                priority="Blocker",
                device="ADT / FireTV",
                program_key="ADTCL",
                issue_type="QA Bug",
                created="28/07/2026 16:38",
                principal="ADTCL-541",
                status_change="09/10/2026 21:23 · QA Validation → Analysis",
                last_comment="13/10/2026 13:32 · Itzayana Castro",
                nested_issues="PLYCL-763, CONCL-1167",
            )
        ]
    )
    book = load_workbook(BytesIO(payload))
    sheet = book.active
    assert [cell.value for cell in sheet[1]] == ISSUE_HEADERS
    assert sheet["A2"].value == "ADTCL-1"
    assert sheet["C2"].value == "QA Bug"
    assert sheet["D2"].value == "Blocker"
    assert sheet["E2"].value == "In Progress"
    assert sheet["F2"].value == "28/07/2026 16:38"
    assert sheet["G2"].value == "ADTCL-541"
    assert sheet["H2"].value == "09/10/2026 21:23 · QA Validation → Analysis"
    assert sheet["I2"].value == "13/10/2026 13:32 · Itzayana Castro"
    assert sheet["J2"].value == "PLYCL-763, CONCL-1167"


def test_rank_blocker_epics_keeps_top_five_and_flags_single_blockers() -> None:
    from app.schemas.sprint_testing import ExecutionBlockerRow
    from app.services.sprint_testing.export import _rank_blocker_epics

    rows: list[ExecutionBlockerRow] = []
    for index, count in enumerate([4, 3, 1, 1, 1, 1, 1], start=1):
        for copy in range(count):
            rows.append(
                ExecutionBlockerRow(
                    key=f"WEBCL-{index}-{copy}",
                    summary="bug",
                    status="Analysis",
                    principal=f"WEBCL-{index}",
                    principal_summary=f"Epic {index}",
                )
            )
    shown, note = _rank_blocker_epics(rows)
    assert [item.key for item in shown] == ["WEBCL-1", "WEBCL-2", "WEBCL-3", "WEBCL-4", "WEBCL-5"]
    assert shown[0].title == "Epic 1"
    assert shown[0].count == 4
    assert note == "+ 2 con 1 blocker"


def test_executive_pptx_has_status_bars_and_blocker_epics() -> None:
    from io import BytesIO

    from pptx import Presentation

    from app.schemas.sprint_testing import (
        ExecutionBlockerRow,
        ExecutionMetrics,
        ExecutionProgramMetrics,
        ProgramMetrics,
        SprintTestingRead,
        SprintTestingSprint,
    )
    from app.services.sprint_testing.export import build_executive_pptx

    payload = SprintTestingRead(
        sprint=SprintTestingSprint(id="44", label="Sprint 44", filter_id="117698"),
        swf="Hitss",
        filter_issue_count=10,
        technical_epic_count=5,
        other_issue_count=0,
        programs=[
            ProgramMetrics(
                program_key="ADTCL",
                display_name="ADT / FireTV",
                total=5,
                todo={"To Do": 1},
                development={"In Progress": 2},
                testing={"QA Validation": 1},
                closed={"Done": 1},
                open=4,
                closed_total=1,
                unclassified={},
                consistency_ok=True,
            )
        ],
        consistency_ok=True,
        execution=ExecutionMetrics(
            filter_id="117704",
            issue_count=3,
            programs=[
                ExecutionProgramMetrics(
                    program_key="ADTCL",
                    display_name="ADT / FireTV",
                    total=3,
                    blocker=1,
                    non_blocker=2,
                )
            ],
            blockers=[
                ExecutionBlockerRow(
                    key="ADTCL-1",
                    summary="Playback VOD",
                    status="In Progress",
                    program_key="ADTCL",
                    display_name="ADT / FireTV",
                    principal="ADTCL-541",
                    principal_summary="Playback en vivo",
                ),
            ],
        ),
    )
    deck = Presentation(BytesIO(build_executive_pptx(payload)))
    assert len(deck.slides) == 2
    assert not any(shape.has_table for slide in deck.slides for shape in slide.shapes)
    status_text = "\n".join(
        shape.text_frame.text for shape in deck.slides[0].shapes if shape.has_text_frame
    )
    blocker_text = "\n".join(
        shape.text_frame.text for shape in deck.slides[1].shapes if shape.has_text_frame
    )
    assert "Estado de Technical Epics" in status_text
    assert "ADT / FireTV" in status_text
    assert "QA Validation" in status_text
    assert "To Do" not in status_text
    assert "ADTCL-541" in blocker_text
    assert "Playback en vivo" in blocker_text
    assert "Observaciones" in blocker_text


def test_export_issues_requires_execution_filter(client) -> None:
    response = client.get("/api/v1/sprint-testing/export/issues", params={"sprint": "46", "swf": "hitss"})
    assert response.status_code == 400
    assert "ejecución" in response.json()["detail"]


def test_tester_cannot_export(client) -> None:
    login_as(client, "tester@test.com", "tester-pass")
    denied = client.get("/api/v1/sprint-testing/export/report", params={"sprint": "44", "swf": "hitss"})
    assert denied.status_code == 403


def test_pack_hidden_without_allowlist(client) -> None:
    denied = client.get("/api/v1/sprint-testing/export/pack", params={"sprint": "44"})
    assert denied.status_code == 403


def test_pack_options_and_zip(client, monkeypatch) -> None:
    from io import BytesIO
    from zipfile import ZipFile

    from openpyxl import load_workbook
    from pptx import Presentation

    from app.services.sprint_testing import service as sprint_service

    monkeypatch.setenv("SPRINT_PACK_EMAILS", "jefe@test.com")
    monkeypatch.setenv("SPRINT_PACK_SWFS", "hitss,neoris")

    options = client.get("/api/v1/sprint-testing/options")
    assert options.status_code == 200
    body = options.json()
    assert body["pack_available"] is True
    assert [item["id"] for item in body["pack_swfs"]] == ["hitss", "neoris"]

    def fake_fetch(filter_id: str, _fields: list[str], expand: list[str] | None = None):
        del expand
        if filter_id == "117698":
            return [
                _issue("ADTCL-1", "ADTCL", "In Progress"),
                _issue("ROKUPR-1", "ROKUPR", "Done"),
            ]
        return [
            _exec_issue("ADTCL-9", "ADTCL", "Analysis", "Blocker", "Hitss blocker", "x"),
            _exec_issue("ADTCL-8", "ADTCL", "Done", "Blocker", "Closed blocker", "x"),
            _exec_issue("ADTCL-7", "ADTCL", "Analysis", "Major", "Not blocker", "x"),
            _exec_issue("ROKUPR-9", "ROKUPR", "Pending Resolution", "Blocker", "Neoris blocker", "x"),
            _exec_issue("ATSCL-1", "ATSCL", "Analysis", "Blocker", "Tata out of pack", "x"),
        ]

    monkeypatch.setattr(sprint_service, "fetch_raw_issues_by_filter", fake_fetch)
    response = client.get("/api/v1/sprint-testing/export/pack", params={"sprint": "44"})
    assert response.status_code == 200
    assert "zip" in (response.headers.get("content-type") or "")
    archive = ZipFile(BytesIO(response.content))
    names = set(archive.namelist())
    pptx_name = next(name for name in names if name.endswith(".pptx"))
    xlsx_name = next(name for name in names if name.endswith(".xlsx"))
    deck = Presentation(BytesIO(archive.read(pptx_name)))
    assert len(deck.slides) == 4
    titles = [
        shape.text_frame.text
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert any("Hitss" in text for text in titles)
    assert any("Neoris" in text for text in titles)
    book = load_workbook(BytesIO(archive.read(xlsx_name)))
    assert book.sheetnames == ["Hitss", "Neoris"]
    hitss_keys = [row[0] for row in book["Hitss"].iter_rows(min_row=2, max_col=1, values_only=True) if row[0]]
    neoris_keys = [row[0] for row in book["Neoris"].iter_rows(min_row=2, max_col=1, values_only=True) if row[0]]
    assert hitss_keys == ["ADTCL-9"]
    assert neoris_keys == ["ROKUPR-9"]
