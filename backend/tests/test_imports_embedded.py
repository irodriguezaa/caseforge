import io

import openpyxl

EMBEDDED_CSV_HEADER = (
    "ID,Grupo Funcional,Caso de Prueba,Tipo de Usuario,Test Steps,Resultado Esperado,"
    "Datos de Prueba,Prioridad,Test Type,Estado\n"
)


def _create_release(client) -> dict:
    return client.post(
        "/api/v1/releases", json={"name": "Claro Video", "version": "8.15", "platform": "tvOS"}
    ).json()


def _upload_csv(client, release_id: int, csv_text: str, sheet_name: str | None = None):
    data = {}
    if sheet_name is not None:
        data["sheet_name"] = sheet_name
    return client.post(
        f"/api/v1/releases/{release_id}/test-cases/import/preview",
        files={"file": ("import.csv", csv_text.encode("utf-8"), "text/csv")},
        data=data,
    )


def test_embedded_matching_line_counts_splits_into_real_steps(client) -> None:
    release = _create_release(client)
    steps = "1. Open app.\n2. Select asset."
    expected = "1. App opens.\n2. Asset plays."
    csv_text = EMBEDDED_CSV_HEADER + (
        f'QC-001,Playback,"Play a VOD asset",Subscriber,"{steps}","{expected}",,Critical,,Pendiente\n'
    )

    response = _upload_csv(client, release["id"], csv_text)

    assert response.status_code == 200
    body = response.json()
    assert body["format"] == "EMBEDDED"
    assert body["valid_count"] == 1
    assert body["error_count"] == 0
    assert body["warnings"] == []
    steps_out = body["valid"][0]["steps"]
    assert len(steps_out) == 2
    assert steps_out[0] == {"step_number": 1, "test_step": "Open app.", "expected_result": "App opens."}
    assert steps_out[1] == {"step_number": 2, "test_step": "Select asset.", "expected_result": "Asset plays."}


def test_embedded_mismatched_line_counts_becomes_single_step_with_warning(client) -> None:
    release = _create_release(client)
    steps = "1. Precondition A.\n2. Precondition B.\n3. Do the action."
    expected = "1. Effect one.\n2. Effect two."
    csv_text = EMBEDDED_CSV_HEADER + (
        f'QC-007,Playback,"Some scenario",Admin,"{steps}","{expected}",,Critical,,Pendiente\n'
    )

    response = _upload_csv(client, release["id"], csv_text)

    assert response.status_code == 200
    body = response.json()
    assert body["valid_count"] == 1
    assert body["error_count"] == 0
    assert len(body["warnings"]) == 1
    assert body["warnings"][0]["test_case_id"] == "QC-007"
    steps_out = body["valid"][0]["steps"]
    assert len(steps_out) == 1
    assert steps_out[0]["step_number"] == 1
    assert steps_out[0]["test_step"] == steps
    assert steps_out[0]["expected_result"] == expected


def test_embedded_duplicate_id_within_file_is_an_error(client) -> None:
    release = _create_release(client)
    csv_text = EMBEDDED_CSV_HEADER + (
        'QC-001,Playback,"Case A",Subscriber,"1. Step",Result,,Critical,,Pendiente\n'
        'QC-001,Playback,"Case B",Subscriber,"1. Step",Result,,Critical,,Pendiente\n'
    )

    response = _upload_csv(client, release["id"], csv_text)

    body = response.json()
    assert body["valid_count"] == 1  # first occurrence is accepted
    assert body["error_count"] == 1  # second occurrence is flagged
    assert "duplicado" in body["errors"][0]["message"]


def test_embedded_datos_de_prueba_folds_into_description(client) -> None:
    release = _create_release(client)
    csv_text = EMBEDDED_CSV_HEADER + (
        'QC-001,Playback,"Case A",Subscriber,"1. Step",Result,"HTTP 500",Critical,,Pendiente\n'
    )

    response = _upload_csv(client, release["id"], csv_text)

    body = response.json()
    assert "HTTP 500" in body["valid"][0]["description"]


def test_flat_and_embedded_are_auto_detected_from_headers(client) -> None:
    release = _create_release(client)

    flat = client.post(
        f"/api/v1/releases/{release['id']}/test-cases/import/preview",
        files={
            "file": (
                "flat.csv",
                (
                    "Test Case ID,Component,Test Case Name,User Type,Step,Test Step,"
                    "Expected Result,Priority,Test Type,Status\n"
                    "QC-001,Playback,Case,Subscriber,1,Do it,It works,Critical,,Pendiente\n"
                ).encode("utf-8"),
                "text/csv",
            )
        },
    )
    assert flat.json()["format"] == "FLAT"

    embedded = _upload_csv(
        client,
        release["id"],
        EMBEDDED_CSV_HEADER + 'QC-002,Playback,"Case",Subscriber,"1. Do it","1. Works",,Critical,,Pendiente\n',
    )
    assert embedded.json()["format"] == "EMBEDDED"


def test_list_sheets_reports_format_per_sheet_and_recommends_one(client) -> None:
    release = _create_release(client)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    unrelated = workbook.create_sheet("Resumen")
    unrelated.append(["Algo", "Otra cosa"])
    unrelated.append(["x", "y"])

    casos = workbook.create_sheet("Casos de Prueba")
    casos.append(
        ["ID", "Grupo Funcional", "Caso de Prueba", "Tipo de Usuario", "Test Steps",
         "Resultado Esperado", "Datos de Prueba", "Prioridad", "Test Type", "Estado"]
    )
    casos.append(["QC-001", "Playback", "Case", "Subscriber", "1. Step", "1. Result", None,
                  "Critical", None, "Pendiente"])

    buffer = io.BytesIO()
    workbook.save(buffer)

    response = client.post(
        f"/api/v1/releases/{release['id']}/test-cases/import/sheets",
        files={
            "file": (
                "workbook.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    names_and_formats = {s["name"]: s["format"] for s in body["sheets"]}
    assert names_and_formats["Resumen"] == "UNKNOWN"
    assert names_and_formats["Casos de Prueba"] == "EMBEDDED"
    assert body["recommended_sheet"] == "Casos de Prueba"


def test_preview_can_target_an_explicit_sheet_overriding_autodetection(client) -> None:
    release = _create_release(client)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    casos = workbook.create_sheet("Casos de Prueba")
    casos.append(
        ["ID", "Grupo Funcional", "Caso de Prueba", "Tipo de Usuario", "Test Steps",
         "Resultado Esperado", "Datos de Prueba", "Prioridad", "Test Type", "Estado"]
    )
    casos.append(["QC-001", "Playback", "From Casos", "Subscriber", "1. Step", "1. Result", None,
                  "Critical", None, "Pendiente"])

    zephyr = workbook.create_sheet("Zephyr Export")
    zephyr.append(
        ["Test Case ID", "Component", "Test Case Name", "User Type", "Step", "Test Step",
         "Expected Result", "Datos de Prueba", "Priority", "Test Type", "Status"]
    )
    zephyr.append(["QC-002", "Playback", "From Zephyr", "Subscriber", 1, "Step", "Result", None,
                   "Critical", None, "Pendiente"])

    buffer = io.BytesIO()
    workbook.save(buffer)

    response = client.post(
        f"/api/v1/releases/{release['id']}/test-cases/import/preview",
        files={
            "file": (
                "workbook.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"sheet_name": "Zephyr Export"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["sheet_name"] == "Zephyr Export"
    assert body["format"] == "FLAT"
    assert body["valid"][0]["test_case_id"] == "QC-002"
