"""QC design: one call per epic, CRN/COV cross, and Excel gaps."""

from __future__ import annotations

import io
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from app.schemas.case_generation import CoverageUnit
from app.services.qc_case_design import (
    LlmQuotaError,
    build_epic_packet,
    call_qc_llm,
    case_errors,
    cases_from_payload,
    coverage_gaps,
    device_context,
    ledger_from,
    parse_story_text,
    pick_best_chat_model,
    qc_model,
    qc_temperature,
    render_user_prompt,
    split_criteria,
    validate_design,
    _design_packet,
)
from app.services.jira_generation import collect_user_stories, field_text
from app.services.test_case_export import build_test_cases_workbook, parcial_estado


def _story_text() -> str:
    return (
        "Feature: Pago\n"
        "Ver el diseño en https://figma.example/pago\n"
        "Background:\n"
        "  Given el usuario tiene sesión\n"
        "# comentario de llegada\n"
        'Scenario: Pagar con "PayPal"\n'
        "  When el usuario selecciona \"PayPal\"\n"
        "  # la llave `home_cdn_config` habilita el botón\n"
        "  Then se muestra \"Pago confirmado\"\n"
    )


def test_background_is_context_and_literals_stay() -> None:
    preamble, background, scenarios = parse_story_text(_story_text())
    assert "figma.example/pago" in preamble
    assert "sesión" in background
    assert len(scenarios) == 1
    title, body = scenarios[0]
    assert title == 'Pagar con "PayPal"'
    assert "`home_cdn_config`" in body
    assert "# la llave" in body
    assert '"Pago confirmado"' in body
    packet = build_epic_packet(
        {
            "key": "WEBCL-1",
            "summary": "Épica pago",
            "children": [
                {
                    "key": "WEBCL-2",
                    "issuetype": "Story",
                    "summary": "Feature pago",
                    "description": _story_text(),
                }
            ],
        },
        [],
    )
    assert packet is not None
    assert [item.title for item in packet.covs] == ['Pagar con "PayPal"']
    prompt = render_user_prompt(packet, device_context("WEB"))
    assert "Sin User Story" in prompt
    assert "`home_cdn_config`" in prompt
    assert "no es un escenario" in prompt
    assert "COV-001 | Pagar con \"PayPal\"" in prompt


def test_user_story_criteria_become_crn() -> None:
    rows = split_criteria("1. El usuario ve el aviso\n2. El usuario puede cerrarlo")
    assert len(rows) == 2
    packet = build_epic_packet(
        {
            "key": "WEBCL-1",
            "summary": "Épica",
            "user_story": {
                "key": "TBRFRE-9",
                "summary": "Aviso de pago",
                "description": "Historia de negocio",
                "acceptance_criteria": "1. El usuario ve el aviso\n2. El usuario puede cerrarlo",
            },
            "children": [
                {
                    "key": "WEBCL-2",
                    "issuetype": "Story",
                    "summary": "Feature",
                    "description": _story_text(),
                }
            ],
        },
        [],
    )
    assert packet is not None
    assert packet.has_user_story
    assert [item.crn_id for item in packet.crns] == ["CRN-001", "CRN-002"]
    nested = split_criteria("1. El usuario ve el aviso\n- el aviso incluye el monto\n2. El usuario puede cerrarlo")
    assert len(nested) == 2
    assert "el aviso incluye el monto" in nested[0][1]


def test_validation_rejects_short_steps_forbidden_actions_and_missing_cov() -> None:
    short = {
        "titulo": "Visualización del aviso",
        "tipo": "HPP",
        "brecha": "ninguna",
        "precondiciones": ["Hay sesión", "La funcionalidad está habilitada"],
        "pasos": [{"accion": "Abrir inicio", "resultado_esperado": "Se muestra inicio"}],
        "cubre_negocio": ["CRN-001"],
        "cubre_tecnico": ["COV-001"],
    }
    assert any("entre 3 y 5 pasos" in item for item in case_errors(short))
    forbidden = {
        **short,
        "pasos": [
            {"accion": "Abrir inicio", "resultado_esperado": "Se muestra inicio"},
            {"accion": "Consumir API de pagos", "resultado_esperado": "Responde 200"},
            {"accion": "Regresar", "resultado_esperado": "Se muestra el home"},
        ],
    }
    assert any("llamada técnica" in item for item in case_errors(forbidden))
    gaps = coverage_gaps(
        {"casos": [], "excluidos": []},
        ["CRN-001"],
        ["COV-001", "COV-002"],
    )
    assert any("COV-001" in item and "COV-002" in item for item in gaps)


def test_quota_aborts_without_retry() -> None:
    response = MagicMock(status_code=429, text='{"error":{"type":"insufficient_quota","message":"no credits"}}')
    client = MagicMock()
    client.post.return_value = response
    with patch("app.services.qc_case_design.httpx.Client") as cls, patch(
        "app.services.qc_case_design.time.sleep"
    ) as sleep:
        cls.return_value.__enter__.return_value = client
        try:
            call_qc_llm("system", "user", 5)
        except LlmQuotaError as exc:
            assert "insufficient_quota" in str(exc)
        else:
            raise AssertionError("debía abortar")
    assert client.post.call_count == 1
    sleep.assert_not_called()


def test_rate_limit_retries_three_times() -> None:
    response = MagicMock(status_code=429, text='{"error":{"code":"rate_limit_exceeded"}}')
    client = MagicMock()
    client.post.return_value = response
    with patch("app.services.qc_case_design.httpx.Client") as cls, patch(
        "app.services.qc_case_design.time.sleep"
    ) as sleep:
        cls.return_value.__enter__.return_value = client
        try:
            call_qc_llm("system", "user", 5)
        except RuntimeError:
            pass
        else:
            raise AssertionError("debía fallar tras los reintentos")
    assert client.post.call_count == 4
    assert sleep.call_count == 3


def test_model_is_not_mini_and_temperature_stays_in_range() -> None:
    with patch("app.services.qc_case_design.settings") as settings, patch(
        "app.services.qc_case_design.account_model_ids", return_value=[]
    ):
        settings.openai_qc_model = "gpt-4o-mini"
        settings.openai_qc_temperature = 0.9
        assert qc_model() == "gpt-4o"
        assert qc_temperature() == 0.3
    assert pick_best_chat_model(
        ["gpt-4o", "gpt-5-mini", "gpt-5.4", "gpt-5.4-pro", "gpt-5-nano", "whisper-1", "gpt-5-pro-2025-10-06"]
    ) == "gpt-5.4-pro"
    assert pick_best_chat_model(["gpt-5.6-luna", "gpt-5.6-terra", "gpt-5.6-sol", "gpt-5.5-pro"]) == "gpt-5.6-sol"
    web = device_context("WEB")
    assert web.dispositivo == "WEB"
    assert "Cursor" in web.modo
    assert "PayPal" in web.metodos


def test_cases_map_covers_tipo_and_brecha() -> None:
    packet = build_epic_packet(
        {
            "key": "WEBCL-1",
            "summary": "Épica",
            "user_story": {
                "key": "TBRFRE-1",
                "summary": "US",
                "acceptance_criteria": "El usuario ve el aviso",
            },
            "children": [
                {
                    "key": "WEBCL-2",
                    "issuetype": "Story",
                    "summary": "Feature",
                    "description": _story_text(),
                }
            ],
        },
        [
            CoverageUnit(
                coverage_id="COV-001",
                role="A",
                behavior="Pagar",
                scenario='Pagar con "PayPal"',
                evidence="WEBCL-2",
                story_key="WEBCL-2",
                jira_key="WEBCL-2",
            )
        ],
    )
    assert packet is not None
    assert packet.covs[0].cov_id == "COV-001"
    payload = {
        "resumen_funcionalidad": "El usuario puede pagar.",
        "casos": [
            {
                "titulo": "Visualización del pago confirmado con PayPal",
                "tipo": "HPP",
                "prioridad": "Blocker",
                "complejidad": "BAJA",
                "precondiciones": ["Hay sesión iniciada", "PayPal está disponible"],
                "pasos": [
                    {"accion": "Abrir el checkout", "resultado_esperado": "Se muestra el checkout"},
                    {"accion": 'Dar clic en "PayPal"', "resultado_esperado": 'Se muestra "PayPal"'},
                    {"accion": "Confirmar el pago", "resultado_esperado": 'Se muestra "Pago confirmado"'},
                ],
                "variantes": [],
                "referencia_tecnica": "`home_cdn_config`",
                "cubre_negocio": ["CRN-001"],
                "cubre_tecnico": ["COV-001"],
                "brecha": "ninguna",
                "riesgo": "El usuario no ve la confirmación",
                "confianza": "alta",
            }
        ],
        "excluidos": [],
    }
    cases = cases_from_payload(packet, payload, "rn.pdf")
    assert cases[0].justification == "covers=CRN-001,COV-001"
    assert cases[0].qc_tipo == "HPP"
    assert cases[0].qc_brecha == "ninguna"
    assert cases[0].qc_complejidad == "BAJA"
    assert cases[0].priority == "BLOCKER"
    assert len(cases[0].steps) == 3


def test_excel_cruce_sheet_highlights_gaps_and_avance_marks_missing_story() -> None:
    payload = build_test_cases_workbook(
        [],
        coverage={
            "cruce_rows": [
                {"epic": "WEBCL-1", "id": "CRN-002", "situacion": "solo_negocio", "casos": "Caso negocio"},
                {"epic": "WEBCL-1", "id": "COV-001", "situacion": "cruce", "casos": "Caso cruce"},
            ],
            "run_warnings": ["Corrida abortada: insufficient_quota"],
        },
        scenario_ledger=[
            {
                "story": "WEBCL-1",
                "epic": "WEBCL-1",
                "scenario": "User Story",
                "estado": "sin_user_story",
                "motivo": "La épica no tiene User Story ligada",
                "case_id": "",
            },
            {
                "story": "WEBCL-2",
                "epic": "WEBCL-1",
                "scenario": "COV-009 Evento interno",
                "estado": "telemetria",
                "motivo": "Solo se verifica en el dashboard",
                "case_id": "",
            },
        ],
    )
    book = load_workbook(io.BytesIO(payload))
    assert "Cruce US - TS" in book.sheetnames
    cruce = book["Cruce US - TS"]
    assert cruce["C2"].value == "solo_negocio"
    assert cruce["A2"].fill.fgColor.rgb.endswith("F4B183")
    assert cruce["C3"].value == "cruce"
    sheet = book["Escenarios de Jira"]
    assert sheet["E3"].value == "Solo se verifica en el dashboard"
    assert parcial_estado("WEBCL-1", [
        {"epic": "WEBCL-1", "estado": "sin_user_story"},
    ]) == "Sin User Story"
    aviso = [book["Avance"].cell(row, 6).value for row in range(1, 6)]
    assert any(item and "insufficient_quota" in str(item) for item in aviso)


def test_product_brief_sends_every_user_story() -> None:
    stories = collect_user_stories(
        {"key": "TBRFRE-1", "parent_key": "BRFRE-1"},
        [
            {"key": "PE-1", "issuetype": "Product Epic"},
            {"key": "PE-2", "issuetype": "Product Epic"},
            {"key": "WEBCL-1586", "issuetype": "Technical Epic"},
        ],
        {
            "PE-1": [
                {"key": "BRFRE-3529", "issuetype": "User Story"},
                {"key": "BRFRE-3511", "issuetype": "User Story"},
            ],
            "PE-2": [
                {"key": "BRFRE-3530", "issuetype": "User Story"},
                {"key": "BRFRE-3502", "issuetype": "User Story"},
                {"key": "WEBCL-9", "issuetype": "Story"},
            ],
        },
    )
    assert [row["key"] for row in stories] == ["BRFRE-3529", "BRFRE-3511", "BRFRE-3530", "BRFRE-3502"]
    packet = build_epic_packet(
        {
            "key": "WEBCL-1586",
            "summary": "Centro de notificaciones",
            "user_stories": [
                {
                    "key": key,
                    "summary": "Historia",
                    "yo_como": "usuario",
                    "quiero": "ver avisos",
                    "para": "decidir",
                    "acceptance_criteria": "1. El usuario ve el aviso",
                }
                for key in ("BRFRE-3529", "BRFRE-3511", "BRFRE-3530", "BRFRE-3502")
            ],
            "children": [
                {
                    "key": "WEBCL-2",
                    "issuetype": "Story",
                    "summary": "Feature",
                    "description": _story_text(),
                }
            ],
        },
        [],
    )
    assert packet is not None
    prompt = render_user_prompt(packet, device_context("WEB"))
    for key in ("BRFRE-3529", "BRFRE-3511", "BRFRE-3530", "BRFRE-3502"):
        assert key in prompt
    assert "Yo como: usuario" in prompt
    assert [item.crn_id for item in packet.crns] == ["CRN-001", "CRN-002", "CRN-003", "CRN-004"]
    rejected = validate_design(
        {
            "resumen_funcionalidad": "x",
            "matriz_cruce": [],
            "casos": [],
            "excluidos": [{"id": "COV-001", "motivo": "otra_epica", "explicacion": "no aplica"}],
        },
        ["CRN-001"],
        ["COV-001"],
    )
    assert any("otra_epica solo aplica" in item for item in rejected)


def test_forty_six_scenarios_stay_in_one_call() -> None:
    body = "\n".join(
        f"Scenario: Escenario {index}\n  When el usuario abre la pantalla {index}\n  Then se muestra el estado {index}"
        for index in range(1, 47)
    )
    packet = build_epic_packet(
        {
            "key": "WEBCL-1586",
            "summary": "Centro",
            "user_story": {"key": "BRFRE-1", "summary": "US", "acceptance_criteria": "1. Ve el aviso"},
            "children": [{"key": "WEBCL-2", "issuetype": "Story", "summary": "Feature", "description": body}],
        },
        [],
    )
    assert packet is not None
    assert len(packet.covs) == 46
    with patch("app.services.qc_case_design._exchange", return_value=({}, {"ok": True}, [])) as exchange:
        _design_packet(packet, device_context("WEB"), 5)
    assert exchange.call_count == 1


def test_validation_gap_keeps_the_required_estado() -> None:
    packet = build_epic_packet(
        {
            "key": "WEBCL-1",
            "summary": "Épica",
            "user_story": {"key": "BRFRE-1", "summary": "US", "acceptance_criteria": "1. Ve el aviso"},
            "children": [
                {"key": "WEBCL-2", "issuetype": "Story", "summary": "Feature", "description": _story_text()}
            ],
        },
        [],
    )
    assert packet is not None
    rows, _accounted = ledger_from(packet, {"casos": [], "excluidos": []}, [], ["respuesta incompleta"])
    assert any(row["estado"] == "sin cobertura (falla de validación)" for row in rows)
    assert "sin cobertura (falla de validación)" in parcial_estado("WEBCL-1", rows)


def test_ordered_acceptance_list_splits_top_level_only() -> None:
    text = field_text(
        {
            "type": "doc",
            "content": [
                {
                    "type": "orderedList",
                    "content": [
                        {
                            "type": "listItem",
                            "content": [
                                {"type": "paragraph", "content": [{"type": "text", "text": "El usuario ve el aviso"}]},
                                {
                                    "type": "orderedList",
                                    "content": [
                                        {
                                            "type": "listItem",
                                            "content": [
                                                {
                                                    "type": "paragraph",
                                                    "content": [{"type": "text", "text": "incluye el monto"}],
                                                }
                                            ],
                                        }
                                    ],
                                },
                            ],
                        },
                        {
                            "type": "listItem",
                            "content": [
                                {"type": "paragraph", "content": [{"type": "text", "text": "El usuario puede cerrarlo"}]}
                            ],
                        },
                    ],
                }
            ],
        }
    )
    rows = split_criteria(text)
    assert len(rows) == 2
    assert "incluye el monto" in rows[0][1]
    assert rows[1][0].startswith("El usuario puede cerrarlo")
