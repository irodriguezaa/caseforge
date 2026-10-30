"""Compara un Excel generado por el motor contra el set dorado.

Uso:  python compare_golden.py golden.json generado.xlsx
Lee las hojas 'Escenarios de Jira' (Historia, Escenario, Estado, ID del caso, Motivo)
y 'Test Cases'. Sale con código 1 si hay faltantes, de más, fragmentados o mal clasificados.
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", s)).strip()


def compare(golden_path, excel_path) -> dict:
    gold = json.loads(Path(golden_path).read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(excel_path, read_only=True)

    tc = {}
    for r in wb["Test Cases"].iter_rows(min_row=2, values_only=True):
        if r[0] and str(r[2]) in gold["epcs"]:
            tc[r[0]] = {"id": r[0], "nombre": r[1], "epc": r[2]}
    by_name = {norm(v["nombre"]): k for k, v in tc.items()}

    def case_id(ref):
        if not ref:
            return None
        m = re.match(r"\s*([A-Z]{2,}-\d+)", str(ref))
        if m and m.group(1) in tc:
            return m.group(1)
        return by_name.get(norm(str(ref).split("·")[-1]))

    s2c = {}
    scen_state = {}
    for r in wb["Escenarios de Jira"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        k = (r[0], norm(r[1]))
        scen_state[k] = norm(r[2])
        cid = case_id(r[3]) if norm(r[2]) == "caso ejecutable" else None
        if cid:
            s2c[k] = cid

    gold_of = {}
    for g in gold["casos"]:
        for c in g["cubre"]:
            gold_of[(c["historia"], norm(c["escenario"]))] = g["id"]
    nocase = {(n["historia"], norm(n["escenario"])): n for n in gold["no_generan_caso"]}

    rep = defaultdict(lambda: defaultdict(list))
    gen_to_gold = defaultdict(set)
    gold_to_gen = defaultdict(set)
    for k, cid in s2c.items():
        if k in gold_of:
            gen_to_gold[cid].add(gold_of[k])
            gold_to_gen[gold_of[k]].add(cid)
        elif k in nocase:
            n = nocase[k]
            rep[n["epc"]]["mal_clasificado"].append(
                f"{k[0]} · {n['escenario']} → {cid} (falla: {n['criterio_que_falla']})"
            )

    for g in gold["casos"]:
        gens = gold_to_gen.get(g["id"], set())
        if not gens:
            rep[g["epc"]]["faltante"].append(f"{g['id']} {g['nombre_referencia']}")
        elif len(gens) > 1:
            rep[g["epc"]]["fragmentado"].append(f"{g['id']} → {', '.join(sorted(gens))}")
    for cid, c in tc.items():
        golds = gen_to_gold.get(cid, set())
        if not golds:
            rep[c["epc"]]["de_mas"].append(f"{cid} {c['nombre']}")
        elif len(golds) > 1:
            rep[c["epc"]]["fusionado"].append(f"{cid} cubre {', '.join(sorted(golds))}")

    fail = False
    lines = [
        f"{'EPC':<12}{'esperado':>9}{'generado':>9}{'faltan':>8}{'de más':>8}{'fragm.':>8}{'mal clas.':>10}"
    ]
    for e in gold["epcs"]:
        r = rep[e]
        gen = sum(1 for c in tc.values() if c["epc"] == e)
        lines.append(
            f"{e:<12}{gold['conteo_esperado'][e]:>9}{gen:>9}{len(r['faltante']):>8}{len(r['de_mas']):>8}"
            f"{len(r['fragmentado']):>8}{len(r['mal_clasificado']):>10}"
        )
        fail |= any(r[x] for x in ("faltante", "de_mas", "fragmentado", "mal_clasificado"))
    for e in gold["epcs"]:
        for kind in ("faltante", "de_mas", "fragmentado", "mal_clasificado", "fusionado"):
            for line in rep[e][kind]:
                lines.append(f"  [{e}] {kind}: {line}")
    lines.append("RESULTADO: " + ("FALLA" if fail else "OK"))
    report = "\n".join(lines)
    return {"ok": not fail, "report": report, "fail": fail}


if __name__ == "__main__":
    result = compare(sys.argv[1], sys.argv[2])
    print(result["report"])
    sys.exit(1 if result["fail"] else 0)
