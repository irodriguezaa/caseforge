"""Publish the official 467 TCs to QCO Test issues. Does not mutate coverage rows."""

from __future__ import annotations

import csv
import json
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.db import SessionLocal
from app.services.zephyr_publish import fingerprint, load_engine_cases, publish_cases

RELEASE_ID = int(os.getenv("PUBLISH_RELEASE_ID", "20"))
OUT_DIR = Path("/app/data/release_notes")
STEM = "CaseForge_Operativas_Zephyr_publish"


def main() -> None:
    with SessionLocal() as session:
        cases = load_engine_cases(session, RELEASE_ID)
        before = {case.id: fingerprint(case) for case in cases}
        records = publish_cases(session, cases)
        session.expire_all()
        after_cases = load_engine_cases(session, RELEASE_ID)
        after = {case.id: fingerprint(case) for case in after_cases}

    unchanged = sum(1 for key, value in before.items() if after.get(key) == value)
    created = sum(1 for row in records if row.resultado == "created")
    errors = sum(1 for row in records if row.resultado == "error")
    duplicates = sum(1 for row in records if row.resultado == "duplicate")
    rejected = errors

    rows = [
        {
            "CaseForge ID": row.caseforge_id,
            "Zephyr ID": row.zephyr_id or "",
            "BRF": row.brf,
            "HN": row.hn,
            "Device/Channel": row.device_channel,
            "Name": row.name,
            "Steps": row.steps,
            "Status": row.status,
            "resultado": row.resultado,
            "detalle": row.detail,
        }
        for row in records
    ]
    summary = {
        "TCs enviados": len(records),
        "creados correctamente": created,
        "errores": errors,
        "duplicados": duplicates,
        "rechazados": rejected,
        "caseforge_sin_modificacion": unchanged == len(before),
        "fingerprints_iguales": unchanged,
        "fingerprints_total": len(before),
    }
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    payload = {"summary": summary, "rows": rows}
    (OUT_DIR / f"{STEM}.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    with (OUT_DIR / f"{STEM}.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "CaseForge ID",
                "Zephyr ID",
                "BRF",
                "HN",
                "Device/Channel",
                "Name",
                "Steps",
                "Status",
                "resultado",
                "detalle",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)
    workbook = Workbook()
    resume = workbook.active
    resume.title = "Resumen"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")
    resume.append(list(summary.keys()))
    resume.append(list(summary.values()))
    sheet = workbook.create_sheet("Publicacion")
    headers = list(rows[0].keys()) if rows else []
    sheet.append(headers)
    for item in rows:
        sheet.append([item[key] for key in headers])
    for sheet_item in (resume, sheet):
        for cell in sheet_item[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap
        sheet_item.freeze_panes = "A2"
    workbook.save(OUT_DIR / f"{STEM}.xlsx")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if errors:
        samples = [row.detail for row in records if row.resultado == "error"][:5]
        print(json.dumps({"error_samples": samples}, ensure_ascii=False))


if __name__ == "__main__":
    main()
