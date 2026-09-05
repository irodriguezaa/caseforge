"""Preview export: matrix → device TCs for the 10 diagnostic BRFs. No persistence."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.db import SessionLocal
from app.models.epc import Epc
from app.services.operativa_engine.coverage_matrix import build_coverage_matrix
from app.services.operativa_engine.matrix_expand import expand_matrix_preview

PATTERNS = {
    "BRF-17442": "HN_CONFLUENCE",
    "BRF-16018": "ALTA_COMPLEJA",
    "BRF-17834": "FUNCIONAL_DEPTH",
    "BRF-17844": "CHANNEL_HIGH",
    "BRF-17465": "NEGATIVE",
    "BRF-17603": "MULTI_BEHAVIOR_HN",
    "BRF-17620": "CHANNEL_SET",
    "BRF-17833": "CHANGE_IMPACT",
    "BRF-17847": "SCOPE_GATE",
    "BRF-17577": "TECH_DEPENDENCY",
}
DEFAULT_OTT = [
    "WEB",
    "AAF",
    "Android",
    "iOS",
    "tvOS",
    "Windows/XBOX",
    "Consolas",
    "Roku",
    "Fire TV",
    "Android TV STV",
]
OUT_DIR = Path("/app/data/release_notes")
STEM = "CaseForge_Operativas_preview_expansion_10BRF"
REPORT_HEADERS = ["BRF", "HN", "behavior", "device", "channel", "generated", "reason"]
SUMMARY_HEADERS = [
    "BRF",
    "filas matriz",
    "filas ejecutables",
    "dispositivos expandidos",
    "TCs preview",
    "QC_REVIEW omitidos",
    "OUT_OF_SCOPE omitidos",
    "warnings",
]


def main() -> None:
    keys = list(PATTERNS)
    with SessionLocal() as session:
        epcs = list(session.scalars(select(Epc).where(Epc.brf_key.in_(keys))).all())
    if not epcs:
        raise SystemExit("No EPCs found for the 10 diagnostic BRFs.")
    for epc in epcs:
        if not epc.dispositivos_aplicables:
            epc.dispositivos_aplicables = list(DEFAULT_OTT)
    matrix = build_coverage_matrix(
        release_id=epcs[0].release_id or 0,
        release_name="OPE-diagnostico-10BRF",
        epcs=epcs,
    )
    matrix.rows = [row for row in matrix.rows if row.brf_key in PATTERNS]
    matrix.row_count = len(matrix.rows)
    preview = expand_matrix_preview(matrix)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "persisted": False,
        "jira": False,
        "zephyr": False,
        "preview_tc_count": preview.preview_tc_count,
        "matrix_row_count": preview.matrix_row_count,
        "warnings": preview.warnings,
        "summary": [item.model_dump() for item in preview.summary],
        "report": [item.model_dump() for item in preview.report],
        "cases": [item.model_dump() for item in preview.cases],
    }
    (OUT_DIR / f"{STEM}.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    with (OUT_DIR / f"{STEM}.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REPORT_HEADERS)
        writer.writeheader()
        for item in preview.report:
            writer.writerow(
                {
                    "BRF": item.brf_key,
                    "HN": item.hn,
                    "behavior": item.behavior,
                    "device": item.device,
                    "channel": item.channel,
                    "generated": item.generated,
                    "reason": item.reason,
                }
            )

    workbook = Workbook()
    resume = workbook.active
    resume.title = "Resumen"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    resume.append(SUMMARY_HEADERS)
    by_brf = {item.brf_key: item for item in preview.summary}
    for brf in PATTERNS:
        item = by_brf.get(brf)
        resume.append(
            [
                brf,
                item.matrix_rows if item else 0,
                item.executable_rows if item else 0,
                item.devices_expanded if item else 0,
                item.preview_tcs if item else 0,
                item.qc_review_omitted if item else 0,
                item.out_of_scope_omitted if item else 0,
                item.warnings if item else 0,
            ]
        )
    report_sheet = workbook.create_sheet("Preview")
    report_sheet.append(REPORT_HEADERS)
    for item in preview.report:
        report_sheet.append(
            [item.brf_key, item.hn, item.behavior, item.device, item.channel, item.generated, item.reason]
        )
    cases_sheet = workbook.create_sheet("TCs_preview")
    case_headers = [
        "BRF",
        "HN/CA",
        "behavior",
        "device",
        "channel",
        "ecosystem",
        "interaction_points",
        "relevant_users",
        "transactional",
        "MDP",
        "EPCs",
        "origin",
        "scope_status",
        "duplicate_risk",
        "test_data",
        "reasoning",
    ]
    cases_sheet.append(case_headers)
    for case in preview.cases:
        cases_sheet.append(
            [
                case.brf_key,
                "; ".join(case.hn_keys),
                case.behavior_title,
                case.device or "—",
                case.channel or "—",
                case.ecosystem or "—",
                "; ".join(case.interaction_points),
                "; ".join(case.relevant_users),
                case.transactional,
                "; ".join(case.mdp),
                "; ".join(case.epc_keys),
                case.origin,
                case.scope_status,
                case.duplicate_risk,
                case.test_data or "",
                case.reasoning or "",
            ]
        )
    for sheet in (resume, report_sheet, cases_sheet):
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
            for cell in row:
                cell.alignment = wrap
                cell.border = thin
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(OUT_DIR / f"{STEM}.xlsx")
    print(
        json.dumps(
            {
                "preview_tc_count": preview.preview_tc_count,
                "matrix_row_count": preview.matrix_row_count,
                "warnings": preview.warnings,
                "summary": [item.model_dump() for item in preview.summary],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
