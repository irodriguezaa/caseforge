"""One-off diagnostic export for the 10 BRF coverage-matrix review. Not wired to generate-cases."""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.db import SessionLocal
from app.models.epc import Epc
from app.services.operativa_engine.coverage_matrix import build_coverage_matrix

PATTERNS = {
    "BRF-17442": "HN_CONFLUENCE",
    "BRF-16018": "ALTA_COMPLEJA",
    "BRF-17834": "FUNCIONAL_DEPTH",
    "BRF-17844": "CHANNEL_HIGH",
    "BRF-17465": "NEGATIVE",
    "BRF-17603": "MULTI_BEHAVIOR_HN",
    "BRF-17620": "CHANNEL_SET",
    "BRF-17833": "CHANGE_IMPACT",
    "BRF-17847": "SCOPE_GATE",
    "BRF-17577": "TECH_DEPENDENCY",
}
DEFAULT_OTT = [
    "WEB",
    "AAF",
    "Android",
    "iOS",
    "tvOS",
    "Windows/XBOX",
    "Consolas",
    "Roku",
    "Fire TV",
    "Android TV STV",
]
OUT_DIR = Path("/app/data/release_notes")
STEM = "CaseForge_Operativas_matriz_diagnostica_10BRF"
HEADERS = [
    "validation_pattern",
    "BRF",
    "EPCs",
    "HN/CA",
    "behavior",
    "behavior_reason",
    "interaction_points",
    "channel",
    "ecosystem",
    "applicable_devices",
    "relevant_users",
    "transactional",
    "MDP",
    "test_data",
    "origin",
    "source",
    "hn_source",
    "scope_status",
    "reasoning",
    "duplicate_risk",
    "duplicate_with",
    "QC_REVIEW",
    "evidence",
]


def _row_dict(pattern: str, row) -> dict:
    return {
        "validation_pattern": pattern,
        "BRF": row.brf_key,
        "EPCs": "; ".join(row.epc_keys),
        "HN/CA": "; ".join(row.hn_keys),
        "behavior": row.behavior_title,
        "behavior_reason": row.behavior_reason or "",
        "interaction_points": "; ".join(row.interaction_points),
        "channel": row.channel or "",
        "ecosystem": row.ecosystem or "",
        "applicable_devices": "; ".join(row.applicable_devices),
        "relevant_users": "; ".join(row.relevant_users),
        "transactional": row.transactional,
        "MDP": "; ".join(row.mdp),
        "test_data": row.test_data or "",
        "origin": row.origin,
        "source": row.source,
        "hn_source": row.hn_source or "",
        "scope_status": row.scope_status,
        "reasoning": row.reasoning or "",
        "duplicate_risk": row.duplicate_risk,
        "duplicate_with": "; ".join(row.duplicate_with),
        "QC_REVIEW": row.origin == "QC_REVIEW",
        "evidence": row.evidence,
    }


def _summary_rows(exported: list[dict], notes: list[str]) -> list[dict]:
    by_brf: dict[str, list[dict]] = defaultdict(list)
    for item in exported:
        by_brf[item["BRF"]].append(item)
    out = []
    for brf, pattern in PATTERNS.items():
        items = by_brf.get(brf, [])
        hn = sorted({key for item in items for key in (item["HN/CA"] or "").split("; ") if key})
        devices = sorted({d for item in items for d in (item["applicable_devices"] or "").split("; ") if d})
        points = sorted({s for item in items for s in (item["interaction_points"] or "").split("; ") if s})
        channels = sorted({item["channel"] for item in items if item["channel"]})
        scopes = sorted({item["scope_status"] for item in items})
        qc = sum(1 for item in items if item["QC_REVIEW"])
        risks = sorted({item["duplicate_risk"] for item in items if item["duplicate_risk"] != "UNIQUE"})
        note_hits = [note for note in notes if note.startswith(brf)]
        out.append(
            {
                "BRF": brf,
                "pattern": pattern,
                "HN/CA": "; ".join(hn) if hn else "—",
                "comportamientos": len(items),
                "scope": "; ".join(scopes) or "—",
                "dispositivos": "; ".join(devices) if devices else "—",
                "canal / punto de interacción": "; ".join(points + channels) if (points or channels) else "—",
                "QC_REVIEW": qc,
                "principales riesgos": "; ".join(risks + note_hits[:3]) or "—",
            }
        )
    return out


def main() -> None:
    keys = list(PATTERNS)
    with SessionLocal() as session:
        epcs = list(session.scalars(select(Epc).where(Epc.brf_key.in_(keys))).all())
    if not epcs:
        raise SystemExit("No EPCs found for the 10 diagnostic BRFs.")
    for epc in epcs:
        if not epc.dispositivos_aplicables:
            epc.dispositivos_aplicables = list(DEFAULT_OTT)
    release_id = epcs[0].release_id or 0
    result = build_coverage_matrix(
        release_id=release_id,
        release_name="OPE-diagnostico-10BRF",
        epcs=epcs,
    )
    by_brf = defaultdict(list)
    for row in result.rows:
        by_brf[row.brf_key].append(row)
    exported = []
    for brf, pattern in PATTERNS.items():
        for row in by_brf.get(brf, []):
            exported.append(_row_dict(pattern, row))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    json_path = OUT_DIR / f"{STEM}.json"
    csv_path = OUT_DIR / f"{STEM}.csv"
    xlsx_path = OUT_DIR / f"{STEM}.xlsx"
    summary = _summary_rows(exported, result.notes)

    json_path.write_text(
        json.dumps(
            {"summary": summary, "notes": result.notes, "rows": exported},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(exported)

    workbook = Workbook()
    resume = workbook.active
    resume.title = "Resumen"
    sum_headers = [
        "BRF",
        "pattern",
        "HN/CA",
        "comportamientos",
        "scope",
        "dispositivos",
        "canal / punto de interacción",
        "QC_REVIEW",
        "principales riesgos",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    resume.append(sum_headers)
    for cell in resume[1]:
        cell.font = header_font
        cell.fill = header_fill
    for item in summary:
        resume.append([item[key] for key in sum_headers])
    matrix = workbook.create_sheet("Matriz")
    matrix.append(HEADERS)
    for cell in matrix[1]:
        cell.font = header_font
        cell.fill = header_fill
    for item in exported:
        matrix.append([item[key] for key in HEADERS])
    for sheet in (resume, matrix):
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
            for cell in row:
                cell.alignment = wrap
                cell.border = thin
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(xlsx_path)
    print(json.dumps({"xlsx": str(xlsx_path), "csv": str(csv_path), "json": str(json_path), "summary": summary}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
