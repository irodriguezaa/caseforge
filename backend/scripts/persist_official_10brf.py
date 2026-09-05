"""Persist official TCs from the frozen 10-BRF matrix. Local DB only. No Jira/Zephyr."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import SessionLocal
from app.models.epc import Epc
from app.models.test_case import TestCase
from app.services.case_persistence import persist_candidates
from app.services.operativa_engine.matrix_official import (
    audit_official_candidates,
    load_approved_matrix,
    materialize_official_candidates,
)

DIAGNOSTIC_KEYS = {
    "BRF-17442",
    "BRF-16018",
    "BRF-17834",
    "BRF-17844",
    "BRF-17465",
    "BRF-17603",
    "BRF-17620",
    "BRF-17833",
    "BRF-17847",
    "BRF-17577",
}
MATRIX_CANDIDATES = [
    Path("/app/data/release_notes/CaseForge_Operativas_matriz_diagnostica_10BRF.json"),
    Path("/Users/rodriguezisr/Documents/caseforge/exports/CaseForge_Operativas_matriz_diagnostica_10BRF.json"),
]
OUT_DIR = Path("/app/data/release_notes")
STEM = "CaseForge_Operativas_TCs_oficiales_10BRF"
HEADERS = [
    "BRF",
    "EPCs",
    "HN/CA",
    "Grupo/Behavior",
    "Canal / Punto de interacción",
    "Ecosystem",
    "Device",
    "User",
    "MDP",
    "Test Case Name",
    "Description",
    "Test Steps",
    "Expected Result",
    "Test Data",
    "Priority",
    "Test Type",
    "Status",
    "Evidence/Notes",
    "Reasoning",
]


def _matrix_path() -> Path:
    for path in MATRIX_CANDIDATES:
        if path.exists():
            return path
    raise SystemExit("Approved diagnostic matrix JSON not found.")


def _delete_diagnostic_engine_cases(db: Session, release_id: int) -> int:
    rows = list(
        db.scalars(
            select(TestCase).where(
                TestCase.release_id == release_id,
                TestCase.generated_by_engine.is_(True),
                TestCase.component.in_(DIAGNOSTIC_KEYS),
            )
        ).all()
    )
    count = len(rows)
    for row in rows:
        db.delete(row)
    if count:
        db.flush()
    return count


def _steps_text(candidate, expected: bool) -> str:
    lines = []
    for step in candidate.steps:
        body = step.expected_result if expected else step.action
        lines.append(f"{step.step_number}. {body}")
    return "\n".join(lines)


def _export_row(candidate) -> dict:
    points = ", ".join(candidate.interaction_points or [])
    channel = ""
    if candidate.device is None and "Email" in (candidate.name or ""):
        channel = "Email"
    if candidate.test_data and "Canal: Email" in candidate.test_data:
        channel = "Email"
    interaction_point = points or channel
    return {
        "BRF": candidate.related_functionality or "",
        "EPCs": candidate.related_jira or "",
        "HN/CA": ", ".join(candidate.hn_keys),
        "Grupo/Behavior": candidate.behavior or candidate.group_id or "",
        "Canal / Punto de interacción": interaction_point,
        "Ecosystem": candidate.ecosystem or "",
        "Device": candidate.device or ("—" if channel == "Email" else ""),
        "User": candidate.user_type or "",
        "MDP": candidate.mdp or "",
        "Test Case Name": candidate.name,
        "Description": candidate.description or "",
        "Test Steps": _steps_text(candidate, False),
        "Expected Result": _steps_text(candidate, True),
        "Test Data": candidate.test_data or "",
        "Priority": candidate.priority or "CRITICAL",
        "Test Type": "FUNCTIONAL",
        "Status": "UNEXECUTED",
        "Evidence/Notes": candidate.evidence or "",
        "Reasoning": candidate.justification or "",
    }


def main() -> None:
    matrix_path = _matrix_path()
    with SessionLocal() as session:
        epcs = list(session.scalars(select(Epc).where(Epc.brf_key.in_(DIAGNOSTIC_KEYS))).all())
        if not epcs:
            raise SystemExit("No diagnostic EPCs in DB.")
        release_id = epcs[0].release_id
        if not release_id:
            raise SystemExit("Diagnostic EPCs are not linked to a QC Release.")
        release_name = "OPE-diagnostico-10BRF"
        matrix = load_approved_matrix(matrix_path, release_id=release_id, release_name=release_name)
        candidates, warnings = materialize_official_candidates(matrix)
        audit = audit_official_candidates(matrix, candidates)
        deleted = _delete_diagnostic_engine_cases(session, release_id)
        persist_candidates(session, release_id, "Operativa", candidates)
        session.commit()

    rows = [_export_row(candidate) for candidate in candidates]

    by_brf = audit["by_brf"]
    summary = []
    for brf in [
        "BRF-17442",
        "BRF-16018",
        "BRF-17834",
        "BRF-17844",
        "BRF-17465",
        "BRF-17603",
        "BRF-17620",
        "BRF-17833",
        "BRF-17847",
        "BRF-17577",
    ]:
        stats = by_brf.get(brf, {})
        summary.append(
            {
                "BRF": brf,
                "TCs generados": stats.get("generated", 0),
                "QC_REVIEW": stats.get("qc_review", 0),
                "OUT_OF_SCOPE": stats.get("out_of_scope", 0),
                "warnings": 0,
            }
        )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "persisted": True,
        "jira": False,
        "zephyr": False,
        "release_id": release_id,
        "deleted_previous_engine_tcs": deleted,
        "tc_count": len(candidates),
        "warnings": warnings,
        "audit": audit,
        "summary": summary,
        "cases": rows,
    }
    (OUT_DIR / f"{STEM}.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    with (OUT_DIR / f"{STEM}.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)

    workbook = Workbook()
    resume = workbook.active
    resume.title = "Resumen"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    resume.append(["BRF", "TCs generados", "QC_REVIEW", "OUT_OF_SCOPE", "warnings"])
    for item in summary:
        resume.append([item["BRF"], item["TCs generados"], item["QC_REVIEW"], item["OUT_OF_SCOPE"], item["warnings"]])
    audit_sheet = workbook.create_sheet("Auditoria")
    audit_sheet.append(["kind", "detail", "BRF", "name"])
    for finding in audit.get("findings") or []:
        audit_sheet.append([finding.get("kind"), finding.get("detail"), finding.get("brf"), finding.get("name")])
    cases_sheet = workbook.create_sheet("TestCases")
    cases_sheet.append(HEADERS)
    for item in rows:
        cases_sheet.append([item[key] for key in HEADERS])
    for sheet in (resume, audit_sheet, cases_sheet):
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
            for cell in row:
                cell.alignment = wrap
                cell.border = thin
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(OUT_DIR / f"{STEM}.xlsx")
    print(
        json.dumps(
            {
                "release_id": release_id,
                "tc_count": len(candidates),
                "deleted": deleted,
                "summary": summary,
                "audit_findings": len(audit.get("findings") or []),
                "findings": audit.get("findings") or [],
                "warnings": warnings,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
