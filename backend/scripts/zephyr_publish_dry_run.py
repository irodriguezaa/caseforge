"""Dry-run CaseForge → Zephyr FLAT mapping for the official 10-BRF lote. No remote writes."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.db import SessionLocal
from app.models.test_case import TestCase
from app.services.zephyr_mapping import dry_run_persisted_cases, preview_zephyr_rows

OFFICIAL_RELEASE_ID = 20

DIAGNOSTIC = {
    "BRF-17442",
    "BRF-16018",
    "BRF-17834",
    "BRF-17844",
    "BRF-17465",
    "BRF-17603",
    "BRF-17620",
    "BRF-17833",
    "BRF-17847",
    "BRF-17577",
}
OUT_DIR = Path("/app/data/release_notes")
STEM = "CaseForge_Operativas_Zephyr_dryrun_10BRF"


def main() -> None:
    with SessionLocal() as session:
        cases = list(
            session.scalars(
                select(TestCase)
                .where(
                    TestCase.release_id == OFFICIAL_RELEASE_ID,
                    TestCase.generated_by_engine.is_(True),
                    TestCase.component.in_(DIAGNOSTIC),
                )
                .options(selectinload(TestCase.steps))
                .order_by(TestCase.test_case_id)
            ).all()
        )
    result = dry_run_persisted_cases(cases)
    preview_rows: list[dict[str, str]] = []
    for case in cases:
        preview_rows.extend(preview_zephyr_rows(case))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "wrote_jira": False,
        "wrote_zephyr": False,
        "model": result.model,
        "summary": {
            "TCs evaluados": result.evaluated,
            "publicables": result.publicable,
            "requieren mapping": result.require_mapping,
            "sin destino": result.without_destination,
            "errores": result.errors,
            "filas FLAT preview": len(preview_rows),
        },
        "notes": result.notes,
        "unsourced_zephyr_fields": result.unsourced_zephyr_fields,
        "mapping": [row.model_dump() for row in result.mapping],
        "cases": [row.model_dump() for row in result.cases],
    }
    (OUT_DIR / f"{STEM}.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    with (OUT_DIR / f"{STEM}.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "caseforge_field",
                "destination_field",
                "transform",
                "required",
                "source",
                "risk",
            ],
        )
        writer.writeheader()
        writer.writerows(payload["mapping"])

    workbook = Workbook()
    mapping_sheet = workbook.active
    mapping_sheet.title = "Mapeo"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")
    mapping_sheet.append(
        ["Campo CaseForge", "Campo destino", "Transformación", "Obligatorio", "Fuente", "Riesgo"]
    )
    for row in result.mapping:
        mapping_sheet.append(
            [
                row.caseforge_field,
                row.destination_field,
                row.transform,
                "sí" if row.required else "no",
                row.source,
                row.risk,
            ]
        )
    summary_sheet = workbook.create_sheet("Resumen")
    summary_sheet.append(["TCs evaluados", "publicables", "requieren mapping", "sin destino", "errores"])
    summary_sheet.append(
        [
            result.evaluated,
            result.publicable,
            result.require_mapping,
            result.without_destination,
            result.errors,
        ]
    )
    summary_sheet.append([])
    summary_sheet.append(["Notas"])
    for note in result.notes:
        summary_sheet.append([note])
    summary_sheet.append([])
    summary_sheet.append(["Campos Zephyr sin fuente CaseForge"])
    for field in result.unsourced_zephyr_fields:
        summary_sheet.append([field])
    cases_sheet = workbook.create_sheet("DryRun_TCs")
    cases_sheet.append(["Test Case ID", "BRF", "Name", "status", "zephyr_step_rows", "reasons"])
    for item in result.cases:
        cases_sheet.append(
            [
                item.test_case_id,
                item.brf,
                item.name,
                item.status,
                item.zephyr_step_rows,
                " | ".join(item.reasons),
            ]
        )
    for sheet in (mapping_sheet, summary_sheet, cases_sheet):
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=sheet.max_column):
            for cell in row:
                cell.alignment = wrap
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 36
        sheet.column_dimensions["C"].width = 48
        sheet.column_dimensions["D"].width = 14
        sheet.column_dimensions["E"].width = 32
        sheet.column_dimensions["F"].width = 40
    workbook.save(OUT_DIR / f"{STEM}.xlsx")
    print(
        json.dumps(
            {
                "wrote_jira": False,
                "wrote_zephyr": False,
                "summary": payload["summary"],
                "notes": result.notes,
                "unsourced_zephyr_fields": result.unsourced_zephyr_fields,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
